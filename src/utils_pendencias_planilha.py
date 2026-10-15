import os
import re
import json
import time
import shutil
from datetime import datetime
from collections import OrderedDict
import tkinter as tk
from tkinter import messagebox, ttk
import customtkinter as ctk
from openpyxl import load_workbook

from utils_os import obter_diretorio_dados

ARQUIVO_PENDENCIAS = "pendencias_planilha.json"


def criar_backup_preventivo(caminho_planilha, prefixo="PRE_GRAVACAO"):
    """
    Cria uma cópia preventiva da planilha na subpasta 'Backup_NEXUS'
    no mesmo diretório da planilha antes de qualquer alteração física.
    """
    if not caminho_planilha or not os.path.exists(caminho_planilha):
        return None
    try:
        pasta_backup = os.path.join(os.path.dirname(caminho_planilha), "Backup_NEXUS")
        os.makedirs(pasta_backup, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        nome_base = os.path.basename(caminho_planilha)
        caminho_bak = os.path.join(pasta_backup, f"BACKUP_{prefixo}_{timestamp}_{nome_base}")
        shutil.copy2(caminho_planilha, caminho_bak)
        return caminho_bak
    except Exception as e:
        print(f"Aviso ao gerar backup preventivo: {e}")
        return None


def normalizar_texto(valor):
    if valor is None:
        return ""
    import unicodedata
    nfkd = unicodedata.normalize("NFKD", str(valor))
    sem_acento = "".join([c for c in nfkd if not unicodedata.combining(c)])
    return re.sub(r"[^A-Z0-9]", "", sem_acento.strip().upper())


def obter_caminho_fila_pendencias():
    pasta_dados = obter_diretorio_dados()
    return os.path.join(pasta_dados, ARQUIVO_PENDENCIAS)


def is_file_writable(caminho):
    """Verifica se o arquivo em disco está acessível para gravação exclusiva (não travado por outro processo)."""
    if not caminho or not os.path.exists(caminho):
        return False
    try:
        with open(caminho, "r+b"):
            pass
        return True
    except (IOError, PermissionError, OSError):
        return False


def carregar_pendencias():
    caminho = obter_caminho_fila_pendencias()
    if not os.path.exists(caminho):
        return []
    try:
        with open(caminho, "r", encoding="utf-8") as f:
            dados = json.load(f)
            if isinstance(dados, list):
                return dados
    except Exception:
        pass
    return []


def salvar_pendencias(lista):
    caminho = obter_caminho_fila_pendencias()
    try:
        os.makedirs(os.path.dirname(caminho), exist_ok=True)
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(lista, f, indent=2, ensure_ascii=False)
        return True
    except Exception:
        return False


def obter_quantidade_pendencias():
    return len(carregar_pendencias())


def adicionar_pendencia(tipo, caminho_planilha, dados_operacao, motivo="Planilha bloqueada por outro usuário"):
    """
    Adiciona uma operação à fila persistente de pendências.
    tipo: 'saida_osme' ou 'transferencia_equipes'
    """
    lista = carregar_pendencias()
    pendencia_id = f"{tipo}_{int(time.time())}_{len(lista) + 1}"
    novo_item = {
        "id": pendencia_id,
        "tipo": tipo,
        "timestamp": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "caminho_planilha": os.path.abspath(caminho_planilha),
        "motivo": str(motivo),
        "tentativas": 1,
        "dados": dados_operacao
    }
    lista.append(novo_item)
    salvar_pendencias(lista)
    return pendencia_id


def remover_pendencia(pendencia_id):
    lista = carregar_pendencias()
    nova_lista = [it for it in lista if it.get("id") != pendencia_id]
    salvar_pendencias(nova_lista)
    return len(nova_lista)


def tentar_gravar_saida_osme(caminho_planilha, equipe, doc_sap, data_hoje, itens):
    """
    Grava uma saída OSME: deleta seriais de ESTOQUE e adiciona em FORNECIDOS.
    Suporta Excel aberto via COM ou arquivo fechado via openpyxl.
    """
    seriais_set = {
        normalizar_texto(s["serial"]): s
        for s in itens
        if s.get("serial") and not s.get("eh_nao_serializado")
    }

    # 0. Backup Preventivo Automático
    criar_backup_preventivo(caminho_planilha, prefixo="PRE_SAIDA_OSME")

    # 1. Tentar via Excel COM aberto localmente
    wb_com = None
    try:
        import win32com.client
        xl = win32com.client.GetActiveObject("Excel.Application")
        for wb in xl.Workbooks:
            try:
                if os.path.abspath(wb.FullName).lower() == os.path.abspath(caminho_planilha).lower():
                    wb_com = wb
                    break
            except Exception:
                pass

        if wb_com is not None:
            # Se estiver aberto como Somente Leitura pelo usuário (porque outra pessoa abriu antes)
            if getattr(wb_com, "ReadOnly", False):
                # Se o arquivo no disco já foi liberado, fecha a visualização local para gravar com openpyxl
                if is_file_writable(caminho_planilha):
                    try:
                        wb_com.Close(SaveChanges=False)
                        wb_com = None
                    except Exception:
                        pass
                else:
                    return False, "Planilha aberta no Excel local em modo 'Somente Leitura' (bloqueada na rede por outro usuário)"
            else:
                # Aberto com permissão de escrita no Excel do operador
                ws_est = None
                ws_forn = None
                for sh in wb_com.Sheets:
                    if sh.Name.upper() == "ESTOQUE":
                        ws_est = sh
                    elif sh.Name.upper() == "FORNECIDOS":
                        ws_forn = sh

                if ws_est is not None:
                    last_est = ws_est.UsedRange.Rows.Count + 1
                    for r in range(last_est, 1, -1):
                        v = str(ws_est.Cells(r, 1).Value or "").strip()
                        if normalizar_texto(v) in seriais_set:
                            ws_est.Rows(r).Delete()

                if ws_forn is not None:
                    next_row = ws_forn.UsedRange.Rows.Count + 1
                    while str(ws_forn.Cells(next_row, 1).Value or "").strip():
                        next_row += 1

                    for s in itens:
                        if s.get("eh_nao_serializado"):
                            continue
                        ws_forn.Cells(next_row, 1).Value = str(s["serial"]).strip()
                        ws_forn.Cells(next_row, 2).Value = str(s["material"]).strip()
                        ws_forn.Cells(next_row, 3).Value = str(s.get("deposito", "CL03")).strip()
                        ws_forn.Cells(next_row, 4).Value = str(equipe).strip()
                        try:
                            ws_forn.Cells(next_row, 5).NumberFormat = "dd/mm/yyyy"
                        except Exception:
                            pass
                        ws_forn.Cells(next_row, 5).Value = _converter_data_para_excel(data_hoje)
                        ws_forn.Cells(next_row, 6).Value = str(doc_sap).strip()
                        next_row += 1

                wb_com.Save()
                return True, "Excel Aberto (win32com)"
    except Exception:
        pass

    # 2. Gravação via openpyxl
    if not is_file_writable(caminho_planilha):
        return False, "Arquivo bloqueado para gravação (em uso por outro usuário)"

    try:
        wb = load_workbook(caminho_planilha, keep_vba=caminho_planilha.lower().endswith(".xlsm"))

        if "ESTOQUE" in wb.sheetnames:
            ws_est = wb["ESTOQUE"]
            linhas_del = []
            for r_idx in range(ws_est.max_row, 1, -1):
                val = str(ws_est.cell(row=r_idx, column=1).value or "").strip()
                if normalizar_texto(val) in seriais_set:
                    linhas_del.append(r_idx)
            for r_del in linhas_del:
                ws_est.delete_rows(r_del, 1)

        if "FORNECIDOS" in wb.sheetnames:
            ws_forn = wb["FORNECIDOS"]
            for s in itens:
                if s.get("eh_nao_serializado"):
                    continue
                ws_forn.append([
                    str(s["serial"]).strip(),
                    str(s["material"]).strip(),
                    str(s.get("deposito", "CL03")).strip(),
                    str(equipe).strip(),
                    _converter_data_para_excel(data_hoje),
                    str(doc_sap).strip()
                ])
                try:
                    ws_forn.cell(row=ws_forn.max_row, column=5).number_format = "dd/mm/yyyy"
                except Exception:
                    pass

        wb.save(caminho_planilha)
        wb.close()
        return True, "openpyxl (Arquivo Direto)"
    except PermissionError:
        return False, "Permissão negada ao salvar (outro usuário está com o arquivo aberto)"
    except Exception as e:
        return False, f"Erro ao gravar planilha: {e}"


def tentar_gravar_transferencia(caminho_planilha, eq_dest, dep_dest, doc_z61, data_hoje, itens):
    """
    Grava uma transferência entre equipes: apaga registros anteriores em FORNECIDOS e insere novos registros.
    """
    seriais_set = set(normalizar_texto(it["serial"]) for it in itens if it.get("serial"))

    # 0. Backup Preventivo Automático
    criar_backup_preventivo(caminho_planilha, prefixo="PRE_TRANSF_EQUIPES")

    # 1. Tentar via Excel COM aberto localmente
    wb_com = None
    try:
        import win32com.client
        xl = win32com.client.GetActiveObject("Excel.Application")
        for wb in xl.Workbooks:
            try:
                if os.path.abspath(wb.FullName).lower() == os.path.abspath(caminho_planilha).lower():
                    wb_com = wb
                    break
            except Exception:
                pass

        if wb_com is not None:
            if getattr(wb_com, "ReadOnly", False):
                if is_file_writable(caminho_planilha):
                    try:
                        wb_com.Close(SaveChanges=False)
                        wb_com = None
                    except Exception:
                        pass
                else:
                    return False, "Planilha aberta no Excel local em modo 'Somente Leitura' (bloqueada na rede)"
            else:
                ws_forn = None
                for sh in wb_com.Sheets:
                    if sh.Name.upper() == "FORNECIDOS":
                        ws_forn = sh
                        break

                if ws_forn is not None:
                    last_row = ws_forn.Cells(ws_forn.Rows.Count, 1).End(3).Row
                    for r in range(last_row, 1, -1):
                        val_serie = str(ws_forn.Cells(r, 1).Value or "")
                        if normalizar_texto(val_serie) in seriais_set:
                            ws_forn.Rows(r).Delete()

                    new_last = ws_forn.Cells(ws_forn.Rows.Count, 1).End(3).Row
                    for it in itens:
                        new_last += 1
                        ws_forn.Cells(new_last, 1).Value = str(it["serial"]).strip()
                        ws_forn.Cells(new_last, 2).Value = str(it["material"]).strip()
                        ws_forn.Cells(new_last, 3).Value = str(dep_dest).strip()
                        ws_forn.Cells(new_last, 4).Value = str(eq_dest).strip()
                        try:
                            ws_forn.Cells(new_last, 5).NumberFormat = "dd/mm/yyyy"
                        except Exception:
                            pass
                        ws_forn.Cells(new_last, 5).Value = _converter_data_para_excel(data_hoje)
                        ws_forn.Cells(new_last, 6).Value = str(doc_z61).strip()

                wb_com.Save()
                return True, "Excel Aberto (win32com)"
    except Exception:
        pass

    # 2. Gravação via openpyxl
    if not is_file_writable(caminho_planilha):
        return False, "Arquivo bloqueado para gravação (em uso por outro usuário)"

    try:
        wb = load_workbook(caminho_planilha, keep_vba=caminho_planilha.lower().endswith(".xlsm"))
        if "FORNECIDOS" in wb.sheetnames:
            ws_forn = wb["FORNECIDOS"]
            idx_ser = None
            for c_idx, val in enumerate(next(ws_forn.iter_rows(max_row=1, values_only=True), ()), start=1):
                if normalizar_texto(val) in ("SERIE", "SERIAL", "SERNR"):
                    idx_ser = c_idx
                    break
            if idx_ser is None:
                idx_ser = 1

            linhas_del = []
            for r_idx, row in enumerate(ws_forn.iter_rows(min_row=2, values_only=True), start=2):
                if row and len(row) >= idx_ser and row[idx_ser - 1]:
                    s_norm = normalizar_texto(row[idx_ser - 1])
                    if s_norm in seriais_set:
                        linhas_del.append(r_idx)

            for r_del in reversed(linhas_del):
                ws_forn.delete_rows(r_del)

            for it in itens:
                ws_forn.append([
                    str(it["serial"]).strip(),
                    str(it["material"]).strip(),
                    str(dep_dest).strip(),
                    str(eq_dest).strip(),
                    _converter_data_para_excel(data_hoje),
                    str(doc_z61).strip()
                ])
                try:
                    ws_forn.cell(row=ws_forn.max_row, column=5).number_format = "dd/mm/yyyy"
                except Exception:
                    pass

        wb.save(caminho_planilha)
        wb.close()
        return True, "openpyxl (Arquivo Direto)"
    except PermissionError:
        return False, "Permissão negada ao salvar (outro usuário está com o arquivo aberto)"
    except Exception as e:
        return False, f"Erro ao gravar planilha: {e}"


def _is_data_valida_helper(data_str):
    if not data_str or str(data_str).strip() in ("—", "-", "None", ""):
        return False
    s = str(data_str).strip()
    if any(p in s.lower() for p in ("fim", "base", "data", "ordem", "caufvd", "lbl", "ctxt")):
        return False
    return bool(re.match(r"^\d{1,2}/\d{1,2}/\d{4}$", s))


def _converter_datetime_helper(data_val):
    if not data_val:
        return None
    if isinstance(data_val, datetime):
        return datetime(data_val.year, data_val.month, data_val.day)
    try:
        from datetime import date
        if isinstance(data_val, date):
            return datetime(data_val.year, data_val.month, data_val.day)
    except Exception:
        pass

    s = str(data_val).strip()
    if any(p in s.lower() for p in ("fim", "base", "data", "ordem", "caufvd", "lbl", "ctxt", "—", "-")):
        return None

    m = re.match(r"^(\d{1,2})[\.\/-](\d{1,2})[\.\/-](\d{4})", s)
    if m:
        d, m_mes, a = m.groups()
        try:
            return datetime(int(a), int(m_mes), int(d))
        except ValueError:
            return None

    m_iso = re.match(r"^(\d{4})[\.\/-](\d{1,2})[\.\/-](\d{1,2})", s)
    if m_iso:
        a, m_mes, d = m_iso.groups()
        try:
            return datetime(int(a), int(m_mes), int(d))
        except ValueError:
            return None

    m2 = re.match(r"^(\d{4})(\d{2})(\d{2})$", s)
    if m2:
        a, m_mes, d = m2.groups()
        try:
            return datetime(int(a), int(m_mes), int(d))
        except ValueError:
            return None

    return None


def tentar_gravar_atualizar_osme(caminho_planilha, itens, col_ordem_idx=7, col_venc_idx=8):
    """
    Grava atualizações de OSME e vencimento na aba CADASTRO da planilha.
    Suporta Excel aberto via COM ou arquivo fechado via openpyxl.
    """
    if not itens:
        return True, "Nenhum item para gravar"

    col_ordem_1based = col_ordem_idx + 1
    col_venc_1based = col_venc_idx + 1

    # 0. Backup Preventivo Automático
    criar_backup_preventivo(caminho_planilha, prefixo="PRE_ATUALIZAR_OSME")

    # 1. Tentar via Excel COM aberto localmente
    wb_com = None
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        xl = win32com.client.GetActiveObject("Excel.Application")
        for wb in xl.Workbooks:
            try:
                if os.path.abspath(wb.FullName).lower() == os.path.abspath(caminho_planilha).lower():
                    wb_com = wb
                    break
            except Exception:
                pass

        if wb_com is not None:
            if getattr(wb_com, "ReadOnly", False):
                if is_file_writable(caminho_planilha):
                    try:
                        wb_com.Close(SaveChanges=False)
                        wb_com = None
                    except Exception:
                        pass
                else:
                    return False, "Planilha aberta no Excel local em modo 'Somente Leitura' (bloqueada na rede)"
            else:
                ws_cad = None
                for sh in wb_com.Sheets:
                    if sh.Name.upper() == "CADASTRO":
                        ws_cad = sh
                        break

                if ws_cad is not None:
                    col_equipe_idx = dados_operacao.get("col_equipe_idx", 4)
                    col_dep_idx = dados_operacao.get("col_dep_idx", 5)
                    col_seg_idx = dados_operacao.get("col_seg_idx", 6)
                    col_ordem_idx = dados_operacao.get("col_ordem_idx", 7)
                    col_venc_idx = dados_operacao.get("col_venc_idx", 8)

                    col_equipe_1based = col_equipe_idx + 1
                    col_dep_1based = col_dep_idx + 1
                    col_seg_1based = col_seg_idx + 1
                    col_ordem_1based = col_ordem_idx + 1
                    col_venc_1based = col_venc_idx + 1
                    mapa_linhas_equipes = {}
                    try:
                        max_r = ws_cad.UsedRange.Rows.Count
                        for r in range(2, max_r + 2):
                            v_eq = ws_cad.Cells(r, col_equipe_1based).Value
                            if v_eq:
                                norm_eq = normalizar_texto(v_eq)
                                if norm_eq and norm_eq not in mapa_linhas_equipes:
                                    mapa_linhas_equipes[norm_eq] = r
                    except Exception:
                        pass

                    for it in itens:
                        norm_it = normalizar_texto(it.get("equipe", ""))
                        r_idx = mapa_linhas_equipes.get(norm_it, it.get("row_idx"))

                        # Validação dupla: checa se a célula contém a equipe
                        val_confere = normalizar_texto(ws_cad.Cells(r_idx, col_equipe_1based).Value)
                        if val_confere != norm_it:
                            achou_linha = False
                            for r_tentativa in range(2, max(300, ws_cad.UsedRange.Rows.Count + 2)):
                                if normalizar_texto(ws_cad.Cells(r_tentativa, col_equipe_1based).Value) == norm_it:
                                    r_idx = r_tentativa
                                    achou_linha = True
                                    break
                            if not achou_linha:
                                continue

                        nova_ordem = str(it.get("ordem_sap") or it.get("ordem_atual") or "").strip()
                        novo_venc = it.get("venc_sap") or it.get("venc_atual")

                        if nova_ordem and nova_ordem not in ("—", "-"):
                            ws_cad.Cells(r_idx, col_ordem_1based).Value = nova_ordem
                        dt_venc = _converter_datetime_helper(novo_venc)
                        if dt_venc is not None:
                            ws_cad.Cells(r_idx, col_venc_1based).Value = dt_venc
                            try:
                                ws_cad.Cells(r_idx, col_venc_1based).NumberFormat = "dd/mm/yyyy"
                            except Exception:
                                pass
                        elif novo_venc and _is_data_valida_helper(novo_venc):
                            ws_cad.Cells(r_idx, col_venc_1based).Value = str(novo_venc).strip()
                            try:
                                ws_cad.Cells(r_idx, col_venc_1based).NumberFormat = "dd/mm/yyyy"
                            except Exception:
                                pass

                    wb_com.Save()
                    return True, "Excel Aberto (win32com)"
    except Exception:
        pass

    # 2. Gravação direta via openpyxl
    if not is_file_writable(caminho_planilha):
        return False, "Arquivo bloqueado para gravação (em uso por outro usuário)"

    try:
        wb = load_workbook(caminho_planilha, keep_vba=caminho_planilha.lower().endswith(".xlsm"))
        if "CADASTRO" not in wb.sheetnames:
            return False, "Aba CADASTRO não encontrada na planilha"

        ws_cad = wb["CADASTRO"]
        col_equipe_idx = dados_operacao.get("col_equipe_idx", 4)
        col_dep_idx = dados_operacao.get("col_dep_idx", 5)
        col_seg_idx = dados_operacao.get("col_seg_idx", 6)
        col_ordem_idx = dados_operacao.get("col_ordem_idx", 7)
        col_venc_idx = dados_operacao.get("col_venc_idx", 8)

        col_equipe_1based = col_equipe_idx + 1
        col_dep_1based = col_dep_idx + 1
        col_seg_1based = col_seg_idx + 1
        col_ordem_1based = col_ordem_idx + 1
        col_venc_1based = col_venc_idx + 1

        mapa_linhas_equipes = {}
        for r in range(2, ws_cad.max_row + 1):
            v_eq = ws_cad.cell(row=r, column=col_equipe_1based).value
            if v_eq:
                norm_eq = normalizar_texto(v_eq)
                if norm_eq and norm_eq not in mapa_linhas_equipes:
                    mapa_linhas_equipes[norm_eq] = r

        for it in itens:
            norm_it = normalizar_texto(it.get("equipe", ""))
            r_idx = mapa_linhas_equipes.get(norm_it, it.get("row_idx"))

            val_confere = normalizar_texto(ws_cad.cell(row=r_idx, column=col_equipe_1based).value)
            if val_confere != norm_it:
                achou_linha = False
                for r_tentativa in range(2, ws_cad.max_row + 1):
                    if normalizar_texto(ws_cad.cell(row=r_tentativa, column=col_equipe_1based).value) == norm_it:
                        r_idx = r_tentativa
                        achou_linha = True
                        break
                if not achou_linha:
                    continue

            nova_ordem = str(it.get("ordem_sap") or it.get("ordem_atual") or "").strip()
            novo_venc = it.get("venc_sap") or it.get("venc_atual")

            if nova_ordem and nova_ordem not in ("—", "-"):
                ws_cad.cell(row=r_idx, column=col_ordem_1based, value=nova_ordem)
            dt_venc = _converter_datetime_helper(novo_venc)
            if dt_venc is not None:
                c_v = ws_cad.cell(row=r_idx, column=col_venc_1based, value=dt_venc)
                c_v.number_format = "dd/mm/yyyy"
            elif novo_venc and _is_data_valida_helper(novo_venc):
                c_v = ws_cad.cell(row=r_idx, column=col_venc_1based, value=str(novo_venc).strip())
                c_v.number_format = "dd/mm/yyyy"

        wb.save(caminho_planilha)
        wb.close()
        return True, "openpyxl (Arquivo Direto)"
    except PermissionError:
        return False, "Permissão negada ao salvar (outro usuário está com o arquivo aberto)"
    except Exception as e:
        return False, f"Erro ao gravar planilha: {e}"


def processar_todas_pendencias():
    """
    Varre a fila de pendências e tenta sincronizar todas as operações cujo arquivo esteja liberado.
    Retorna: (sucessos, total_restantes, detalhes)
    """
    lista = carregar_pendencias()
    if not lista:
        return 0, 0, "Nenhuma pendência na fila."

    itens_restantes = []
    sucessos = 0
    detalhes_sucessos = []

    for item in lista:
        tipo = item.get("tipo")
        caminho = item.get("caminho_planilha")
        dados = item.get("dados", {})

        sucesso = False
        msg_erro = ""

        if tipo == "saida_osme":
            sucesso, msg_metodo = tentar_gravar_saida_osme(
                caminho,
                dados.get("equipe"),
                dados.get("doc_sap"),
                dados.get("data"),
                dados.get("itens", [])
            )
            if sucesso:
                detalhes_sucessos.append(f"Saída OSME {dados.get('equipe')} (Doc {dados.get('doc_sap')}) via {msg_metodo}")
            else:
                msg_erro = msg_metodo

        elif tipo == "transferencia_equipes":
            sucesso, msg_metodo = tentar_gravar_transferencia(
                caminho,
                dados.get("eq_dest"),
                dados.get("dep_dest"),
                dados.get("doc_z61"),
                dados.get("data"),
                dados.get("itens", [])
            )
            if sucesso:
                detalhes_sucessos.append(f"Transferência {dados.get('eq_dest')} (Doc {dados.get('doc_z61')}) via {msg_metodo}")
            else:
                msg_erro = msg_metodo

        elif tipo == "atualizar_osme":
            sucesso, msg_metodo = tentar_gravar_atualizar_osme(
                caminho,
                dados.get("itens", []),
                dados.get("col_ordem_idx", 7),
                dados.get("col_venc_idx", 8)
            )
            if sucesso:
                detalhes_sucessos.append(f"Atualizar OSME ({len(dados.get('itens', []))} equipes) via {msg_metodo}")
            else:
                msg_erro = msg_metodo
        else:
            sucesso = False
            msg_erro = f"Tipo de operação desconhecido: {tipo}"

        if sucesso:
            sucessos += 1
        else:
            item["tentativas"] = item.get("tentativas", 1) + 1
            item["ultimo_erro"] = msg_erro
            itens_restantes.append(item)

    salvar_pendencias(itens_restantes)
    resumo = "; ".join(detalhes_sucessos) if detalhes_sucessos else "Nenhum arquivo liberado no momento."
    return sucessos, len(itens_restantes), resumo


class ModalVisualizarPendencias(ctk.CTkToplevel):
    """Janela moderna para visualização e gerenciamento da fila de pendências da planilha."""
    def __init__(self, parent, callback_apos_sincronizacao=None, paletas=None, obter_tema_atual=None):
        super().__init__(parent)
        self.callback_apos_sincronizacao = callback_apos_sincronizacao
        self.title("⏳ Fila de Sincronização Pendente — Planilha Excel")
        self.geometry("640x520")
        self.minsize(580, 440)
        self.attributes("-topmost", True)
        self.after(300, lambda: self.attributes("-topmost", False))
        self.configure(fg_color="#18181b")

        # Topo
        frame_topo = ctk.CTkFrame(self, fg_color="#78350f", corner_radius=0, height=60)
        frame_topo.pack(side="top", fill="x")
        frame_topo.pack_propagate(False)

        lbl_tit = ctk.CTkLabel(
            frame_topo,
            text="⏳ FILA DE SINCRONIZAÇÃO PENDENTE",
            font=("Segoe UI", 13, "bold"),
            text_color="#fef3c7"
        )
        lbl_tit.pack(side="left", padx=16, pady=6)

        lbl_sub = ctk.CTkLabel(
            frame_topo,
            text="Operações concluídas no SAP aguardando liberação da planilha pelo outro usuário",
            font=("Segoe UI", 10),
            text_color="#fde68a"
        )
        lbl_sub.pack(side="left", padx=16)

        # Status Bar
        self.f_status = ctk.CTkFrame(self, fg_color="#27272a", height=38, corner_radius=0)
        self.f_status.pack(fill="x", padx=16, pady=(12, 6))
        self.lbl_info_status = ctk.CTkLabel(
            self.f_status,
            text="Carregando pendências...",
            font=("Segoe UI", 11),
            text_color="#e4e4e7"
        )
        self.lbl_info_status.pack(side="left", padx=12, pady=6)

        # Scrollable Frame
        self.scroll = ctk.CTkScrollableFrame(self, fg_color="#27272a", corner_radius=8)
        self.scroll.pack(fill="both", expand=True, padx=16, pady=6)

        # Barra de Ações Inferior
        f_botoes = ctk.CTkFrame(self, fg_color="transparent")
        f_botoes.pack(fill="x", padx=16, pady=(8, 14))

        self.btn_sincronizar = ctk.CTkButton(
            f_botoes,
            text="🔄 Sincronizar Agora (Tentar Gravar)",
            font=("Segoe UI", 11, "bold"),
            fg_color="#10b981",
            hover_color="#059669",
            height=38,
            command=self._acao_sincronizar_agora
        )
        self.btn_sincronizar.pack(side="left", fill="x", expand=True, padx=(0, 8))

        self.btn_copiar = ctk.CTkButton(
            f_botoes,
            text="📋 Copiar Tabela",
            font=("Segoe UI", 11),
            fg_color="#0284c7",
            hover_color="#0369a1",
            height=38,
            command=self._acao_copiar_tabela
        )
        self.btn_copiar.pack(side="left", padx=(0, 8))

        btn_fechar = ctk.CTkButton(
            f_botoes,
            text="Fechar",
            width=80,
            height=38,
            fg_color="transparent",
            border_width=1,
            command=self.destroy
        )
        btn_fechar.pack(side="left")

        self._renderizar_itens()

    def _renderizar_itens(self):
        for w in self.scroll.winfo_children():
            w.destroy()

        pendencias = carregar_pendencias()
        qtd = len(pendencias)
        self.lbl_info_status.configure(
            text=f"Total de operações pendentes: {qtd} | O N.E.X.U.S tenta gravar automaticamente a cada 20 segundos"
        )

        if not pendencias:
            f_vazio = ctk.CTkFrame(self.scroll, fg_color="transparent")
            f_vazio.pack(fill="both", expand=True, pady=40)
            ctk.CTkLabel(
                f_vazio,
                text="🎉 Nenhuma pendência encontrada!\nTodas as operações anteriores já foram salvas na planilha.",
                font=("Segoe UI", 12),
                text_color="#10b981",
                justify="center"
            ).pack()
            self.btn_sincronizar.configure(state="disabled")
            self.btn_copiar.configure(state="disabled")
            return

        self.btn_sincronizar.configure(state="normal")
        self.btn_copiar.configure(state="normal")

        for idx, item in enumerate(pendencias):
            tipo = item.get("tipo", "saida_osme")
            dados = item.get("dados", {})
            itens_op = dados.get("itens", [])
            doc_sap = dados.get("doc_sap") or dados.get("doc_z61") or "Sem Doc"
            eq = dados.get("equipe") or dados.get("eq_dest") or "-"
            data_op = dados.get("data") or item.get("timestamp", "-")
            motivo = item.get("motivo") or item.get("ultimo_erro") or "Bloqueio de rede"

            if tipo == "saida_osme":
                tipo_lbl = "Saída OSME"
                tipo_cor = "#0284c7"
            elif tipo == "atualizar_osme":
                tipo_lbl = "Atualizar OSME"
                tipo_cor = "#f59e0b"
            else:
                tipo_lbl = "Transferência Equipes"
                tipo_cor = "#8b5cf6"

            if tipo == "atualizar_osme":
                equipes_lista = [it.get("equipe") for it in itens_op if it.get("equipe")]
                eq = f"{len(equipes_lista)} equipes" if len(equipes_lista) > 1 else (equipes_lista[0] if equipes_lista else "CADASTRO")
                doc_sap = f"{len(itens_op)} novas OSMEs"
                seriais_txt = ", ".join(equipes_lista[:8])
                if len(equipes_lista) > 8:
                    seriais_txt += f"... (+{len(equipes_lista) - 8})"
                label_prefix = "Equipes"
            else:
                seriais_txt = ", ".join([str(it.get("serial")) for it in itens_op if it.get("serial")])
                if not seriais_txt:
                    seriais_txt = f"{len(itens_op)} item(ns) não-serializados"
                label_prefix = "Seriais"

            card = ctk.CTkFrame(self.scroll, fg_color="#3f3f46", corner_radius=6)
            card.pack(fill="x", pady=4, padx=2)

            f_topo_card = ctk.CTkFrame(card, fg_color="transparent")
            f_topo_card.pack(fill="x", padx=10, pady=(6, 2))

            ctk.CTkLabel(
                f_topo_card,
                text=f"#{idx + 1} [{tipo_lbl}] • Equipe: {eq}",
                font=("Segoe UI", 11, "bold"),
                text_color=tipo_cor
            ).pack(side="left")

            ctk.CTkLabel(
                f_topo_card,
                text=f"Doc SAP: {doc_sap} | {data_op}",
                font=("Consolas", 10),
                text_color="#e4e4e7"
            ).pack(side="right")

            f_info = ctk.CTkFrame(card, fg_color="transparent")
            f_info.pack(fill="x", padx=10, pady=(2, 6))

            ctk.CTkLabel(
                f_info,
                text=f"{label_prefix}: {seriais_txt}\nStatus: ⚠️ {motivo}",
                font=("Segoe UI", 10),
                text_color="#fcd34d",
                justify="left"
            ).pack(side="left", anchor="w")

            btn_excluir = ctk.CTkButton(
                f_info,
                text="Descartar",
                width=65,
                height=26,
                font=("Segoe UI", 9),
                fg_color="#ef4444",
                hover_color="#dc2626",
                command=lambda pid=item.get("id"): self._descartar_item(pid)
            )
            btn_excluir.pack(side="right", padx=(6, 0))

    def _descartar_item(self, pendencia_id):
        if messagebox.askyesno("Descartar Pendência", "Deseja realmente remover esta pendência da fila? Ela não será gravada na planilha."):
            remover_pendencia(pendencia_id)
            self._renderizar_itens()
            if self.callback_apos_sincronizacao:
                self.callback_apos_sincronizacao()

    def _acao_sincronizar_agora(self):
        self.btn_sincronizar.configure(text="⏳ Gravando...", state="disabled")
        self.update_idletasks()

        sucessos, restantes, msg = processar_todas_pendencias()
        self._renderizar_itens()

        if sucessos > 0 and restantes == 0:
            messagebox.showinfo(
                "Sincronização Concluída",
                f"🎉 Sucesso total!\n\nForam gravados {sucessos} registro(s) pendente(s) na planilha com êxito!"
            )
        elif sucessos > 0 and restantes > 0:
            messagebox.showinfo(
                "Sincronização Parcial",
                f"✓ {sucessos} operação(ões) gravada(s) com sucesso.\nRestam {restantes} que continuam bloqueadas na rede."
            )
        else:
            messagebox.showwarning(
                "Planilha Ainda Bloqueada",
                "A planilha ainda não foi liberada pelo outro usuário no Excel da rede.\n\n"
                "O N.E.X.U.S continuará tentando gravar automaticamente a cada 20 segundos em segundo plano."
            )

        if self.callback_apos_sincronizacao:
            self.callback_apos_sincronizacao()

    def _acao_copiar_tabela(self):
        pendencias = carregar_pendencias()
        if not pendencias:
            return

        linhas = ["TIPO\tEQUIPE\tDOC_OU_OSME\tVENC_OU_DATA\tDETALHES"]
        for p in pendencias:
            tipo = p.get("tipo", "")
            d = p.get("dados", {})
            if tipo == "atualizar_osme":
                for it in d.get("itens", []):
                    linhas.append(f"ATUALIZAR_OSME\t{it.get('equipe')}\t{it.get('ordem_sap')}\t{it.get('venc_sap')}\tLinha CADASTRO {it.get('row_idx')}")
            else:
                eq = d.get("equipe") or d.get("eq_dest", "")
                doc = d.get("doc_sap") or d.get("doc_z61", "")
                data_op = d.get("data", "")
                for it in d.get("itens", []):
                    s = it.get("serial", "")
                    m = it.get("material", "")
                    linhas.append(f"{tipo.upper()}\t{eq}\t{doc}\t{data_op}\t{s} - {m}")

        texto = "\r\n".join(linhas)
        try:
            self.clipboard_clear()
            self.clipboard_append(texto)
            self.btn_copiar.configure(text="✓ Copiado!", fg_color="#10b981")
            self.after(2000, lambda: self.btn_copiar.configure(text="📋 Copiar Tabela", fg_color="#0284c7"))
        except Exception:
            pass

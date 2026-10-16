# -*- coding: utf-8 -*-
"""
Módulo: Contagem de Depósito (Inventário Rotativo — Contagem Física ⟷ ZM29)
N.E.X.U.S — Núcleo de Estoque e Extração Unificada SAP

Compara a contagem física do depósito (formato ZM29 com 'Estoque Livre' preenchido
ou extração automática do Lote 0 da planilha beta) com o saldo oficial de Estoque Livre
da transação ZM29 do SAP para os depósitos CL03 e CL04.
"""

import os
import io
import re
import json
import threading
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import customtkinter as ctk
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils_os import obter_diretorio_executavel, obter_caminho_config
from utils_logger import log_error, log_info
from utils_ui import show_toast, play_sound

DEFAULT_PALETAS = {
    "Nexus": {
        "bg": "#090d16", "container": "#111827", "texto": "#f8fafc",
        "accent": "#00d2ff", "hover": "#0284c7", "input": "#1a2436",
        "card": "#131d2e", "border": "#1e293b"
    },
    "Azul Delray": {
        "bg": "#0a192f", "container": "#172a45", "texto": "#ffffff",
        "accent": "#38bdf8", "hover": "#0284c7", "input": "#1f3a60",
        "card": "#15243b", "border": "#233d60"
    },
    "Dark Mode": {
        "bg": "#1a202c", "container": "#2d3748", "texto": "#ffffff",
        "accent": "#319795", "hover": "#2c7a7b", "input": "#3a4454",
        "card": "#242d3d", "border": "#3b475a"
    },
    "Light Clean": {
        "bg": "#f8f9fa", "container": "#ffffff", "texto": "#1e293b",
        "accent": "#0284c7", "hover": "#0369a1", "input": "#edf2f7",
        "card": "#f1f5f9", "border": "#cbd5e1"
    }
}
DEFAULT_PALETAS["Conecta Blue"] = DEFAULT_PALETAS["Azul Delray"]


def _limpar_quantidade(val):
    """Converte valores com unidades (ex: '100 PC', '1.000 UN', '2.295 M', '0') em float/int."""
    if val is None or pd.isna(val):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    txt = str(val).strip()
    if not txt:
        return 0.0
    # Remove qualquer unidade de medida do final (letras, ex: PC, UN, M, CJ)
    txt = re.sub(r"[A-Za-z]+$", "", txt).strip()
    # Trata formato brasileiro (1.000 ou 1.000,50)
    if "." in txt and "," in txt:
        txt = txt.replace(".", "").replace(",", ".")
    elif "." in txt:
        partes = txt.split(".")
        # Se for milhar comum do SAP brasileiro tipo 1.000 ou 2.295 com 3 dígitos finais
        if len(partes) == 2 and len(partes[1]) == 3 and partes[1].isdigit():
            txt = partes[0] + partes[1]
        elif len(partes) > 2:
            txt = txt.replace(".", "")
    elif "," in txt:
        txt = txt.replace(",", ".")
    try:
        return float(txt)
    except Exception:
        return 0.0


def _normalizar_codigo(val):
    """Padroniza código do material como string limpa em maiúsculo sem pontuação residual."""
    if val is None or pd.isna(val):
        return ""
    txt = str(val).strip().upper()
    if txt.endswith(".0"):
        txt = txt[:-2]
    return txt


class AbaContagemDeposito:
    def __init__(self, parent_frame, label_status_global=None, paletas=None, obter_tema_atual=None):
        self.parent_frame = parent_frame
        self.label_status_global = label_status_global
        self.paletas = paletas or DEFAULT_PALETAS
        self.obter_tema_atual = obter_tema_atual or (lambda: "Nexus")

        self.deposito_atual = "CL03"
        self.caminho_zm29 = ""
        self.caminho_controle = ""
        self.caminho_beta_padrao = self._obter_caminho_controle_padrao()
        
        self.itens_comparados = []  # Lista de dicts com todos os itens comparados
        self.itens_exibidos = []    # Lista filtrada exibida na Treeview
        self.filtro_ativo = "TODOS" # TODOS, DIVERGENTES, FALTAS, SOBRAS, OK
        self.processando = False

        self._construir_interface()
        self.aplicar_tema(self.obter_tema_atual())

    def _cores(self, tema_nome):
        return self.paletas.get(tema_nome, self.paletas.get("Nexus", DEFAULT_PALETAS["Nexus"]))

    def _status(self, msg):
        if self.label_status_global:
            try:
                self.label_status_global.configure(text=msg)
            except Exception:
                pass
        try:
            self.lbl_status_card.configure(text=msg)
        except Exception:
            pass

    def _status_threadsafe(self, msg):
        self.parent_frame.after(0, lambda: self._status(msg))

    def _obter_caminho_controle_padrao(self):
        """Busca caminho salvo em config_lotes.json ou config_controle_material.json."""
        for cfg_nome in ("config_lotes.json", "config_controle_material.json"):
            caminho_cfg = os.path.join(obter_diretorio_executavel(), "data", cfg_nome)
            if os.path.exists(caminho_cfg):
                try:
                    with open(caminho_cfg, "r", encoding="utf-8") as f:
                        data = json.load(f)
                        p = data.get("caminho") or data.get("caminho_planilha")
                        if p and os.path.exists(p):
                            return p
                except Exception:
                    pass
        # Fallback para rede padrão ENEL / Elpa
        fallback = r"X:\obras\Elpa - Almox CR Mauá\CONTROLES\CONTROLE_SERIAIS.vBeta.xlsm"
        if os.path.exists(fallback):
            return fallback
        return ""

    def _construir_interface(self):
        # Container principal
        self.container = ctk.CTkFrame(self.parent_frame, corner_radius=10)
        self.container.pack(fill=tk.BOTH, expand=True, padx=16, pady=12)

        # 1. CABEÇALHO
        f_header = ctk.CTkFrame(self.container, fg_color="transparent")
        f_header.pack(fill=tk.X, padx=16, pady=(12, 6))

        lbl_tit = ctk.CTkLabel(
            f_header,
            text="🏢 Contagem de Depósito (Inventário Rotativo — Contagem Física ⟷ ZM29)",
            font=("Segoe UI", 16, "bold")
        )
        lbl_tit.pack(side=tk.LEFT)

        self.lbl_sub_header = ctk.CTkLabel(
            f_header,
            text="• Cruzamento instantâneo do Saldo SAP ZM29 vs. Estoque Livre Contado no Depósito",
            font=("Segoe UI", 11),
            text_color="#94a3b8"
        )
        self.lbl_sub_header.pack(side=tk.LEFT, padx=(10, 0), pady=(3, 0))

        # 2. CARD DE CONFIGURAÇÃO E ENTRADAS DE DADOS
        self.card_top = ctk.CTkFrame(self.container, corner_radius=8)
        self.card_top.pack(fill=tk.X, padx=16, pady=(4, 8))

        # Linha 1: Seletor de Depósito, Lote e Puxar Lote 0 Beta
        f_top_ctrl = ctk.CTkFrame(self.card_top, fg_color="transparent")
        f_top_ctrl.pack(fill=tk.X, padx=14, pady=(10, 6))

        lbl_dep = ctk.CTkLabel(f_top_ctrl, text="Depósito Ativo:", font=("Segoe UI", 12, "bold"))
        lbl_dep.pack(side=tk.LEFT, padx=(0, 8))

        self.btn_dep_cl03 = ctk.CTkButton(
            f_top_ctrl,
            text="🏢 CL03",
            width=90,
            height=30,
            font=("Segoe UI", 11, "bold"),
            command=lambda: self._selecionar_deposito("CL03")
        )
        self.btn_dep_cl03.pack(side=tk.LEFT, padx=(0, 6))

        self.btn_dep_cl04 = ctk.CTkButton(
            f_top_ctrl,
            text="🏢 CL04",
            width=90,
            height=30,
            font=("Segoe UI", 11, "bold"),
            command=lambda: self._selecionar_deposito("CL04")
        )
        self.btn_dep_cl04.pack(side=tk.LEFT, padx=(0, 16))

        lbl_centro = ctk.CTkLabel(f_top_ctrl, text="Centro:", font=("Segoe UI", 12, "bold"))
        lbl_centro.pack(side=tk.LEFT, padx=(0, 6))

        self.ent_centro = ctk.CTkEntry(f_top_ctrl, width=70, height=30, font=("Segoe UI", 11, "bold"))
        self.ent_centro.pack(side=tk.LEFT, padx=(0, 16))
        self.ent_centro.insert(0, "EMPR")
        self.ent_centro.configure(state="disabled")

        # Botão de extração automática de Lote 0 da planilha beta
        self.btn_extrair_lote0 = ctk.CTkButton(
            f_top_ctrl,
            text="📥 Puxar Lote 0 da Planilha Beta",
            height=30,
            font=("Segoe UI", 11, "bold"),
            command=self._extrair_lote0_planilha_beta
        )
        self.btn_extrair_lote0.pack(side=tk.LEFT, padx=(0, 10))

        # Botão principal de Comparar
        self.btn_comparar = ctk.CTkButton(
            f_top_ctrl,
            text="⚡ Cruzar Contagem ⟷ ZM29",
            height=32,
            font=("Segoe UI", 12, "bold"),
            command=self.executar_comparacao
        )
        self.btn_comparar.pack(side=tk.RIGHT)

        # Linha 2: Entradas de arquivos (ZM29 e Contagem Física)
        f_arquivos = ctk.CTkFrame(self.card_top, fg_color="transparent")
        f_arquivos.pack(fill=tk.X, padx=14, pady=(0, 10))

        # Arquivo 1: ZM29 SAP
        f_zm = ctk.CTkFrame(f_arquivos, fg_color="transparent")
        f_zm.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))

        lbl_zm = ctk.CTkLabel(f_zm, text="1. Bruto ZM29 (Exportação SAP / Saldo Sistema):", font=("Segoe UI", 11, "bold"))
        lbl_zm.pack(anchor="w", pady=(0, 2))

        f_zm_box = ctk.CTkFrame(f_zm, fg_color="transparent")
        f_zm_box.pack(fill=tk.X)

        self.ent_zm29 = ctk.CTkEntry(f_zm_box, height=28, placeholder_text="Selecione o arquivo da transação ZM29 (.xlsx)...")
        self.ent_zm29.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 6))

        self.btn_sel_zm29 = ctk.CTkButton(
            f_zm_box,
            text="📂 Buscar ZM29",
            width=120,
            height=28,
            font=("Segoe UI", 11),
            command=self._selecionar_zm29
        )
        self.btn_sel_zm29.pack(side=tk.RIGHT)

        # Arquivo 2: Contagem Física Depósito (Formato ZM29 com 'Estoque Livre')
        f_ctrl = ctk.CTkFrame(f_arquivos, fg_color="transparent")
        f_ctrl.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(8, 0))

        lbl_ctrl = ctk.CTkLabel(f_ctrl, text="2. Contagem Física Depósito (Formato ZM29 / 'Estoque Livre'):", font=("Segoe UI", 11, "bold"))
        lbl_ctrl.pack(anchor="w", pady=(0, 2))

        f_ctrl_box = ctk.CTkFrame(f_ctrl, fg_color="transparent")
        f_ctrl_box.pack(fill=tk.X)

        self.ent_controle = ctk.CTkEntry(
            f_ctrl_box,
            height=28,
            placeholder_text="Planilha formato ZM29 ou com aba Estoque Livre preenchida..."
        )
        self.ent_controle.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 6))

        self.btn_sel_controle = ctk.CTkButton(
            f_ctrl_box,
            text="📂 Buscar Contagem",
            width=135,
            height=28,
            font=("Segoe UI", 11),
            command=self._selecionar_controle
        )
        self.btn_sel_controle.pack(side=tk.RIGHT)

        # 3. CARDS DE INDICADORES / KPIS DO ROTATIVO
        self.f_kpis = ctk.CTkFrame(self.container, fg_color="transparent")
        self.f_kpis.pack(fill=tk.X, padx=16, pady=(0, 8))

        self.card_total = self._criar_card_kpi(self.f_kpis, "📦 Total Itens", "0", "#38bdf8")
        self.card_ok = self._criar_card_kpi(self.f_kpis, "🟢 Bateu 100%", "0", "#10b981")
        self.card_faltou = self._criar_card_kpi(self.f_kpis, "🔴 Faltou (Déficit)", "0", "#ef4444")
        self.card_sobrou = self._criar_card_kpi(self.f_kpis, "🟡 Sobrando (Excesso)", "0", "#f59e0b")
        self.card_acuracia = self._criar_card_kpi(self.f_kpis, "🎯 Acuracidade", "0.0%", "#a855f7")

        # 4. BARRA DE FILTROS RÁPIDOS E BUSCA
        self.f_filtros = ctk.CTkFrame(self.container, fg_color="transparent")
        self.f_filtros.pack(fill=tk.X, padx=16, pady=(0, 6))

        # Campo de Busca
        self.ent_busca = ctk.CTkEntry(
            self.f_filtros,
            placeholder_text="🔍 Filtrar por código ou descrição em tempo real...",
            height=30,
            width=320,
            font=("Segoe UI", 11)
        )
        self.ent_busca.pack(side=tk.LEFT, padx=(0, 12))
        self.ent_busca.bind("<KeyRelease>", lambda e: self._aplicar_filtro_busca())

        # Botões de Filtros
        self.btn_filtro_todos = self._criar_btn_filtro("Todos (0)", "TODOS")
        self.btn_filtro_div = self._criar_btn_filtro("🚨 Divergências (0)", "DIVERGENTES")
        self.btn_filtro_falta = self._criar_btn_filtro("🔴 Faltas (0)", "FALTAS")
        self.btn_filtro_sobra = self._criar_btn_filtro("🟡 Sobras (0)", "SOBRAS")
        self.btn_filtro_ok = self._criar_btn_filtro("🟢 OK (0)", "OK")

        # 5. TABELA TREEVIEW COM ROLAGEM
        self.f_tabela_container = ctk.CTkFrame(self.container, corner_radius=8)
        self.f_tabela_container.pack(fill=tk.BOTH, expand=True, padx=16, pady=(0, 8))

        colunas = (
            "material", "descricao", "deposito", "centro",
            "umb", "qtd_sap", "qtd_fisico", "diferenca", "status"
        )
        self.tree = ttk.Treeview(
            self.f_tabela_container,
            columns=colunas,
            show="headings",
            selectmode="extended",
            style="Deposito.Treeview"
        )

        headers = [
            ("material", "Material", 100, "center"),
            ("descricao", "Descrição do Material", 300, "w"),
            ("deposito", "Depósito", 75, "center"),
            ("centro", "Centro", 70, "center"),
            ("umb", "UMB", 60, "center"),
            ("qtd_sap", "Qtd SAP (ZM29)", 125, "center"),
            ("qtd_fisico", "Qtd Físico (Estoque Livre)", 140, "center"),
            ("diferenca", "Diferença", 95, "center"),
            ("status", "Status", 130, "center"),
        ]

        for col_id, col_text, col_width, col_anchor in headers:
            self.tree.heading(col_id, text=col_text, command=lambda c=col_id: self._ordenar_coluna(c, False))
            self.tree.column(col_id, width=col_width, anchor=col_anchor)

        scroll_y = ttk.Scrollbar(self.f_tabela_container, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(self.f_tabela_container, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)

        scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # Tags visuais de linha
        self.tree.tag_configure("tag_ok", foreground="#10b981")
        self.tree.tag_configure("tag_falta", foreground="#ef4444")
        self.tree.tag_configure("tag_sobra", foreground="#f59e0b")
        self.tree.tag_configure("tag_sem_sap", foreground="#a855f7")

        # 6. RODAPÉ DE AÇÕES
        f_rodape = ctk.CTkFrame(self.container, fg_color="transparent")
        f_rodape.pack(fill=tk.X, padx=16, pady=(0, 10))

        self.btn_exportar = ctk.CTkButton(
            f_rodape,
            text="💾 Exportar Relatório Excel (.xlsx)",
            font=("Segoe UI", 11, "bold"),
            height=34,
            command=self.exportar_excel
        )
        self.btn_exportar.pack(side=tk.LEFT, padx=(0, 8))

        self.btn_copiar_resumo = ctk.CTkButton(
            f_rodape,
            text="📋 Copiar Resumo",
            font=("Segoe UI", 11),
            height=34,
            command=self.copiar_resumo
        )
        self.btn_copiar_resumo.pack(side=tk.LEFT, padx=(0, 8))

        self.btn_limpar = ctk.CTkButton(
            f_rodape,
            text="🗑️ Limpar Dados",
            font=("Segoe UI", 11),
            height=34,
            command=self.limpar_tela
        )
        self.btn_limpar.pack(side=tk.LEFT)

        self.lbl_status_card = ctk.CTkLabel(
            f_rodape,
            text="Pronto para auditar. Selecione o arquivo ZM29 e clique em Cruzar.",
            font=("Segoe UI", 11)
        )
        self.lbl_status_card.pack(side=tk.RIGHT, padx=10)

        self._atualizar_botoes_deposito()

    def _criar_card_kpi(self, parent, titulo, valor_padrao, cor_valor):
        card = ctk.CTkFrame(parent, corner_radius=8, height=65)
        card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=4)
        card.pack_propagate(False)

        lbl_tit = ctk.CTkLabel(card, text=titulo, font=("Segoe UI", 10, "bold"), text_color="#94a3b8")
        lbl_tit.pack(anchor="w", padx=12, pady=(6, 0))

        lbl_val = ctk.CTkLabel(card, text=valor_padrao, font=("Segoe UI", 18, "bold"), text_color=cor_valor)
        lbl_val.pack(anchor="w", padx=12, pady=(0, 6))

        card.lbl_val = lbl_val
        card.lbl_tit = lbl_tit
        return card

    def _criar_btn_filtro(self, texto, tipo):
        btn = ctk.CTkButton(
            self.f_filtros,
            text=texto,
            height=30,
            font=("Segoe UI", 11),
            command=lambda: self._trocar_filtro(tipo)
        )
        btn.pack(side=tk.LEFT, padx=(0, 6))
        return btn

    def _selecionar_deposito(self, dep):
        self.deposito_atual = dep
        self._atualizar_botoes_deposito()
        if self.itens_comparados:
            self._status(f"Depósito alterado para {dep}. Clique em Cruzar para reprocessar.")

    def _atualizar_botoes_deposito(self):
        c = self._cores(self.obter_tema_atual())
        if self.deposito_atual == "CL03":
            self.btn_dep_cl03.configure(fg_color=c.get("accent", "#00d2ff"), text_color="#090d16" if "Light" not in self.obter_tema_atual() else "#ffffff")
            self.btn_dep_cl04.configure(fg_color=c.get("input", "#1a2436"), text_color=c["texto"])
        else:
            self.btn_dep_cl04.configure(fg_color=c.get("accent", "#00d2ff"), text_color="#090d16" if "Light" not in self.obter_tema_atual() else "#ffffff")
            self.btn_dep_cl03.configure(fg_color=c.get("input", "#1a2436"), text_color=c["texto"])

    def _selecionar_zm29(self):
        caminho = filedialog.askopenfilename(
            title="Selecione a Planilha ZM29 (Exportação SAP)",
            filetypes=[("Arquivos Excel / CSV", "*.xlsx;*.xls;*.csv"), ("Todos os arquivos", "*.*")],
            parent=self.parent_frame.winfo_toplevel()
        )
        if caminho:
            self.caminho_zm29 = caminho
            self.ent_zm29.delete(0, tk.END)
            self.ent_zm29.insert(0, caminho)
            self._status(f"ZM29 selecionada: {os.path.basename(caminho)}")

    def _selecionar_controle(self):
        caminho = filedialog.askopenfilename(
            title="Selecione a Planilha de Contagem Física (Formato ZM29 / Estoque Livre)",
            filetypes=[("Arquivos Excel / CSV", "*.xlsx;*.xlsm;*.xls;*.csv"), ("Todos os arquivos", "*.*")],
            parent=self.parent_frame.winfo_toplevel()
        )
        if caminho:
            self.caminho_controle = caminho
            self.ent_controle.delete(0, tk.END)
            self.ent_controle.insert(0, caminho)
            self._status(f"Contagem física selecionada: {os.path.basename(caminho)}")

    def _agendar_ui(self, func):
        """Executa callback na thread principal de forma resiliente."""
        try:
            self.parent_frame.after(0, func)
        except Exception:
            try:
                func()
            except Exception:
                pass

    def _extrair_lote0_nucleo(self, caminho_beta, dep_alvo, lote_alvo):
        """
        Lógica de leitura da planilha de seriais (aba ESTOQUE) e consolidação
        das quantidades por material para o depósito e lote selecionados.
        Gera e salva a planilha no padrão ZM29 com a coluna 'Estoque Livre'.
        Retorna: (caminho_arquivo_gerado, dict_dados_agrupados)
        """
        with open(caminho_beta, "rb") as f_raw:
            buf = io.BytesIO(f_raw.read())

        dados_agrupados = {}  # { cod: {"desc": str, "umb": str, "qtd": int} }
        try:
            wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)
            if "ESTOQUE" not in wb.sheetnames:
                raise RuntimeError("Aba 'ESTOQUE' não encontrada na planilha de seriais selecionada.")

            ws = wb["ESTOQUE"]
            header_row = next(ws.iter_rows(max_row=1, values_only=True), ())
            idx_mat, idx_desc, idx_dep, idx_lote, idx_umb = None, None, None, None, None

            for i, col_name in enumerate(header_row):
                cn = str(col_name or "").strip().upper()
                if cn in ("MATERIAL", "MAT", "CODIGO", "MATNR") and idx_mat is None:
                    idx_mat = i
                elif any(k in cn for k in ("DESCRI", "TEXTO", "DENOMINA")) and idx_desc is None:
                    idx_desc = i
                elif cn in ("DEPOSITO", "DEP", "LGORT") and idx_dep is None:
                    idx_dep = i
                elif cn in ("LOTE", "LOT", "CHARG") and idx_lote is None:
                    idx_lote = i
                elif cn in ("UMB", "UNID", "UNIDADE") and idx_umb is None:
                    idx_umb = i

            if idx_mat is None: idx_mat = 1
            if idx_dep is None: idx_dep = 2
            if idx_lote is None: idx_lote = 3

            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row:
                    continue
                dep_val = str(row[idx_dep] or "").strip().upper() if len(row) > idx_dep else ""
                lote_val = str(row[idx_lote] or "").strip().upper() if len(row) > idx_lote else ""

                if dep_alvo in dep_val and lote_alvo in lote_val:
                    m_cod = _normalizar_codigo(row[idx_mat]) if len(row) > idx_mat else ""
                    if m_cod:
                        m_desc = str(row[idx_desc] or "").strip() if idx_desc is not None and len(row) > idx_desc else ""
                        m_umb = str(row[idx_umb] or "").strip().upper() if idx_umb is not None and len(row) > idx_umb else "UN"
                        if m_cod in dados_agrupados:
                            dados_agrupados[m_cod]["qtd"] += 1
                            if not dados_agrupados[m_cod]["desc"] and m_desc:
                                dados_agrupados[m_cod]["desc"] = m_desc
                        else:
                            dados_agrupados[m_cod] = {
                                "desc": m_desc,
                                "umb": m_umb or "UN",
                                "qtd": 1
                            }
        finally:
            try:
                wb.close()
            except Exception:
                pass
            try:
                buf.close()
            except Exception:
                pass

        if not dados_agrupados:
            raise RuntimeError(f"Nenhum serial/item com Depósito '{dep_alvo}' e Lote '{lote_alvo}' localizado na aba ESTOQUE.")

        # Cria arquivo Excel com formato oficial de Contagem do Depósito (Material, Depósito, Estoque, QNT, UMB)
        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "Estoque Livre"

        headers = ["Material", "Depósito", "Estoque", "QNT", "UMB"]
        ws_out.append(headers)

        hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF")
        )

        for col_idx in range(1, 6):
            cell = ws_out.cell(row=1, column=col_idx)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws_out.row_dimensions[1].height = 24

        for row_idx, (cod, item) in enumerate(sorted(dados_agrupados.items()), start=2):
            c_mat = ws_out.cell(row=row_idx, column=1, value=cod)
            c_mat.alignment = Alignment(horizontal="center")
            c_mat.border = thin_border

            c_dep = ws_out.cell(row=row_idx, column=2, value=dep_alvo)
            c_dep.alignment = Alignment(horizontal="center")
            c_dep.border = thin_border

            c_est = ws_out.cell(row=row_idx, column=3, value=item["qtd"])
            c_est.alignment = Alignment(horizontal="center")
            c_est.border = thin_border

            c_qnt = ws_out.cell(row=row_idx, column=4, value="")
            c_qnt.alignment = Alignment(horizontal="center")
            c_qnt.border = thin_border

            c_umb = ws_out.cell(row=row_idx, column=5, value=item["umb"])
            c_umb.alignment = Alignment(horizontal="center")
            c_umb.border = thin_border

        ws_out.column_dimensions["A"].width = 16
        ws_out.column_dimensions["B"].width = 12
        ws_out.column_dimensions["C"].width = 12
        ws_out.column_dimensions["D"].width = 10
        ws_out.column_dimensions["E"].width = 10

        pasta_data = os.path.join(obter_diretorio_executavel(), "data")
        os.makedirs(pasta_data, exist_ok=True)
        arq_saida = os.path.join(pasta_data, f"CONTAGEM_FISICA_LOTE0_{dep_alvo}.xlsx")
        wb_out.save(arq_saida)
        return arq_saida, dados_agrupados

    def _extrair_lote0_planilha_beta(self):
        """
        Dispara a extração do Lote 0 em background a partir da planilha beta de seriais.
        """
        if self.processando:
            messagebox.showwarning("Aviso", "Aguarde a operação em andamento terminar.")
            return

        caminho_beta = self.caminho_beta_padrao
        if not caminho_beta or not os.path.exists(caminho_beta):
            caminho_beta = filedialog.askopenfilename(
                title="Selecione a Planilha de Seriais (CONTROLE_SERIAIS.vBeta.xlsm)",
                filetypes=[("Planilhas Excel", "*.xlsm;*.xlsx;*.xls"), ("Todos os arquivos", "*.*")],
                parent=self.parent_frame.winfo_toplevel()
            )
            if not caminho_beta:
                return
            self.caminho_beta_padrao = caminho_beta

        dep_alvo = self.deposito_atual
        lote_alvo = "LT0"

        self.btn_extrair_lote0.configure(state="disabled")
        self._status(f"Puxando Lote 0 de {os.path.basename(caminho_beta)} para {dep_alvo}...")

        def _worker_extracao():
            try:
                arq_saida, dados_agrupados = self._extrair_lote0_nucleo(caminho_beta, dep_alvo, lote_alvo)

                def _sucesso():
                    self.ent_controle.delete(0, tk.END)
                    self.ent_controle.insert(0, arq_saida)
                    self.caminho_controle = arq_saida
                    msg = f"Lote 0 ({dep_alvo}) extraído: {len(dados_agrupados)} materiais consolidados!"
                    self._status(msg)
                    show_toast(self.parent_frame, msg, type="success")
                    play_sound("done")

                self._agendar_ui(_sucesso)

            except Exception as e:
                log_error(f"Erro ao extrair Lote 0 da planilha beta: {e}")
                self._status_threadsafe(f"Erro: {str(e)}")
                play_sound("error")
                self._agendar_ui(lambda: messagebox.showerror("Erro na Extração do Lote 0", str(e), parent=self.parent_frame.winfo_toplevel()))
            finally:
                self._agendar_ui(lambda: self.btn_extrair_lote0.configure(state="normal"))

        threading.Thread(target=_worker_extracao, daemon=True).start()

    def _carregar_dados_zm29(self, zm_path, dep_alvo):
        """
        Carrega dados da ZM29 exportada do SAP.
        Retorna dicionário: { cod: {"desc": str, "umb": str, "qtd_sap": float} }
        """
        if zm_path.lower().endswith(".csv"):
            try:
                df = pd.read_csv(zm_path, sep=";", encoding="latin1")
            except Exception:
                df = pd.read_csv(zm_path, sep=",", encoding="utf-8")
        else:
            xl = pd.ExcelFile(zm_path, engine="openpyxl")
            sheet_target = xl.sheet_names[0]
            for s in xl.sheet_names:
                if "ZM29" in s.upper() or "SAP" in s.upper():
                    sheet_target = s
                    break
            df = xl.parse(sheet_target)

        # Identifica colunas flexíveis da ZM29
        col_mat = next((c for c in df.columns if re.search(r"MATERIA|CODIGO|MATNR", str(c).upper())), df.columns[0])
        col_desc = next((c for c in df.columns if re.search(r"DESCRI|TEXTO|MAKTX", str(c).upper())), df.columns[1] if len(df.columns) > 1 else None)
        col_dep = next((c for c in df.columns if re.search(r"\bDEP|\bDEPOSITO\b|LGORT", str(c).upper())), df.columns[3] if len(df.columns) > 3 else None)
        col_est = next((c for c in df.columns if re.search(r"ESTOQUE\s*LIVRE|LIVRE|UTILIZ|SALDO|QTD", str(c).upper())), df.columns[4] if len(df.columns) > 4 else None)
        col_umb = next((c for c in df.columns if re.search(r"\bUMB\b|UNID|MEINS", str(c).upper())), df.columns[5] if len(df.columns) > 5 else None)

        dados_zm = {}
        for _, row in df.iterrows():
            if col_dep and pd.notna(row[col_dep]):
                dep_row = str(row[col_dep]).strip().upper()
                if dep_row and dep_alvo not in dep_row:
                    continue
            m_cod = _normalizar_codigo(row[col_mat])
            if not m_cod:
                continue
            desc = str(row[col_desc]).strip() if col_desc and pd.notna(row[col_desc]) else ""
            umb = str(row[col_umb]).strip().upper() if col_umb and pd.notna(row[col_umb]) else "UN"
            qtd_val = _limpar_quantidade(row[col_est]) if col_est and pd.notna(row[col_est]) else 0.0

            if m_cod in dados_zm:
                dados_zm[m_cod]["qtd_sap"] += qtd_val
            else:
                dados_zm[m_cod] = {
                    "desc": desc,
                    "umb": umb,
                    "qtd_sap": qtd_val
                }
        return dados_zm

    def _carregar_dados_contagem_fisica(self, ctrl_path, dep_alvo, lote_alvo, zm_path=""):
        """
        Carrega dados da contagem física do depósito. Suporta:
        1. Planilha formato ZM29 com coluna de 'Estoque Livre' preenchida.
        2. Planilha com aba chamada 'ESTOQUE LIVRE' ou 'CONTAGEM'.
        3. Caso único: mesmo arquivo da ZM29 com aba separada de contagem.
        4. Fallback legado: Planilha de seriais beta (aba ESTOQUE com linhas de Lote/Série).
        Retorna dicionário: { cod: {"desc": str, "umb": str, "qtd_fisico": float} }
        """
        # Se ctrl_path não foi informado, checa se zm_path tem aba de contagem física
        caminho_leitura = ctrl_path
        aba_alvo = None

        if not caminho_leitura or not os.path.exists(caminho_leitura):
            if zm_path and os.path.exists(zm_path) and not zm_path.lower().endswith(".csv"):
                try:
                    xl_zm = pd.ExcelFile(zm_path, engine="openpyxl")
                    for s in xl_zm.sheet_names:
                        s_up = s.upper()
                        if any(k in s_up for k in ("ESTOQUE LIVRE", "LIVRE", "CONTAGEM", "FISICO", "LOTE 0", "LOTE0")):
                            caminho_leitura = zm_path
                            aba_alvo = s
                            break
                except Exception:
                    pass
            if not caminho_leitura or not os.path.exists(caminho_leitura):
                raise RuntimeError("Planilha de Contagem Física não encontrada. Selecione a planilha formato ZM29 ou use o botão 'Puxar Lote 0 da Planilha Beta'.")

        dados_fisico = {}

        if caminho_leitura.lower().endswith(".csv"):
            try:
                df = pd.read_csv(caminho_leitura, sep=";", encoding="latin1")
            except Exception:
                df = pd.read_csv(caminho_leitura, sep=",", encoding="utf-8")
        else:
            xl = pd.ExcelFile(caminho_leitura, engine="openpyxl")
            if aba_alvo and aba_alvo in xl.sheet_names:
                aba_nome = aba_alvo
            else:
                # Procura aba prioritária de contagem
                aba_nome = None
                for s in xl.sheet_names:
                    s_up = s.upper()
                    if any(k in s_up for k in ("ESTOQUE LIVRE", "CONTAGEM", "FISICO", "LOTE 0", "LOTE0")):
                        aba_nome = s
                        break
                if not aba_nome:
                    # Se houver aba 'ESTOQUE' e for a planilha de seriais tradicional
                    if "ESTOQUE" in xl.sheet_names and len(xl.sheet_names) > 1:
                        aba_nome = "ESTOQUE"
                    else:
                        # Se for o mesmo arquivo da ZM29 com 2 abas, usa a 2ª aba
                        if caminho_leitura == zm_path and len(xl.sheet_names) > 1:
                            aba_nome = xl.sheet_names[1]
                        else:
                            aba_nome = xl.sheet_names[0]
            df = xl.parse(aba_nome)

        # Verifica se o DataFrame possui colunas de formato de contagem (Material, Depósito, Estoque, QNT, UMB)
        col_mat = next((c for c in df.columns if re.search(r"MATERIA|CODIGO|MATNR", str(c).upper())), None)
        col_desc = next((c for c in df.columns if re.search(r"DESCRI|TEXTO|MAKTX|DENOMINA", str(c).upper())), None)
        col_dep = next((c for c in df.columns if re.search(r"\bDEP|\bDEPOSITO\b|LGORT", str(c).upper())), None)
        col_est = next((c for c in df.columns if re.search(r"ESTOQUE\s*LIVRE|\bESTOQUE\b|LIVRE|UTILIZ|SALDO|CONTA|FISIC", str(c).upper())), None)
        col_qnt = next((c for c in df.columns if re.search(r"\bQNT\b|\bQTD\b|QUANT", str(c).upper())), None)
        col_umb = next((c for c in df.columns if re.search(r"\bUMB\b|UNID|MEINS", str(c).upper())), None)

        col_qtd_efetiva = col_est or col_qnt

        if col_mat is not None and col_qtd_efetiva is not None:
            # FORMATO DE CONTAGEM FÍSICA (Material, Depósito, Estoque, QNT, UMB ou ZM29)
            for _, row in df.iterrows():
                if col_dep and pd.notna(row[col_dep]):
                    dep_row = str(row[col_dep]).strip().upper()
                    if dep_row and dep_alvo not in dep_row:
                        continue
                m_cod = _normalizar_codigo(row[col_mat])
                if not m_cod:
                    continue

                # Prioriza QNT se o operador preencheu nela; caso contrário usa o Estoque
                qtd_val = 0.0
                if col_qnt and pd.notna(row[col_qnt]) and str(row[col_qnt]).strip() != "":
                    qtd_val = _limpar_quantidade(row[col_qnt])
                elif col_est and pd.notna(row[col_est]) and str(row[col_est]).strip() != "":
                    qtd_val = _limpar_quantidade(row[col_est])

                desc = str(row[col_desc]).strip() if col_desc and pd.notna(row[col_desc]) else ""
                umb = str(row[col_umb]).strip().upper() if col_umb and pd.notna(row[col_umb]) else "UN"

                if m_cod in dados_fisico:
                    dados_fisico[m_cod]["qtd_fisico"] += qtd_val
                else:
                    dados_fisico[m_cod] = {
                        "desc": desc,
                        "umb": umb,
                        "qtd_fisico": qtd_val
                    }
        else:
            # FALLBACK LEGADO: Planilha com linhas de seriais individuais (aba ESTOQUE)
            col_lote = next((c for c in df.columns if re.search(r"LOTE|LOT|CHARG", str(c).upper())), None)
            if col_mat is None:
                col_mat = df.columns[1] if len(df.columns) > 1 else df.columns[0]
            if col_dep is None and len(df.columns) > 2:
                col_dep = df.columns[2]
            if col_lote is None and len(df.columns) > 3:
                col_lote = df.columns[3]

            for _, row in df.iterrows():
                dep_row = str(row[col_dep]).strip().upper() if col_dep and pd.notna(row[col_dep]) else ""
                lote_row = str(row[col_lote]).strip().upper() if col_lote and pd.notna(row[col_lote]) else ""

                if dep_alvo in dep_row and lote_alvo in lote_row:
                    m_cod = _normalizar_codigo(row[col_mat])
                    if not m_cod:
                        continue
                    desc = str(row[col_desc]).strip() if col_desc and pd.notna(row[col_desc]) else ""
                    umb = str(row[col_umb]).strip().upper() if col_umb and pd.notna(row[col_umb]) else "UN"

                    if m_cod in dados_fisico:
                        dados_fisico[m_cod]["qtd_fisico"] += 1.0
                    else:
                        dados_fisico[m_cod] = {
                            "desc": desc,
                            "umb": umb,
                            "qtd_fisico": 1.0
                        }

        return dados_fisico

    def executar_comparacao(self):
        if self.processando:
            messagebox.showwarning("Aviso", "Já existe uma comparação sendo processada.")
            return

        zm_path = self.ent_zm29.get().strip()
        ctrl_path = self.ent_controle.get().strip()
        lote_alvo = "LT0"
        dep_alvo = self.deposito_atual

        if not zm_path or not os.path.exists(zm_path):
            messagebox.showwarning("Aviso", "Por favor, selecione o arquivo da transação ZM29 exportada do SAP.")
            return

        self.processando = True
        self.btn_comparar.configure(state="disabled")
        self._status(f"Cruzando inventário do {dep_alvo} (Saldo SAP vs. Estoque Livre Contado)...")

        def _worker():
            try:
                # 1. Carrega dados do SAP ZM29
                dados_zm = self._carregar_dados_zm29(zm_path, dep_alvo)

                # 2. Carrega dados da Contagem Física
                dados_fisico = self._carregar_dados_contagem_fisica(ctrl_path, dep_alvo, lote_alvo, zm_path)

                # 3. Cruzamento e Consolidação dos Resultados
                todos_materiais = sorted(set(dados_zm.keys()) | set(dados_fisico.keys()))
                resultado_itens = []

                total_ok = 0
                total_falta = 0
                total_sobra = 0

                for m in todos_materiais:
                    info_zm = dados_zm.get(m, {})
                    info_fis = dados_fisico.get(m, {})

                    qtd_sap = float(info_zm.get("qtd_sap", 0.0))
                    qtd_fisico = float(info_fis.get("qtd_fisico", 0.0))
                    desc = info_zm.get("desc") or info_fis.get("desc") or "(Material sem descrição)"
                    umb = info_zm.get("umb") or info_fis.get("umb") or "UN"
                    dif = qtd_fisico - qtd_sap

                    if dif == 0:
                        status_str = "🟢 Bateu"
                        status_tipo = "OK"
                        total_ok += 1
                    elif dif < 0:
                        status_str = "🔴 Faltou"
                        status_tipo = "FALTA"
                        total_falta += 1
                    else:
                        status_str = "🟡 Sobrando"
                        status_tipo = "SOBRA"
                        total_sobra += 1

                    resultado_itens.append({
                        "material": m,
                        "descricao": desc,
                        "deposito": dep_alvo,
                        "centro": "EMPR",
                        "umb": umb,
                        "qtd_sap": int(qtd_sap) if qtd_sap.is_integer() else qtd_sap,
                        "qtd_fisico": int(qtd_fisico) if qtd_fisico.is_integer() else qtd_fisico,
                        "diferenca": int(dif) if dif.is_integer() else dif,
                        "status": status_str,
                        "tipo": status_tipo
                    })

                self.itens_comparados = resultado_itens
                total_geral = len(resultado_itens)
                acuracia = (total_ok / total_geral * 100.0) if total_geral > 0 else 0.0

                def _atualizar_ui():
                    self.card_total.lbl_val.configure(text=f"{total_geral}")
                    self.card_ok.lbl_val.configure(text=f"{total_ok}")
                    self.card_faltou.lbl_val.configure(text=f"{total_falta}")
                    self.card_sobrou.lbl_val.configure(text=f"{total_sobra}")
                    self.card_acuracia.lbl_val.configure(text=f"{acuracia:.1f}%")

                    self.btn_filtro_todos.configure(text=f"Todos ({total_geral})")
                    self.btn_filtro_div.configure(text=f"🚨 Divergências ({total_falta + total_sobra})")
                    self.btn_filtro_falta.configure(text=f"🔴 Faltas ({total_falta})")
                    self.btn_filtro_sobra.configure(text=f"🟡 Sobras ({total_sobra})")
                    self.btn_filtro_ok.configure(text=f"🟢 OK ({total_ok})")

                    self._aplicar_filtro_busca()
                    self._status(f"Concluído! {dep_alvo}: {total_ok} certos, {total_falta} faltas, {total_sobra} sobras.")
                    show_toast(self.parent_frame, f"Auditoria concluída para o {dep_alvo}!", type="success")
                    play_sound("done")

                self.parent_frame.after(0, _atualizar_ui)

            except Exception as e:
                log_error(f"Erro na comparação de depósito: {e}")
                self._status_threadsafe(f"Erro: {str(e)}")
                play_sound("error")
                self.parent_frame.after(0, lambda: messagebox.showerror("Erro na Auditoria", str(e), parent=self.parent_frame.winfo_toplevel()))
            finally:
                def _fim():
                    self.processando = False
                    self.btn_comparar.configure(state="normal")
                self.parent_frame.after(0, _fim)

        threading.Thread(target=_worker, daemon=True).start()

    def _trocar_filtro(self, tipo):
        self.filtro_ativo = tipo
        self._aplicar_filtro_busca()

    def _aplicar_filtro_busca(self):
        termo = self.ent_busca.get().strip().upper()
        filtrados = []

        for item in self.itens_comparados:
            # Filtro por tipo
            if self.filtro_ativo == "DIVERGENTES" and item["tipo"] == "OK":
                continue
            elif self.filtro_ativo == "FALTAS" and item["tipo"] != "FALTA":
                continue
            elif self.filtro_ativo == "SOBRAS" and item["tipo"] != "SOBRA":
                continue
            elif self.filtro_ativo == "OK" and item["tipo"] != "OK":
                continue

            # Filtro por texto
            if termo:
                if termo not in item["material"] and termo not in item["descricao"].upper():
                    continue

            filtrados.append(item)

        self.itens_exibidos = filtrados
        self._preencher_treeview(filtrados)

    def _preencher_treeview(self, itens):
        self.tree.delete(*self.tree.get_children())
        for it in itens:
            tag = "tag_ok" if it["tipo"] == "OK" else ("tag_falta" if it["tipo"] == "FALTA" else "tag_sobra")
            dif_str = f"+{it['diferenca']}" if it["diferenca"] > 0 else str(it["diferenca"])

            self.tree.insert(
                "", tk.END,
                values=(
                    it["material"], it["descricao"], it["deposito"], it.get("centro", "EMPR"),
                    it["umb"], it["qtd_sap"], it["qtd_fisico"], dif_str, it["status"]
                ),
                tags=(tag,)
            )

    def _ordenar_coluna(self, col, reverse):
        try:
            lista = [(self.tree.set(k, col), k) for k in self.tree.get_children("")]
            # Converte para número se possível para ordenação natural
            try:
                lista.sort(key=lambda t: float(str(t[0]).replace("+", "")), reverse=reverse)
            except ValueError:
                lista.sort(reverse=reverse)

            for index, (val, k) in enumerate(lista):
                self.tree.move(k, "", index)

            self.tree.heading(col, command=lambda: self._ordenar_coluna(col, not reverse))
        except Exception:
            pass

    def exportar_excel(self):
        if not self.itens_comparados:
            messagebox.showwarning("Aviso", "Não há dados comparados para exportar.")
            return

        caminho_salvar = filedialog.asksaveasfilename(
            title="Salvar Relatório de Contagem do Depósito",
            defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")],
            initialfile=f"CONTAGEM_DEPOSITO_{self.deposito_atual}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            parent=self.parent_frame.winfo_toplevel()
        )
        if not caminho_salvar:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Inventário {self.deposito_atual}"

            # Título
            ws.merge_cells("A1:I1")
            ws["A1"] = f"RELATÓRIO DE INVENTÁRIO ROTATIVO — DEPÓSITO {self.deposito_atual} (ESTOQUE LIVRE ⟷ ZM29)"
            ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 36

            # Subtítulo com métricas
            ws.merge_cells("A2:I2")
            total_ok = sum(1 for i in self.itens_comparados if i["tipo"] == "OK")
            total_falta = sum(1 for i in self.itens_comparados if i["tipo"] == "FALTA")
            total_sobra = sum(1 for i in self.itens_comparados if i["tipo"] == "SOBRA")
            ws["A2"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Total: {len(self.itens_comparados)} | 🟢 Bateu: {total_ok} | 🔴 Faltou: {total_falta} | 🟡 Sobrando: {total_sobra}"
            ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
            ws["A2"].fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 22

            # Cabeçalhos
            headers = ["Material", "Descrição", "Depósito", "Centro", "UMB", "Qtd SAP (ZM29)", "Qtd Físico (Estoque Livre)", "Diferença", "Status"]
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=3, column=col_idx, value=h)
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[3].height = 26

            fill_falta = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            fill_sobra = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            fill_ok = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
            font_falta = Font(name="Segoe UI", size=10, color="991B1B", bold=True)
            font_sobra = Font(name="Segoe UI", size=10, color="92400E", bold=True)
            font_ok = Font(name="Segoe UI", size=10, color="065F46")

            # Linhas de dados
            for row_idx, it in enumerate(self.itens_comparados, 4):
                ws.cell(row=row_idx, column=1, value=it["material"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=2, value=it["descricao"]).alignment = Alignment(horizontal="left")
                ws.cell(row=row_idx, column=3, value=it["deposito"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=4, value=it.get("centro", "EMPR")).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=5, value=it["umb"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=6, value=it["qtd_sap"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=7, value=it["qtd_fisico"]).alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=8, value=it["diferenca"]).alignment = Alignment(horizontal="center")
                cell_st = ws.cell(row=row_idx, column=9, value=it["status"])
                cell_st.alignment = Alignment(horizontal="center")

                # Destaque de cor
                f_cor = fill_ok if it["tipo"] == "OK" else (fill_falta if it["tipo"] == "FALTA" else fill_sobra)
                font_cor = font_ok if it["tipo"] == "OK" else (font_falta if it["tipo"] == "FALTA" else font_sobra)
                cell_st.fill = f_cor
                cell_st.font = font_cor

            # Ajuste de largura das colunas (ignora linhas 1 e 2 que são títulos mesclados para não alargar a Coluna A)
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col if cell.row >= 3)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 60)

            wb.save(caminho_salvar)
            show_toast(self.parent_frame, "Relatório exportado com sucesso!", type="success")
            play_sound("done")
            msg_ok = f"Relatório salvo com sucesso em:\n{caminho_salvar}"
            messagebox.showinfo("Sucesso", msg_ok)

        except Exception as e:
            log_error(f"Erro ao exportar Excel: {e}")
            messagebox.showerror("Erro ao Exportar", str(e))

    def copiar_resumo(self):
        if not self.itens_comparados:
            messagebox.showwarning("Aviso", "Não há dados para copiar.")
            return

        total = len(self.itens_comparados)
        ok = sum(1 for i in self.itens_comparados if i["tipo"] == "OK")
        faltas = [i for i in self.itens_comparados if i["tipo"] == "FALTA"]
        sobras = [i for i in self.itens_comparados if i["tipo"] == "SOBRA"]

        txt = [
            f"📊 *RESUMO DO INVENTÁRIO ROTATIVO — DEPÓSITO {self.deposito_atual}*",
            f"📅 Data: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            f"📦 Total de Itens: {total}",
            f"🟢 Bateu 100%: {ok}",
            f"🔴 Faltou: {len(faltas)} itens",
            f"🟡 Sobrando: {len(sobras)} itens",
            ""
        ]

        if faltas:
            txt.append("🚨 *PRINCIPAIS FALTAS (Déficit):*")
            for f in faltas[:8]:
                txt.append(f"  • {f['material']} ({f['descricao'][:25]}): {f['diferenca']} {f['umb']}")
            if len(faltas) > 8:
                txt.append(f"  ... e mais {len(faltas) - 8} itens com falta.")
            txt.append("")

        if sobras:
            txt.append("⚠️ *PRINCIPAIS SOBRAS (Excesso):*")
            for s in sobras[:8]:
                txt.append(f"  • {s['material']} ({s['descricao'][:25]}): +{s['diferenca']} {s['umb']}")
            if len(sobras) > 8:
                txt.append(f"  ... e mais {len(sobras) - 8} itens sobrando.")

        resumo_str = "\n".join(txt)
        self.parent_frame.clipboard_clear()
        self.parent_frame.clipboard_append(resumo_str)
        show_toast(self.parent_frame, "Resumo copiado para a área de transferência!", type="success")
        self._status("Resumo copiado para WhatsApp/E-mail.")

    def limpar_tela(self):
        self.itens_comparados = []
        self.itens_exibidos = []
        self.tree.delete(*self.tree.get_children())
        self.ent_busca.delete(0, tk.END)

        self.card_total.lbl_val.configure(text="0")
        self.card_ok.lbl_val.configure(text="0")
        self.card_faltou.lbl_val.configure(text="0")
        self.card_sobrou.lbl_val.configure(text="0")
        self.card_acuracia.lbl_val.configure(text="0.0%")

        self.btn_filtro_todos.configure(text="Todos (0)")
        self.btn_filtro_div.configure(text="🚨 Divergências (0)")
        self.btn_filtro_falta.configure(text="🔴 Faltas (0)")
        self.btn_filtro_sobra.configure(text="🟡 Sobras (0)")
        self.btn_filtro_ok.configure(text="🟢 OK (0)")

        self._status("Dados limpos.")

    def aplicar_tema(self, tema_nome):
        c = self._cores(tema_nome)
        self.parent_frame.configure(fg_color=c["bg"])
        self.container.configure(fg_color=c["container"])
        self.card_top.configure(fg_color=c.get("card", "#131d2e"))
        self.f_tabela_container.configure(fg_color=c.get("card", "#131d2e"))

        # Inputs
        self.ent_centro.configure(fg_color=c["input"], text_color=c["texto"], border_color=c.get("border", "#1e293b"))
        self.ent_zm29.configure(fg_color=c["input"], text_color=c["texto"], border_color=c.get("border", "#1e293b"))
        self.ent_controle.configure(fg_color=c["input"], text_color=c["texto"], border_color=c.get("border", "#1e293b"))
        self.ent_busca.configure(fg_color=c["input"], text_color=c["texto"], border_color=c.get("border", "#1e293b"))

        # Botões principais
        self.btn_comparar.configure(fg_color=c["accent"], hover_color=c["hover"], text_color="#090d16" if "Light" not in tema_nome else "#ffffff")
        self.btn_extrair_lote0.configure(fg_color=c["input"], hover_color=c["hover"], text_color=c["texto"])
        self.btn_sel_zm29.configure(fg_color=c["input"], hover_color=c["hover"], text_color=c["texto"])
        self.btn_sel_controle.configure(fg_color=c["input"], hover_color=c["hover"], text_color=c["texto"])
        self.btn_exportar.configure(fg_color=c["accent"], hover_color=c["hover"], text_color="#090d16" if "Light" not in tema_nome else "#ffffff")
        self.btn_copiar_resumo.configure(fg_color=c["input"], hover_color=c["hover"], text_color=c["texto"])
        self.btn_limpar.configure(fg_color=c["input"], hover_color=c["hover"], text_color=c["texto"])

        # Filtros
        for b in (self.btn_filtro_todos, self.btn_filtro_div, self.btn_filtro_falta, self.btn_filtro_sobra, self.btn_filtro_ok):
            b.configure(fg_color=c["input"], hover_color=c["hover"], text_color=c["texto"])

        # Cards KPIs
        for card in (self.card_total, self.card_ok, self.card_faltou, self.card_sobrou, self.card_acuracia):
            card.configure(fg_color=c.get("card", "#131d2e"))

        self._atualizar_botoes_deposito()

        # Estilo Treeview
        style = ttk.Style()
        tv_bg = c.get("card", "#131d2e")
        tv_fg = c["texto"]
        hdr_bg = c["container"]
        hdr_fg = c["texto"]

        style.theme_use("default")
        style.configure(
            "Deposito.Treeview",
            background=tv_bg,
            foreground=tv_fg,
            fieldbackground=tv_bg,
            borderwidth=0,
            font=("Segoe UI", 10),
            rowheight=26
        )
        style.configure(
            "Deposito.Treeview.Heading",
            background=hdr_bg,
            foreground=hdr_fg,
            font=("Segoe UI", 10, "bold"),
            borderwidth=1,
            relief="flat"
        )
        style.map("Deposito.Treeview", background=[("selected", c["accent"])], foreground=[("selected", "white")])
        style.map("Deposito.Treeview.Heading", background=[("active", c["hover"])])


def construir_aba_contagem_deposito(parent_frame, label_status_global=None, paletas=None, obter_tema_atual=None):
    """Ponto de entrada padrão para o Conecta Hub / N.E.X.U.S."""
    return AbaContagemDeposito(parent_frame, label_status_global, paletas, obter_tema_atual)

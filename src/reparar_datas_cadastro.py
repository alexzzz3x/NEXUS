# -*- coding: utf-8 -*-
"""
Script de Restauração e Padronização Segura de Ordens e Vencimentos na Planilha CONTROLE_SERIAIS.vBeta.xlsm (Aba CADASTRO).
Restaura os dados originais íntegros a partir do backup de referência com formato dd/mm/yyyy para todas as 261 equipes.
"""

import os
import sys
import shutil
import json
from datetime import datetime, date
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

CAMINHO_CONFIG = os.path.expandvars("%APPDATA%/ConectaHub/config_saida_osme_caminho.json")
BACKUP_REF_NOME = "BACKUP_PRE_ATUALIZAR_OSME_20260921_101924_CONTROLE_SERIAIS.vBeta.xlsm"


def obter_caminho_planilha():
    if os.path.exists(CAMINHO_CONFIG):
        with open(CAMINHO_CONFIG, "r", encoding="utf-8") as f:
            cfg = json.load(f)
            c = cfg.get("caminho_planilha")
            if c and os.path.exists(c):
                return c
    c_padrao = "X:/obras/Elpa - Almox CR Mauá/CONTROLES/CONTROLE_SERIAIS.vBeta.xlsm"
    if os.path.exists(c_padrao):
        return c_padrao
    raise FileNotFoundError("Não foi possível localizar o arquivo CONTROLE_SERIAIS.vBeta.xlsm.")


def executar_reparo():
    caminho = obter_caminho_planilha()
    pasta_dir = os.path.dirname(caminho)
    pasta_backup = os.path.join(pasta_dir, "Backup_NEXUS")
    os.makedirs(pasta_backup, exist_ok=True)

    caminho_bak_ref = os.path.join(pasta_backup, BACKUP_REF_NOME)
    if not os.path.exists(caminho_bak_ref):
        raise FileNotFoundError(f"Backup de referência não encontrado: {caminho_bak_ref}")

    print(f"Planilha alvo: {caminho}")
    print(f"Backup de referência: {caminho_bak_ref}")

    # 1. Backup Preventivo Pré-Correção
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho_bak_seguranca = os.path.join(pasta_backup, f"BACKUP_PRE_CORRECAO_DATAS_{timestamp}_CONTROLE_SERIAIS.vBeta.xlsm")
    shutil.copy2(caminho, caminho_bak_seguranca)
    print(f"Backup preventivo criado: {os.path.basename(caminho_bak_seguranca)}")

    # 2. Carrega Backup de Referência (dados íntegros)
    wb_ref = openpyxl.load_workbook(caminho_bak_ref, read_only=True, data_only=False)
    ws_ref = wb_ref["CADASTRO"]

    bak_dict = {}
    for r in range(2, 263):
        eq = ws_ref.cell(r, 5).value
        ord_val = ws_ref.cell(r, 8).value
        venc_val = ws_ref.cell(r, 9).value
        if eq:
            eq_norm = str(eq).strip().upper()
            bak_dict[eq_norm] = (ord_val, venc_val)
    wb_ref.close()

    print(f"Total de equipes carregadas do backup de referência: {len(bak_dict)}")

    # 3. Carrega Planilha Oficial para Restauração
    wb = openpyxl.load_workbook(caminho, keep_vba=True)
    ws = wb["CADASTRO"]

    atualizadas = 0
    exemplos_corrigidos = []

    print("\n--- INICIANDO RESTAURAÇÃO ---")
    for r in range(2, 263):
        eq = ws.cell(r, 5).value
        if not eq:
            continue
        eq_norm = str(eq).strip().upper()

        if eq_norm in bak_dict:
            ord_ref, venc_ref = bak_dict[eq_norm]

            # Atualiza Coluna H (ORDEM) se vazia ou se diferente
            if ord_ref is not None:
                ws.cell(r, 8).value = ord_ref

            # Atualiza Coluna I (VENCIMENTO) com datetime seguro e formato dd/mm/yyyy
            dt_correto = None
            if isinstance(venc_ref, datetime):
                dt_correto = datetime(venc_ref.year, venc_ref.month, venc_ref.day)
            elif isinstance(venc_ref, str) and "/" in venc_ref:
                partes = venc_ref.strip().split("/")
                if len(partes) == 3:
                    try:
                        d, m, y = int(partes[0]), int(partes[1]), int(partes[2])
                        dt_correto = datetime(y, m, d)
                    except Exception:
                        pass

            if dt_correto is not None:
                cell_venc = ws.cell(r, 9)
                cell_venc.value = dt_correto
                cell_venc.number_format = "dd/mm/yyyy"

            atualizadas += 1
            if len(exemplos_corrigidos) < 10 or eq_norm in ("CSC308", "CSC300", "CSC304", "CDC317", "CDC400"):
                v_str = dt_correto.strftime("%d/%m/%Y") if dt_correto else str(venc_ref)
                exemplos_corrigidos.append(f"Linha {r} ({eq_norm}): Ordem={ord_ref} | Vencimento={v_str}")

    wb.save(caminho)
    wb.close()

    print(f"\nSUCESSO:")
    print(f"Total de equipes restauradas com Ordem e Vencimento: {atualizadas}")
    print("\nAmostra de equipes restauradas:")
    for ex in exemplos_corrigidos:
        print(f"  {ex}")


if __name__ == "__main__":
    executar_reparo()

import os
import re
import json
import time
import subprocess
import tempfile
import threading
from collections import OrderedDict
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import customtkinter as ctk
from openpyxl import load_workbook
import win32com.client

import unicodedata
from utils_os import obter_caminho_config, obter_diretorio_dados, obter_diretorio_executavel
from modulo_extrator import extrair_lista_seriais
try:
    from utils_workforce import _carregar_banco_nomes, executar_fluxo_workforce, carregar_credenciais_workforce
except ImportError:
    from modulo_fichas import _carregar_banco_nomes, executar_fluxo_workforce, carregar_credenciais_workforce
from utils_pendencias_planilha import (
    adicionar_pendencia,
    obter_quantidade_pendencias,
    processar_todas_pendencias,
    ModalVisualizarPendencias,
    is_file_writable,
    criar_backup_preventivo
)

# Compartilha a mesma configuração de caminho de planilha de controle usada na Saída OSME
CAMINHO_CONFIG = obter_caminho_config("config_saida_osme_caminho.json")

DEFAULT_PALETAS = {
    "Nexus": {
        "bg": "#090d16", "container": "#111827", "texto": "#f8fafc",
        "accent": "#00d2ff", "hover": "#0284c7", "input": "#1a2436",
        "sucesso": "#10b981", "alerta": "#ef4444"
    },
    "Azul Delray": {
        "bg": "#0a192f", "container": "#172a45", "texto": "#ffffff",
        "accent": "#38bdf8", "hover": "#0284c7", "input": "#1f3a60",
        "sucesso": "#10b981", "alerta": "#ef4444"
    },
    "Conecta Blue": {
        "bg": "#0a192f", "container": "#172a45", "texto": "#ffffff",
        "accent": "#38bdf8", "hover": "#0284c7", "input": "#1f3a60",
        "sucesso": "#10b981", "alerta": "#ef4444"
    }
}


def normalizar_texto(valor):
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto_sem_acento = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", texto_sem_acento.strip().upper())


def normalizar_serial_chave(valor):
    """
    Normaliza o número de série para chave única de comparação e deduplicação.
    Remove zeros à esquerda caso seja numérico (ex: '00000000018349138' -> '18349138').
    Mantém intactos seriais alfanuméricos (ex: 'UQM9HDN26L01533345').
    """
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto_sem_acento = "".join(c for c in texto if not unicodedata.combining(c))
    limpo = re.sub(r"[^A-Z0-9]", "", texto_sem_acento.strip().upper())
    if limpo.isdigit():
        return limpo.lstrip("0") or "0"
    return limpo


def formatar_data_excel(valor):
    if valor is None:
        return "-"
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    try:
        num = int(valor)
        if 40000 < num < 60000:
            return (datetime(1899, 12, 30) + timedelta(days=num)).strftime("%d/%m/%Y")
    except Exception:
        pass
    return str(valor).strip()


def copiar_texto_clipboard_resiliente(texto):
    """Copia texto para o Clipboard do Windows de forma resiliente."""
    if not texto:
        return False
    try:
        import win32clipboard
        win32clipboard.OpenClipboard(0)
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardText(str(texto), win32clipboard.CF_UNICODETEXT)
        win32clipboard.CloseClipboard()
        return True
    except Exception:
        pass
    try:
        import subprocess
        subprocess.run(["clip.exe"], input=str(texto), text=True, check=True, timeout=2)
        return True
    except Exception:
        pass
    return False


class ModalAlertaNaoFornecidos(ctk.CTkToplevel):
    """Pop-up modal alertando que determinados seriais NÃO constam como fornecidos."""
    def __init__(self, parent, itens_nao_fornecidos, paletas=None, obter_tema_atual=None, callback_prosseguir=None, total_liberados=0):
        super().__init__(parent)
        self.itens = itens_nao_fornecidos
        self.callback_prosseguir = callback_prosseguir
        self.total_liberados = total_liberados
        self.title("⛔ ALERTA: Seriais Não Encontrados em FORNECIDOS!")
        self.geometry("580x440")
        self.minsize(500, 360)
        self.attributes("-topmost", True)
        self.after(300, lambda: self.attributes("-topmost", False))
        self.configure(fg_color="#18181b")

        # Topo
        frame_topo = ctk.CTkFrame(self, fg_color="#7f1d1d", corner_radius=0, height=55)
        frame_topo.pack(side="top", fill="x")
        frame_topo.pack_propagate(False)

        lbl_tit = ctk.CTkLabel(
            frame_topo,
            text="⛔ SERIAIS NÃO FORNECIDOS DETECTADOS!",
            font=("Segoe UI", 13, "bold"),
            text_color="#fca5a5"
        )
        lbl_tit.pack(side="left", padx=16, pady=14)

        lbl_desc = ctk.CTkLabel(
            self,
            text=(
                f"Foram identificados {len(self.itens)} número(s) de série que NÃO constam na aba FORNECIDOS.\n"
                "Para transferir entre equipes, o medidor precisa estar sob custódia de uma equipe.\n"
                "Estes itens foram bloqueados e NÃO serão enviados ao SAP:"
            ),
            font=("Segoe UI", 11),
            text_color="#e2e8f0",
            justify="left",
            wraplength=540
        )
        lbl_desc.pack(padx=16, pady=(12, 8), anchor="w")

        # Lista com scroll
        scroll_itens = ctk.CTkScrollableFrame(self, fg_color="#27272a", corner_radius=8)
        scroll_itens.pack(fill="both", expand=True, padx=16, pady=(0, 12))

        for it in self.itens:
            card = ctk.CTkFrame(scroll_itens, fg_color="#3f3f46", corner_radius=6)
            card.pack(fill="x", pady=3)

            f_lin1 = ctk.CTkFrame(card, fg_color="transparent")
            f_lin1.pack(fill="x", padx=10, pady=(6, 2))

            serial_txt = it.get("serial", "-")
            ctk.CTkLabel(
                f_lin1,
                text=f"Serial: {serial_txt}",
                font=("Consolas", 12, "bold"),
                text_color="#f87171"
            ).pack(side="left")

            ctk.CTkLabel(
                f_lin1,
                text="Status: Não localizado em FORNECIDOS",
                font=("Segoe UI", 10),
                text_color="#fecaca"
            ).pack(side="right")

        f_botoes = ctk.CTkFrame(self, fg_color="transparent")
        f_botoes.pack(fill="x", padx=16, pady=(0, 14))

        if self.callback_prosseguir and self.total_liberados > 0:
            btn_prosseguir = ctk.CTkButton(
                f_botoes,
                text=f"🚀 Prosseguir Transferência ({self.total_liberados} Liberados)",
                font=("Segoe UI", 11, "bold"),
                fg_color="#10b981",
                hover_color="#059669",
                height=38,
                command=self._prosseguir_acao
            )
            btn_prosseguir.pack(side="left", fill="x", expand=True, padx=(0, 8))

        btn_fechar = ctk.CTkButton(
            f_botoes,
            text="Ajustar Lista / Fechar" if (self.callback_prosseguir and self.total_liberados > 0) else "Entendido (Remover ou Corrigir Seriais)",
            font=("Segoe UI", 11, "bold" if not (self.callback_prosseguir and self.total_liberados > 0) else "normal"),
            fg_color="#334155" if (self.callback_prosseguir and self.total_liberados > 0) else "#dc2626",
            hover_color="#475569" if (self.callback_prosseguir and self.total_liberados > 0) else "#b91c1c",
            height=38,
            command=self.destroy
        )
        btn_fechar.pack(side="right" if (self.callback_prosseguir and self.total_liberados > 0) else "left", fill="x" if not (self.callback_prosseguir and self.total_liberados > 0) else False)

    def _prosseguir_acao(self):
        cb = self.callback_prosseguir
        self.destroy()
        if cb:
            cb()



class ModalVisualizarWorkforce(ctk.CTkToplevel):
    """Pop-up modal exibindo os seriais agrupados por material no formato WorkForce com ';'."""
    def __init__(self, parent, texto_workforce, totais_materiais, paletas=None, obter_tema_atual=None):
        super().__init__(parent)
        self.texto_workforce = texto_workforce
        self.totais_materiais = totais_materiais
        self.title("📋 Lista de Seriais para WorkForce")
        self.geometry("500x580")
        self.minsize(440, 480)
        self.attributes("-topmost", True)
        self.after(300, lambda: self.attributes("-topmost", False))
        self.configure(fg_color="#18181b")

        frame_topo = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=0, height=55)
        frame_topo.pack(side="top", fill="x")
        frame_topo.pack_propagate(False)

        lbl_tit = ctk.CTkLabel(
            frame_topo,
            text="📋 LISTA FORMATADA — WORKFORCE",
            font=("Segoe UI", 13, "bold"),
            text_color="#38bdf8"
        )
        lbl_tit.pack(side="left", padx=16, pady=14)

        total_seriais = sum(len(itens) for itens in totais_materiais.values()) if totais_materiais else 0
        lbl_sub = ctk.CTkLabel(
            self,
            text=f"Total: {total_seriais} seriais separados por material (1 por linha terminando com ';'):",
            font=("Segoe UI", 11),
            text_color="#cbd5e1"
        )
        lbl_sub.pack(anchor="w", padx=16, pady=(10, 6))

        self.txt = ctk.CTkTextbox(self, font=("Consolas", 12), fg_color="#27272a", corner_radius=8)
        self.txt.pack(fill="both", expand=True, padx=16, pady=(0, 10))
        self.txt.insert("1.0", self.texto_workforce)

        f_botoes = ctk.CTkFrame(self, fg_color="transparent")
        f_botoes.pack(fill="x", padx=16, pady=(0, 14))

        self.btn_copiar = ctk.CTkButton(
            f_botoes,
            text="📋 Copiar Tudo para a Área de Transferência",
            font=("Segoe UI", 12, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            height=38,
            command=self._copiar
        )
        self.btn_copiar.pack(side="left", fill="x", expand=True, padx=(0, 8))

        btn_fechar = ctk.CTkButton(
            f_botoes,
            text="Fechar",
            width=80,
            height=38,
            fg_color="transparent",
            border_width=1,
            command=self.destroy
        )
        btn_fechar.pack(side="left")

    def _copiar(self):
        self.clipboard_clear()
        self.clipboard_append(self.texto_workforce)
        self.btn_copiar.configure(text="✓ Copiado com Sucesso!", fg_color="#10b981")
        self.after(2000, lambda: self.btn_copiar.configure(text="📋 Copiar Tudo para a Área de Transferência", fg_color="#0284c7"))


class ModalSucessoTransferenciaEquipes(ctk.CTkToplevel):
    """Pop-up modal exibindo os dois Documentos SAP (Z62 e Z61) com botão de cópia rápida e atalho WorkForce."""
    def __init__(self, parent, doc_z62, doc_z61, origem_resumo, destino, osme_destino, todos_itens, sucesso_planilha, modo_planilha, paletas=None, obter_tema_atual=None, callback_abrir_workforce=None):
        super().__init__(parent)
        self.doc_z62 = str(doc_z62).strip()
        self.doc_z61 = str(doc_z61).strip()
        self.origem_resumo = str(origem_resumo).strip()
        self.destino = str(destino).strip()
        self.osme_destino = str(osme_destino).strip()
        self.todos_itens = todos_itens or []
        self.sucesso_planilha = sucesso_planilha
        self.modo_planilha = modo_planilha
        self.callback_abrir_workforce = callback_abrir_workforce

        self.title("🎉 Transferência entre Equipes Concluída com Sucesso!")
        self.geometry("540x510")
        self.minsize(480, 460)
        self.attributes("-topmost", True)
        self.after(300, lambda: self.attributes("-topmost", False))
        self.configure(fg_color="#18181b")

        # Auto-cópia para o clipboard do documento Z61 (nova saída)
        try:
            self.clipboard_clear()
            self.clipboard_append(self.doc_z61)
        except Exception:
            pass

        # Topo
        frame_topo = ctk.CTkFrame(self, fg_color="#064e3b", corner_radius=0, height=55)
        frame_topo.pack(side="top", fill="x")
        frame_topo.pack_propagate(False)

        lbl_tit = ctk.CTkLabel(
            frame_topo,
            text="🎉 TRANSFERÊNCIA ENTRE EQUIPES CONCLUÍDA!",
            font=("Segoe UI", 13, "bold"),
            text_color="#a7f3d0"
        )
        lbl_tit.pack(side="left", padx=16, pady=14)

        # Container dos Documentos SAP
        f_docs = ctk.CTkFrame(self, fg_color="transparent")
        f_docs.pack(fill="x", padx=16, pady=(12, 6))

        # Card Doc Z62
        card_z62 = ctk.CTkFrame(f_docs, fg_color="#1e293b", corner_radius=8, border_color="#38bdf8", border_width=1)
        card_z62.pack(side="left", fill="both", expand=True, padx=(0, 6))

        ctk.CTkLabel(card_z62, text="DOC. SAP Z62 (ESTORNO ORIGEM)", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").pack(pady=(8, 1))
        ctk.CTkLabel(card_z62, text=self.doc_z62, font=("Consolas", 18, "bold"), text_color="#38bdf8").pack(pady=(0, 6))
        ctk.CTkButton(
            card_z62, text="📋 Copiar Z62", font=("Segoe UI", 10, "bold"), height=26,
            fg_color="#0284c7", hover_color="#0369a1", command=lambda: self._copiar_texto(self.doc_z62)
        ).pack(padx=10, pady=(0, 8), fill="x")

        # Card Doc Z61 (Principal - Nova Saída)
        card_z61 = ctk.CTkFrame(f_docs, fg_color="#1e293b", corner_radius=8, border_color="#10b981", border_width=2)
        card_z61.pack(side="left", fill="both", expand=True, padx=(6, 0))

        ctk.CTkLabel(card_z61, text="DOC. SAP Z61 (SAÍDA DESTINO)", font=("Segoe UI", 9, "bold"), text_color="#a7f3d0").pack(pady=(8, 1))
        ctk.CTkLabel(card_z61, text=self.doc_z61, font=("Consolas", 18, "bold"), text_color="#34d399").pack(pady=(0, 6))
        self.btn_copiar_z61 = ctk.CTkButton(
            card_z61, text="📋 Copiar Z61 (Novo Doc)", font=("Segoe UI", 10, "bold"), height=26,
            fg_color="#10b981", hover_color="#059669", command=lambda: self._copiar_texto(self.doc_z61, self.btn_copiar_z61)
        )
        self.btn_copiar_z61.pack(padx=10, pady=(0, 8), fill="x")

        # Card de Detalhes
        card_detalhes = ctk.CTkFrame(self, fg_color="#27272a", corner_radius=8)
        card_detalhes.pack(fill="both", expand=True, padx=16, pady=(6, 12))

        f_l1 = ctk.CTkFrame(card_detalhes, fg_color="transparent")
        f_l1.pack(fill="x", padx=12, pady=(10, 2))
        ctk.CTkLabel(f_l1, text=f"• Origem(ns): {self.origem_resumo}", font=("Segoe UI", 11), text_color="#f1f5f9").pack(side="left")

        f_l2 = ctk.CTkFrame(card_detalhes, fg_color="transparent")
        f_l2.pack(fill="x", padx=12, pady=(2, 2))
        ctk.CTkLabel(f_l2, text=f"• Destino: {self.destino} (OSME {self.osme_destino})", font=("Segoe UI", 11, "bold"), text_color="#34d399").pack(side="left")

        f_l3 = ctk.CTkFrame(card_detalhes, fg_color="transparent")
        f_l3.pack(fill="x", padx=12, pady=(2, 6))
        ctk.CTkLabel(f_l3, text=f"• Total de Seriais Transferidos: {len(self.todos_itens)} item(ns)", font=("Segoe UI", 11), text_color="#cbd5e1").pack(side="left")

        # Status da Planilha
        f_plan = ctk.CTkFrame(card_detalhes, fg_color="#18181b", corner_radius=6)
        f_plan.pack(fill="x", padx=12, pady=(4, 8))

        if self.sucesso_planilha:
            txt_plan = f"Planilha atualizada com sucesso ({self.modo_planilha})!\nOs registros antigos foram removidos e os novos registros foram gravados em FORNECIDOS."
            cor_plan = "#34d399"
        elif "Fila" in str(self.modo_planilha):
            txt_plan = "Planilha em uso por outro usuário no Excel da rede.\nSalvo na Fila de Sincronização: o N.E.X.U.S gravará automaticamente assim que for fechada!"
            cor_plan = "#f59e0b"
        else:
            txt_plan = "Planilha em uso na rede.\nSalvo na contingência local para sincronização posterior."
            cor_plan = "#f59e0b"

        ctk.CTkLabel(
            f_plan, text=txt_plan, font=("Segoe UI", 10), text_color=cor_plan, justify="left"
        ).pack(anchor="w", padx=8, pady=6)

        if self.callback_abrir_workforce:
            ctk.CTkButton(
                self, text="🌐 Abrir / Preencher Ficha no WorkForce", font=("Segoe UI", 11, "bold"), height=36,
                fg_color="#0284c7", hover_color="#0369a1", command=self.callback_abrir_workforce
            ).pack(fill="x", padx=16, pady=(0, 8))

        # Botão Fechar
        ctk.CTkButton(
            self, text="Fechar", font=("Segoe UI", 11, "bold"), height=36,
            fg_color="#3f3f46", hover_color="#52525b", command=self.destroy
        ).pack(fill="x", padx=16, pady=(0, 14))

    def _copiar_texto(self, texto, btn=None):
        try:
            self.clipboard_clear()
            self.clipboard_append(texto)
            if btn:
                btn.configure(text="✓ Copiado!", fg_color="#10b981")
                self.after(1500, lambda: btn.configure(text="📋 Copiar Z61 (Novo Doc)", fg_color="#10b981"))
        except Exception:
            pass


class AbaTransferenciaEquipes:
    ALIASES_CADASTRO_EQUIPE = ("EQUIPES", "EQUIPE", "RECURSO", "RECURSOS")
    ALIASES_CADASTRO_ORDEM = ("ORDEM", "OSME", "NUMERO_OSME", "NUMORDEM")
    ALIASES_CADASTRO_DEP = ("DEPOSITO", "DEP", "LGORT", "DEPSITO")
    ALIASES_CADASTRO_SEG = ("SEGMENTO", "SEG", "SETOR")
    ALIASES_CADASTRO_VENC = ("VENCIMENTO", "VENCTO", "DATA_VENC")
    ALIASES_CADASTRO_MAT = ("MATERIAIS", "MATERIAL", "CODIGO", "CODIGOMATERIAL", "MATNR")
    ALIASES_CADASTRO_DESC = ("DESCRICAO", "DESCRIÇÃO", "DENOMINACAO", "DENOMINAÇÃO", "NOME", "DESCRIO")
    ALIASES_CADASTRO_UND = ("UND", "UNIDADE", "UN", "UM", "MEINS")

    def __init__(self, parent_frame, label_status_global=None, paletas=None, obter_tema_atual=None):
        self.parent_frame = parent_frame
        self.lbl_status_global = label_status_global
        self.paletas = paletas or DEFAULT_PALETAS
        self.obter_tema_atual = obter_tema_atual or (lambda: "Nexus")

        self.caminho_planilha = self._carregar_caminho_salvo()
        self.dados_cadastro = {}
        self.dados_fornecidos = {}
        self.dados_estoque = {}
        self.dados_materiais_cadastro = {}

        self.seriais_validados_atuais = []
        self.totais_materiais = OrderedDict()
        self.texto_workforce = ""
        self.processando_sap = False
        self.carregando = False
        self.mapa_nomes = {}

        threading.Thread(target=self._carregar_banco_nomes_async, daemon=True).start()

        self._construir_ui()
        if self.caminho_planilha and os.path.exists(self.caminho_planilha):
            self._iniciar_leitura_planilha_thread()

        # Inicia o verificador automático de pendências de sincronização
        self._iniciar_timer_pendencias()

    def _carregar_caminho_salvo(self):
        if os.path.exists(CAMINHO_CONFIG):
            try:
                with open(CAMINHO_CONFIG, "r", encoding="utf-8") as f:
                    dados = json.load(f)
                    return dados.get("caminho_planilha", "")
            except Exception:
                pass
        return ""

    def _salvar_caminho(self, caminho):
        self.caminho_planilha = caminho
        try:
            with open(CAMINHO_CONFIG, "w", encoding="utf-8") as f:
                json.dump({"caminho_planilha": caminho}, f, indent=2, ensure_ascii=False)
        except Exception:
            pass

    def _status(self, texto):
        if self.lbl_status_global and hasattr(self.lbl_status_global, "configure"):
            self.lbl_status_global.configure(text=f" {texto}")

    def _log(self, texto):
        agora = datetime.now().strftime("%H:%M:%S")
        self.txt_log.configure(state="normal")
        self.txt_log.insert(tk.END, f"[{agora}] {texto}\n")
        self.txt_log.see(tk.END)
        self.txt_log.configure(state="disabled")

    def _definir_estado_botoes_processar(self, executando=False):
        for btn in (getattr(self, "btn_processar_sap", None), getattr(self, "btn_processar_topo", None), getattr(self, "btn_processar_tabela", None)):
            if btn and hasattr(btn, "configure"):
                if executando:
                    btn.configure(state="disabled", text="⏳ Executando Z62 ➡️ Z61...", fg_color="#f59e0b")
                else:
                    texto = "🚀 Processar Transferência no SAP (Z62 ➡️ Z61)" if btn == getattr(self, "btn_processar_sap", None) else ("🚀 Processar Transferência (SAP)" if btn == getattr(self, "btn_processar_topo", None) else "🚀 Processar Transferência no SAP")
                    btn.configure(state="normal", text=texto, fg_color="#10b981")

    def _construir_ui(self):
        # 1. Barra de Ações Fixa no Rodapé (SEMPRE VISÍVEL - Não oculta na rolagem)
        self.f_barra_fixa = ctk.CTkFrame(self.parent_frame, corner_radius=10, fg_color="#172a45")
        self.f_barra_fixa.pack(side="bottom", fill="x", padx=10, pady=(4, 10))

        f_botoes_acao = ctk.CTkFrame(self.f_barra_fixa, fg_color="transparent")
        f_botoes_acao.pack(fill="x", padx=14, pady=10)

        self.btn_copiar_workforce = ctk.CTkButton(
            f_botoes_acao,
            text="📋 Copiar p/ WorkForce (;)",
            font=("Segoe UI", 11, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            height=38,
            command=self._copiar_workforce
        )
        self.btn_copiar_workforce.pack(side="left", padx=(0, 8))

        self.btn_ver_workforce = ctk.CTkButton(
            f_botoes_acao,
            text="👁️ Ver Lista WorkForce",
            font=("Segoe UI", 11),
            width=140,
            height=38,
            fg_color="#334155",
            hover_color="#475569",
            command=self._abrir_modal_workforce
        )
        self.btn_ver_workforce.pack(side="left", padx=(0, 8))

        self.btn_copiar_tabela = ctk.CTkButton(
            f_botoes_acao,
            text="📋 Copiar Tabela",
            font=("Segoe UI", 11),
            width=110,
            height=38,
            fg_color="#334155",
            hover_color="#475569",
            command=self._copiar_tabela
        )
        self.btn_copiar_tabela.pack(side="left")

        self.btn_processar_sap = ctk.CTkButton(
            f_botoes_acao,
            text="🚀 Processar Transferência no SAP (Z62 ➡️ Z61)",
            font=("Segoe UI", 13, "bold"),
            fg_color="#10b981",
            hover_color="#059669",
            height=38,
            command=self.processar_transferencia_sap_acao
        )
        self.btn_processar_sap.pack(side="right")

        # 2. Frame com Scroll para o Conteúdo
        self.frame_conteudo = ctk.CTkScrollableFrame(self.parent_frame, fg_color="transparent")
        self.frame_conteudo.pack(side="top", fill="both", expand=True, padx=10, pady=(10, 0))

        # 1. Card de Configuração da Planilha
        card_planilha = ctk.CTkFrame(self.frame_conteudo, corner_radius=10, fg_color="#172a45")
        card_planilha.pack(fill="x", pady=(0, 10))

        f_topo_plan = ctk.CTkFrame(card_planilha, fg_color="transparent")
        f_topo_plan.pack(fill="x", padx=14, pady=(10, 6))

        ctk.CTkLabel(
            f_topo_plan,
            text="📂 PLANILHA DE CONTROLE DE SERIAIS (ESTOQUE / FORNECIDOS / CADASTRO):",
            font=("Segoe UI", 11, "bold"),
            text_color="#38bdf8"
        ).pack(side="left")

        f_input_plan = ctk.CTkFrame(card_planilha, fg_color="transparent")
        f_input_plan.pack(fill="x", padx=14, pady=(0, 10))

        self.entry_caminho = ctk.CTkEntry(
            f_input_plan,
            placeholder_text="Selecione a planilha CONTROLE_SERIAIS.xlsx...",
            font=("Segoe UI", 11),
            height=34
        )
        self.entry_caminho.pack(side="left", fill="x", expand=True, padx=(0, 8))
        if self.caminho_planilha:
            self.entry_caminho.insert(0, self.caminho_planilha)

        btn_selecionar = ctk.CTkButton(
            f_input_plan,
            text="📁 Selecionar",
            font=("Segoe UI", 11, "bold"),
            width=100,
            height=34,
            command=self._selecionar_planilha
        )
        btn_selecionar.pack(side="left", padx=(0, 6))

        self.btn_recarregar = ctk.CTkButton(
            f_input_plan,
            text="🔄 Recarregar Dados",
            font=("Segoe UI", 11, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            width=135,
            height=34,
            command=self._iniciar_leitura_planilha_thread
        )
        self.btn_recarregar.pack(side="left")

        # Badges informativos
        f_badges = ctk.CTkFrame(card_planilha, fg_color="transparent")
        f_badges.pack(fill="x", padx=14, pady=(0, 10))

        self.badge_fornecidos = ctk.CTkLabel(f_badges, text="🏷️ Itens Fornecidos: 0", font=("Segoe UI", 11, "bold"), text_color="#10b981")
        self.badge_fornecidos.pack(side="left", padx=(0, 15))

        self.badge_equipes = ctk.CTkLabel(f_badges, text="📋 Equipes Cadastradas: 0", font=("Segoe UI", 11, "bold"), text_color="#38bdf8")
        self.badge_equipes.pack(side="left", padx=(0, 15))

        self.badge_estoque = ctk.CTkLabel(f_badges, text="📦 Itens Estoque: 0", font=("Segoe UI", 11), text_color="#94a3b8")
        self.badge_estoque.pack(side="left")

        # Banner de Pendências de Sincronização (Auto-Retry quando a planilha estiver em uso)
        self.f_banner_pendencias = ctk.CTkFrame(card_planilha, fg_color="#78350f", corner_radius=6)
        
        self.lbl_banner_pendencias = ctk.CTkLabel(
            self.f_banner_pendencias,
            text="",
            font=("Segoe UI", 11, "bold"),
            text_color="#fef3c7"
        )
        self.lbl_banner_pendencias.pack(side="left", padx=12, pady=6)

        self.btn_sinc_banner = ctk.CTkButton(
            self.f_banner_pendencias,
            text="🔄 Sincronizar Agora",
            font=("Segoe UI", 10, "bold"),
            fg_color="#10b981",
            hover_color="#059669",
            height=28,
            width=135,
            command=self._sincronizar_pendencias_manual
        )
        self.btn_sinc_banner.pack(side="right", padx=(6, 10), pady=4)

        self.btn_ver_banner = ctk.CTkButton(
            self.f_banner_pendencias,
            text="📋 Ver Detalhes",
            font=("Segoe UI", 10),
            fg_color="#0284c7",
            hover_color="#0369a1",
            height=28,
            width=95,
            command=self._abrir_modal_pendencias
        )
        self.btn_ver_banner.pack(side="right", padx=(0, 4), pady=4)

        # 2. Frame com Coluna Esquerda (Equipe Destino) e Direita (Seriais)
        self.frame_entradas = ctk.CTkFrame(self.frame_conteudo, corner_radius=10, fg_color="#172a45")
        self.frame_entradas.pack(fill="x", pady=(0, 10))
        self.frame_entradas.columnconfigure(0, weight=1)
        self.frame_entradas.columnconfigure(1, weight=1)

        # Coluna Esquerda: Equipe de Destino
        card_destino = ctk.CTkFrame(self.frame_entradas, fg_color="transparent")
        card_destino.grid(row=0, column=0, padx=14, pady=12, sticky="nsew")

        ctk.CTkLabel(
            card_destino,
            text="1. EQUIPE DE DESTINO (RECEBEDORA):",
            font=("Segoe UI", 11, "bold"),
            text_color="#38bdf8"
        ).pack(anchor="w", pady=(0, 4))

        f_eq = ctk.CTkFrame(card_destino, fg_color="transparent")
        f_eq.pack(fill="x", pady=(0, 8))

        self.entry_equipe_destino = ctk.CTkEntry(
            f_eq,
            placeholder_text="Código da Equipe de Destino (ex: CVC395, CSC300)",
            font=("Segoe UI", 12, "bold"),
            height=36
        )
        self.entry_equipe_destino.pack(side="left", fill="x", expand=True, padx=(0, 6))
        self.entry_equipe_destino.bind("<KeyRelease>", self._ao_digitar_equipe_destino)

        btn_consultar_eq = ctk.CTkButton(
            f_eq,
            text="Buscar",
            font=("Segoe UI", 11, "bold"),
            width=70,
            height=36,
            command=self._consultar_equipe_destino_atual
        )
        btn_consultar_eq.pack(side="left")

        # Data de Lançamento (Opcional - Fechamento de Mês no SAP)
        f_dt_lcto = ctk.CTkFrame(card_destino, fg_color="transparent")
        f_dt_lcto.pack(fill="x", pady=(0, 6))

        lbl_dt_lcto = ctk.CTkLabel(
            f_dt_lcto,
            text="📅 Data Lçto. (Fechamento SAP):",
            font=("Segoe UI", 11, "bold"),
            text_color="#f59e0b"
        )
        lbl_dt_lcto.pack(side="left", padx=(0, 8))

        self.entry_data_lancamento = ctk.CTkEntry(
            f_dt_lcto,
            placeholder_text="dd.mm.aaaa (opcional)",
            font=("Consolas", 11),
            width=150,
            height=28
        )
        self.entry_data_lancamento.pack(side="left")

        lbl_dt_help = ctk.CTkLabel(
            f_dt_lcto,
            text="* Deixe vazio para usar a data de hoje",
            font=("Segoe UI", 9, "italic"),
            text_color="#94a3b8"
        )
        lbl_dt_help.pack(side="left", padx=(8, 0))

        # Card informativo da OSME de Destino
        card_osme_info = ctk.CTkFrame(card_destino, corner_radius=8, fg_color="#0d1b2e")
        card_osme_info.pack(fill="x", pady=(0, 6))

        f_osme_top = ctk.CTkFrame(card_osme_info, fg_color="transparent")
        f_osme_top.pack(fill="x", padx=12, pady=(8, 2))

        ctk.CTkLabel(f_osme_top, text="OSME DESTINO ATIVA:", font=("Segoe UI", 10, "bold"), text_color="#94a3b8").pack(side="left")
        self.lbl_osme_status = ctk.CTkLabel(f_osme_top, text="Aguardando Equipe", font=("Segoe UI", 10, "bold"), text_color="#cbd5e1")
        self.lbl_osme_status.pack(side="right")

        self.lbl_osme_numero = ctk.CTkLabel(
            card_osme_info,
            text="—",
            font=("Consolas", 22, "bold"),
            text_color="#f57c00"
        )
        self.lbl_osme_numero.pack(padx=12, pady=(0, 4), anchor="w")

        f_osme_sub = ctk.CTkFrame(card_osme_info, fg_color="transparent")
        f_osme_sub.pack(fill="x", padx=12, pady=(0, 8))

        self.lbl_deposito_destino = ctk.CTkLabel(f_osme_sub, text="Depósito: —", font=("Segoe UI", 11), text_color="#e2e8f0")
        self.lbl_deposito_destino.pack(side="left", padx=(0, 15))

        self.lbl_segmento_destino = ctk.CTkLabel(f_osme_sub, text="Segmento: —", font=("Segoe UI", 11), text_color="#e2e8f0")
        self.lbl_segmento_destino.pack(side="left", padx=(0, 15))

        self.lbl_vencimento_destino = ctk.CTkLabel(f_osme_sub, text="Vencimento: —", font=("Segoe UI", 11), text_color="#e2e8f0")
        self.lbl_vencimento_destino.pack(side="left")

        # 3. Card WorkForce Completo (Ficha de Solicitação & Materiais)
        card_wf = ctk.CTkFrame(card_destino, corner_radius=10, fg_color="#10233b")
        card_wf.pack(fill="x", pady=(8, 0))

        f_wf_top = ctk.CTkFrame(card_wf, fg_color="transparent")
        f_wf_top.pack(fill="x", padx=12, pady=(8, 2))

        ctk.CTkLabel(
            f_wf_top,
            text="3. WORKFORCE — FICHA DE SOLICITAÇÃO & MATERIAIS:",
            font=("Segoe UI", 10, "bold"),
            text_color="#38bdf8"
        ).pack(side="left")

        self.chk_abrir_wf = ctk.CTkCheckBox(
            f_wf_top,
            text="Abrir Ficha no WorkForce",
            font=("Segoe UI", 10, "bold"),
            text_color="#e2e8f0",
            checkbox_width=18,
            checkbox_height=18
        )
        self.chk_abrir_wf.pack(side="right")
        self.chk_abrir_wf.select()

        # Grid interno com todos os campos da ficha original
        f_wf_grid = ctk.CTkFrame(card_wf, fg_color="transparent")
        f_wf_grid.pack(fill="x", padx=10, pady=(2, 6))
        f_wf_grid.columnconfigure(0, weight=1)
        f_wf_grid.columnconfigure(1, weight=1)

        creds_ini = carregar_credenciais_workforce()
        re_almox_default = creds_ini.get("usuario", "19595")

        # Linha 1: Data e Base
        ctk.CTkLabel(f_wf_grid, text="Data (DD-MM-AAAA):", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=0, column=0, sticky="w", padx=(0, 4), pady=(1, 0))
        ctk.CTkLabel(f_wf_grid, text="Base:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=0, column=1, sticky="w", padx=(4, 0), pady=(1, 0))

        self.entry_wf_data = ctk.CTkEntry(f_wf_grid, height=26, font=("Segoe UI", 10))
        self.entry_wf_data.grid(row=1, column=0, sticky="ew", padx=(0, 4), pady=(0, 4))
        self.entry_wf_data.insert(0, datetime.now().strftime("%d-%m-%Y"))

        self.combo_wf_base = ctk.CTkComboBox(
            f_wf_grid,
            height=26,
            values=["SANTO ANDRE", "MAUA", "NORTE", "INTERLAGOS", "SALVADOR", "SAO PAULO", "OUTRA"],
            font=("Segoe UI", 10)
        )
        self.combo_wf_base.grid(row=1, column=1, sticky="ew", padx=(4, 0), pady=(0, 4))
        self.combo_wf_base.set("SANTO ANDRE")

        # Linha 2: RE Resp Almox e RE Almox Entregue
        ctk.CTkLabel(f_wf_grid, text="RE Resp. Almox. *:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=2, column=0, sticky="w", padx=(0, 4), pady=(1, 0))
        ctk.CTkLabel(f_wf_grid, text="RE Almox. Entregue:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=2, column=1, sticky="w", padx=(4, 0), pady=(1, 0))

        self.entry_wf_re_almox = ctk.CTkEntry(f_wf_grid, height=26, font=("Segoe UI", 10), placeholder_text="Ex: 19595")
        self.entry_wf_re_almox.grid(row=3, column=0, sticky="ew", padx=(0, 4), pady=(0, 4))
        self.entry_wf_re_almox.insert(0, re_almox_default)

        self.entry_wf_re_almox_entregue = ctk.CTkEntry(f_wf_grid, height=26, font=("Segoe UI", 10), placeholder_text="Ex: 19595")
        self.entry_wf_re_almox_entregue.grid(row=3, column=1, sticky="ew", padx=(4, 0), pady=(0, 4))
        self.entry_wf_re_almox_entregue.insert(0, re_almox_default)

        # Linha 3: RE Eletricista 1 e RE Eletricista 2
        ctk.CTkLabel(f_wf_grid, text="RE Eletricista 1 *:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=4, column=0, sticky="w", padx=(0, 4), pady=(1, 0))
        ctk.CTkLabel(f_wf_grid, text="RE Eletricista 2:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=4, column=1, sticky="w", padx=(4, 0), pady=(1, 0))

        self.entry_re_eletricista = ctk.CTkEntry(f_wf_grid, height=26, font=("Segoe UI", 10), placeholder_text="Ex: 21475")
        self.entry_re_eletricista.grid(row=5, column=0, sticky="ew", padx=(0, 4), pady=(0, 1))
        self.entry_re_eletricista.bind("<KeyRelease>", self._ao_digitar_re_eletricista)

        self.entry_wf_re_eletricista_2 = ctk.CTkEntry(f_wf_grid, height=26, font=("Segoe UI", 10), placeholder_text="Ex: 22739")
        self.entry_wf_re_eletricista_2.grid(row=5, column=1, sticky="ew", padx=(4, 0), pady=(0, 1))
        self.entry_wf_re_eletricista_2.bind("<KeyRelease>", self._ao_digitar_re_eletricista_2)

        # Linha 4: Nomes dos Eletricistas
        self.lbl_nome_eletricista = ctk.CTkLabel(f_wf_grid, text="👤 Eletricista 1: não informado", font=("Segoe UI", 9), text_color="#94a3b8", anchor="w")
        self.lbl_nome_eletricista.grid(row=6, column=0, sticky="w", padx=(0, 4), pady=(0, 4))

        self.lbl_nome_eletricista_2 = ctk.CTkLabel(f_wf_grid, text="👤 Eletricista 2: não informado", font=("Segoe UI", 9), text_color="#94a3b8", anchor="w")
        self.lbl_nome_eletricista_2.grid(row=6, column=1, sticky="w", padx=(4, 0), pady=(0, 4))

        # Linha 5: Recurso (Destino) e Checkbox Salvar Automático
        ctk.CTkLabel(f_wf_grid, text="Recurso:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=7, column=0, sticky="w", padx=(0, 4), pady=(1, 0))

        self.chk_wf_salvar_auto = ctk.CTkCheckBox(
            f_wf_grid,
            text="Salvar auto no site",
            font=("Segoe UI", 9),
            text_color="#cbd5e1",
            checkbox_width=16,
            checkbox_height=16
        )
        self.chk_wf_salvar_auto.grid(row=7, column=1, sticky="e", padx=(4, 0), pady=(1, 0))

        self.entry_wf_recurso = ctk.CTkEntry(f_wf_grid, height=26, font=("Segoe UI", 10, "bold"), placeholder_text="Ex: CVC334")
        self.entry_wf_recurso.grid(row=8, column=0, columnspan=2, sticky="ew", pady=(0, 4))

        # Linha 6: Observação
        ctk.CTkLabel(f_wf_grid, text="Observação da Ficha:", font=("Segoe UI", 9, "bold"), text_color="#94a3b8").grid(row=9, column=0, columnspan=2, sticky="w", pady=(1, 0))

        self.txt_wf_observacao = ctk.CTkTextbox(f_wf_grid, height=44, font=("Segoe UI", 10))
        self.txt_wf_observacao.grid(row=10, column=0, columnspan=2, sticky="ew", pady=(0, 6))
        self.txt_wf_observacao.insert("1.0", "Transferência de medidores entre equipes")

        # Linha 7: Botões de Ação
        f_wf_botoes = ctk.CTkFrame(card_wf, fg_color="transparent")
        f_wf_botoes.pack(fill="x", padx=10, pady=(0, 8))

        self.btn_abrir_wf_agora = ctk.CTkButton(
            f_wf_botoes,
            text="🌐 Abrir Ficha no WorkForce",
            font=("Segoe UI", 10, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            height=28,
            command=self.abrir_ficha_workforce_manual
        )
        self.btn_abrir_wf_agora.pack(side="left", fill="x", expand=True, padx=(0, 6))

        self.btn_config_wf_login = ctk.CTkButton(
            f_wf_botoes,
            text="🔑 Credenciais",
            font=("Segoe UI", 10, "bold"),
            fg_color="#334155",
            hover_color="#475569",
            height=28,
            width=95,
            command=self._modal_credenciais_workforce
        )
        self.btn_config_wf_login.pack(side="right")

        # Coluna Direita: Seriais / Medidores a Transferir
        card_seriais = ctk.CTkFrame(self.frame_entradas, fg_color="transparent")
        card_seriais.grid(row=0, column=1, padx=14, pady=12, sticky="nsew")

        ctk.CTkLabel(
            card_seriais,
            text="2. SERIAIS / MEDIDORES A TRANSFERIR:",
            font=("Segoe UI", 11, "bold"),
            text_color="#38bdf8"
        ).pack(anchor="w", pady=(0, 4))

        self.txt_seriais = ctk.CTkTextbox(
            card_seriais,
            font=("Consolas", 12),
            height=100
        )
        self.txt_seriais.pack(fill="both", expand=True, pady=(0, 6))

        f_btn_seriais = ctk.CTkFrame(card_seriais, fg_color="transparent")
        f_btn_seriais.pack(fill="x")

        self.btn_validar = ctk.CTkButton(
            f_btn_seriais,
            text="🔍 Validar Planilha",
            font=("Segoe UI", 11, "bold"),
            fg_color="#0284c7",
            hover_color="#0369a1",
            height=34,
            command=self.validar_seriais_acao
        )
        self.btn_validar.pack(side="left", padx=(0, 6))

        self.btn_consultar_zmmcse = ctk.CTkButton(
            f_btn_seriais,
            text="🔬 Ordem SAP (ZMMCSE001)",
            font=("Segoe UI", 11, "bold"),
            fg_color="#4f46e5",
            hover_color="#4338ca",
            height=34,
            command=self.consultar_zmmcse001_acao
        )
        self.btn_consultar_zmmcse.pack(side="left", padx=(0, 6))

        btn_limpar = ctk.CTkButton(
            f_btn_seriais,
            text="Limpar",
            font=("Segoe UI", 11),
            width=65,
            height=34,
            fg_color="transparent",
            border_width=1,
            command=self._limpar_campos
        )
        btn_limpar.pack(side="left", padx=(0, 6))

        self.btn_processar_topo = ctk.CTkButton(
            f_btn_seriais,
            text="🚀 Processar Transferência (SAP)",
            font=("Segoe UI", 11, "bold"),
            fg_color="#10b981",
            hover_color="#059669",
            height=34,
            command=self.processar_transferencia_sap_acao
        )
        self.btn_processar_topo.pack(side="left", fill="x", expand=True)

        # 3. Chips de Resumo Agrupado por Material
        self.frame_resumo_materiais = ctk.CTkScrollableFrame(
            self.frame_conteudo, height=42, orientation="horizontal", corner_radius=8, fg_color="#172a45"
        )
        self.frame_resumo_materiais.pack(fill="x", pady=(0, 10))

        self.lbl_resumo_vazio = ctk.CTkLabel(
            self.frame_resumo_materiais,
            text="Valide os seriais para ver o agrupamento consolidado por material...",
            font=("Segoe UI", 10, "italic"),
            text_color="#94a3b8"
        )
        self.lbl_resumo_vazio.pack(side="left", padx=10, pady=6)

        # 4. Tabela de Conferência (Treeview)
        card_tabela = ctk.CTkFrame(self.frame_conteudo, corner_radius=10, fg_color="#172a45")
        card_tabela.pack(fill="both", expand=True, pady=(0, 10))

        f_tit_tab = ctk.CTkFrame(card_tabela, fg_color="transparent")
        f_tit_tab.pack(fill="x", padx=14, pady=(10, 6))

        ctk.CTkLabel(
            f_tit_tab,
            text="📊 CONFERÊNCIA DE TRANSFERÊNCIA (ORIGEM ➡️ DESTINO):",
            font=("Segoe UI", 11, "bold"),
            text_color="#38bdf8"
        ).pack(side="left")

        self.btn_processar_tabela = ctk.CTkButton(
            f_tit_tab,
            text="🚀 Processar Transferência no SAP",
            font=("Segoe UI", 10, "bold"),
            fg_color="#10b981",
            hover_color="#059669",
            height=28,
            command=self.processar_transferencia_sap_acao
        )
        self.btn_processar_tabela.pack(side="right", padx=(8, 0))

        self.lbl_contagem_tabela = ctk.CTkLabel(
            f_tit_tab,
            text="0 itens preparados",
            font=("Segoe UI", 11, "bold"),
            text_color="#cbd5e1"
        )
        self.lbl_contagem_tabela.pack(side="right")

        f_tree_container = ctk.CTkFrame(card_tabela, fg_color="transparent")
        f_tree_container.pack(fill="both", expand=True, padx=14, pady=(0, 10))

        colunas = ("item", "serial", "material", "origem", "osme_origem", "destino", "osme_destino", "status")
        self.tree = ttk.Treeview(f_tree_container, columns=colunas, show="headings", height=8)

        self.tree.heading("item", text="#")
        self.tree.heading("serial", text="Número de Série")
        self.tree.heading("material", text="Material")
        self.tree.heading("origem", text="Equipe Origem")
        self.tree.heading("osme_origem", text="OSME Origem")
        self.tree.heading("destino", text="Equipe Destino")
        self.tree.heading("osme_destino", text="OSME Destino")
        self.tree.heading("status", text="Validação")

        self.tree.column("item", width=40, anchor="center")
        self.tree.column("serial", width=140, anchor="center")
        self.tree.column("material", width=90, anchor="center")
        self.tree.column("origem", width=95, anchor="center")
        self.tree.column("osme_origem", width=110, anchor="center")
        self.tree.column("destino", width=95, anchor="center")
        self.tree.column("osme_destino", width=110, anchor="center")
        self.tree.column("status", width=220, anchor="w")

        scrollbar = ttk.Scrollbar(f_tree_container, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.tree.tag_configure("ok", foreground="#34d399")
        self.tree.tag_configure("bloqueado", foreground="#f87171")
        self.tree.tag_configure("alerta", foreground="#fbbf24")


        # 6. Painel de Log de Execução
        card_log = ctk.CTkFrame(self.frame_conteudo, corner_radius=10, fg_color="#172a45")
        card_log.pack(fill="x", pady=(0, 5))

        f_tit_log = ctk.CTkFrame(card_log, fg_color="transparent")
        f_tit_log.pack(fill="x", padx=14, pady=(8, 2))

        ctk.CTkLabel(f_tit_log, text="📜 LOG DE AUDITORIA E EXECUÇÃO:", font=("Segoe UI", 10, "bold"), text_color="#94a3b8").pack(side="left")

        self.txt_log = ctk.CTkTextbox(card_log, height=70, font=("Consolas", 10))
        self.txt_log.pack(fill="x", padx=14, pady=(0, 10))
        self.txt_log.configure(state="disabled")

        self._log("Módulo Transferência entre Equipes inicializado.")

    def _selecionar_planilha(self):
        caminho = filedialog.askopenfilename(
            title="Selecione a Planilha de Controle de Seriais",
            filetypes=[("Arquivos Excel", "*.xlsx *.xlsm"), ("Todos os Arquivos", "*.*")]
        )
        if caminho:
            self.entry_caminho.delete(0, tk.END)
            self.entry_caminho.insert(0, caminho)
            self._salvar_caminho(caminho)
            self._iniciar_leitura_planilha_thread()

    def _iniciar_leitura_planilha_thread(self):
        caminho = self.entry_caminho.get().strip()
        if not caminho or not os.path.exists(caminho):
            messagebox.showwarning("Transferência entre Equipes", "Selecione uma planilha válida.")
            return

        if self.carregando:
            return

        self.carregando = True
        self.btn_recarregar.configure(state="disabled", text="⏳ Lendo...")
        self._status("Lendo dados da planilha...")
        self._log("Iniciando leitura da planilha em segundo plano...")

        threading.Thread(target=self._ler_planilha_background, args=(caminho,), daemon=True).start()

    def _ler_planilha_background(self, caminho):
        try:
            wb = load_workbook(filename=caminho, read_only=True, data_only=True)

            cadastro_lido = {}
            materiais_lidos = {}
            if "CADASTRO" in wb.sheetnames:
                ws_cad = wb["CADASTRO"]
                cabecalho = None
                idx_equipe, idx_ordem, idx_dep, idx_seg, idx_venc = None, None, None, None, None
                idx_mat, idx_desc, idx_und = None, None, None

                # Detecta linha de cabeçalho dinamicamente
                for r_num, row in enumerate(ws_cad.iter_rows(max_row=10, values_only=True), 1):
                    for c_idx, val in enumerate(row):
                        norm = normalizar_texto(val)
                        if norm in self.ALIASES_CADASTRO_EQUIPE and idx_equipe is None:
                            idx_equipe = c_idx
                        elif norm in self.ALIASES_CADASTRO_ORDEM and idx_ordem is None:
                            idx_ordem = c_idx
                        elif norm in self.ALIASES_CADASTRO_DEP and idx_dep is None:
                            idx_dep = c_idx
                        elif norm in self.ALIASES_CADASTRO_SEG and idx_seg is None:
                            idx_seg = c_idx
                        elif norm in self.ALIASES_CADASTRO_VENC and idx_venc is None:
                            idx_venc = c_idx
                        elif norm in self.ALIASES_CADASTRO_MAT and idx_mat is None:
                            idx_mat = c_idx
                        elif norm in self.ALIASES_CADASTRO_DESC and idx_desc is None:
                            idx_desc = c_idx
                        elif norm in self.ALIASES_CADASTRO_UND and idx_und is None:
                            idx_und = c_idx

                    if idx_equipe is not None and idx_ordem is not None:
                        cabecalho = r_num
                        break

                # Padrão de contingência se os cabeçalhos não forem localizados (Layout padrão: E=Equipes, F=Depósito, G=Segmento, H=Ordem, I=Vencimento)
                if idx_equipe is None: idx_equipe = 4
                if idx_dep is None: idx_dep = idx_equipe + 1 if idx_equipe is not None else 5
                if idx_seg is None: idx_seg = idx_equipe + 2 if idx_equipe is not None else 6
                if idx_ordem is None: idx_ordem = idx_equipe + 3 if idx_equipe is not None else 7
                if idx_venc is None: idx_venc = idx_equipe + 4 if idx_equipe is not None else 8
                if idx_mat is None: idx_mat = 10
                if idx_desc is None: idx_desc = 11
                if idx_und is None: idx_und = 12
                if cabecalho is None: cabecalho = 1

                for row in ws_cad.iter_rows(min_row=cabecalho + 1, values_only=True):
                    if not row:
                        continue
                    # Leitura de Equipes
                    if len(row) > idx_equipe and row[idx_equipe]:
                        eq_val = row[idx_equipe]
                        eq_norm = normalizar_texto(eq_val)
                        ordem_val = row[idx_ordem] if len(row) > idx_ordem else ""
                        dep_val = row[idx_dep] if len(row) > idx_dep else ""
                        seg_val = row[idx_seg] if len(row) > idx_seg else ""
                        venc_val = row[idx_venc] if len(row) > idx_venc else ""

                        if eq_norm:
                            cadastro_lido[eq_norm] = {
                                "equipe_orig": str(eq_val).strip(),
                                "ordem": str(ordem_val or "").strip(),
                                "deposito": str(dep_val or "").strip(),
                                "segmento": str(seg_val or "").strip(),
                                "vencimento": formatar_data_excel(venc_val)
                            }

                    # Leitura de Materiais Cadastrados
                    if len(row) > idx_mat and row[idx_mat]:
                        mat_val = row[idx_mat]
                        mat_norm = normalizar_texto(mat_val)
                        desc_val = row[idx_desc] if len(row) > idx_desc else ""
                        und_val = row[idx_und] if len(row) > idx_und else "UN"

                        if mat_norm:
                            materiais_lidos[mat_norm] = {
                                "codigo": str(mat_val).strip(),
                                "descricao": str(desc_val or "").strip(),
                                "unidade": str(und_val or "UN").strip()
                            }

            fornecidos_lidos = {}
            if "FORNECIDOS" in wb.sheetnames:
                ws_forn = wb["FORNECIDOS"]
                idx_f_ser, idx_f_mat, idx_f_dep, idx_f_eq, idx_f_dt, idx_f_doc = None, None, None, None, None, None
                for c_idx, val in enumerate(next(ws_forn.iter_rows(max_row=1, values_only=True), ())):
                    norm = normalizar_texto(val)
                    if norm in ("SERIE", "SERIAL", "SERNR", "NUMEROSERIE") and idx_f_ser is None:
                        idx_f_ser = c_idx
                    elif norm in ("MATERIAL", "MAT", "MATNR", "CODIGO") and idx_f_mat is None:
                        idx_f_mat = c_idx
                    elif norm in ("DEPOSITO", "DEP", "LGORT") and idx_f_dep is None:
                        idx_f_dep = c_idx
                    elif norm in ("EQUIPE", "RECURSO") and idx_f_eq is None:
                        idx_f_eq = c_idx
                    elif norm in ("DATA", "DATADOC", "DATE") and idx_f_dt is None:
                        idx_f_dt = c_idx
                    elif norm in ("DOCSAP", "DOC", "DOCUMENTO", "NUMDOC") and idx_f_doc is None:
                        idx_f_doc = c_idx

                if idx_f_ser is None: idx_f_ser = 0
                if idx_f_mat is None: idx_f_mat = 1
                if idx_f_dep is None: idx_f_dep = 2
                if idx_f_eq is None: idx_f_eq = 3
                if idx_f_dt is None: idx_f_dt = 4
                if idx_f_doc is None: idx_f_doc = 5

                for row in ws_forn.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) <= idx_f_ser or not row[idx_f_ser]:
                        continue
                    serial_val = str(row[idx_f_ser]).strip()
                    s_norm = normalizar_texto(serial_val)
                    if not s_norm:
                        continue
                    mat_val = row[idx_f_mat] if len(row) > idx_f_mat else ""
                    dep_val = row[idx_f_dep] if len(row) > idx_f_dep else ""
                    eq_val = row[idx_f_eq] if len(row) > idx_f_eq else ""
                    data_val = row[idx_f_dt] if len(row) > idx_f_dt else ""
                    doc_val = row[idx_f_doc] if len(row) > idx_f_doc else ""

                    item_forn = {
                        "serial": serial_val,
                        "material": str(mat_val or "").strip(),
                        "deposito": str(dep_val or "").strip(),
                        "equipe": str(eq_val or "").strip(),
                        "data": formatar_data_excel(data_val),
                        "doc_sap": str(doc_val or "").strip()
                    }
                    s_chave = normalizar_serial_chave(serial_val)
                    fornecidos_lidos[s_norm] = item_forn
                    fornecidos_lidos[s_chave] = item_forn
                    if s_norm.isdigit():
                        fornecidos_lidos[s_norm.lstrip("0") or "0"] = item_forn
                        fornecidos_lidos[s_norm.zfill(18)] = item_forn

            estoque_lido = {}
            if "ESTOQUE" in wb.sheetnames:
                ws_est = wb["ESTOQUE"]
                idx_est_ser, idx_est_mat, idx_est_dep = 0, 1, 2
                for row in ws_est.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) <= idx_est_ser or not row[idx_est_ser]:
                        continue
                    serial_val = str(row[idx_est_ser]).strip()
                    s_norm = normalizar_texto(serial_val)
                    if not s_norm:
                        continue
                    mat_val = row[idx_est_mat] if len(row) > idx_est_mat else ""
                    dep_val = row[idx_est_dep] if len(row) > idx_est_dep else ""
                    estoque_lido[s_norm] = {
                        "serial": serial_val,
                        "material": str(mat_val or "").strip(),
                        "deposito": str(dep_val or "").strip()
                    }

            wb.close()

            # Notifica a interface na thread principal
            self.parent_frame.after(0, lambda: self._concluir_leitura_planilha(cadastro_lido, materiais_lidos, fornecidos_lidos, estoque_lido))

        except Exception as e:
            self.parent_frame.after(0, lambda e=e: self._erro_leitura_planilha(str(e)))

    def _concluir_leitura_planilha(self, cad, mats, forn, est):
        self.carregando = False
        self.dados_cadastro = cad
        self.dados_materiais = mats
        self.dados_fornecidos = forn
        self.dados_estoque = est

        self.btn_recarregar.configure(state="normal", text="🔄 Recarregar Dados")
        self.badge_fornecidos.configure(text=f"🏷️ Itens Fornecidos: {len(forn)}")
        self.badge_equipes.configure(text=f"📋 Equipes Cadastradas: {len(cad)}")
        self.badge_estoque.configure(text=f"📦 Itens Estoque: {len(est)}")

        self._status("Planilha carregada com sucesso!")
        self._log(f"Planilha carregada: {len(forn)} fornecidos, {len(cad)} equipes, {len(est)} estoque.")

        # Revalida se já havia algo digitado
        if self.entry_equipe_destino.get().strip():
            self._consultar_equipe_destino_atual()
        if self.txt_seriais.get("1.0", tk.END).strip():
            self.validar_seriais_acao()

    def _erro_leitura_planilha(self, erro_msg):
        self.carregando = False
        self.btn_recarregar.configure(state="normal", text="🔄 Recarregar Dados")
        self._status("Erro ao carregar planilha.")
        self._log(f"❌ Erro ao ler planilha: {erro_msg}")
        messagebox.showerror("Transferência entre Equipes — Erro Planilha", f"Não foi possível ler a planilha:\n\n{erro_msg}")

    def _carregar_banco_nomes_async(self):
        try:
            mapa = _carregar_banco_nomes()
            self.mapa_nomes = mapa or {}
            self.parent_frame.after(0, self._ao_digitar_re_eletricista)
        except Exception:
            self.mapa_nomes = {}

    def _ao_digitar_re_eletricista(self, event=None):
        re_txt = self.entry_re_eletricista.get().strip()
        if not re_txt:
            self.lbl_nome_eletricista.configure(
                text="👤 Eletricista 1: não informado",
                text_color="#94a3b8"
            )
            return

        if re_txt in self.mapa_nomes:
            nome = self.mapa_nomes[re_txt]
            self.lbl_nome_eletricista.configure(
                text=f"👤 {nome}",
                text_color="#34d399"
            )
        else:
            self.lbl_nome_eletricista.configure(
                text="👤 RE não localizado no Banco de Nomes",
                text_color="#fbbf24"
            )

    def _ao_digitar_re_eletricista_2(self, event=None):
        re_txt = self.entry_wf_re_eletricista_2.get().strip()
        if not re_txt:
            self.lbl_nome_eletricista_2.configure(
                text="👤 Eletricista 2: não informado",
                text_color="#94a3b8"
            )
            return

        if re_txt in self.mapa_nomes:
            nome = self.mapa_nomes[re_txt]
            self.lbl_nome_eletricista_2.configure(
                text=f"👤 {nome}",
                text_color="#34d399"
            )
        else:
            self.lbl_nome_eletricista_2.configure(
                text="👤 RE não localizado no Banco de Nomes",
                text_color="#fbbf24"
            )

    def _ao_digitar_equipe_destino(self, event=None):
        texto = self.entry_equipe_destino.get().strip()
        if hasattr(self, "entry_wf_recurso"):
            self.entry_wf_recurso.delete(0, tk.END)
            self.entry_wf_recurso.insert(0, texto.upper())
        if len(texto) >= 3 or not texto:
            self._consultar_equipe_destino_atual()

    def _consultar_equipe_destino_atual(self):
        eq_digitada = self.entry_equipe_destino.get().strip()
        if not eq_digitada:
            self.lbl_osme_numero.configure(text="—", text_color="#f57c00")
            self.lbl_osme_status.configure(text="Aguardando Equipe", text_color="#cbd5e1")
            self.lbl_deposito_destino.configure(text="Depósito: —")
            self.lbl_segmento_destino.configure(text="Segmento: —")
            self.lbl_vencimento_destino.configure(text="Vencimento: —")
            return

        eq_norm = normalizar_texto(eq_digitada)
        if eq_norm in self.dados_cadastro:
            dados = self.dados_cadastro[eq_norm]
            ordem = dados.get("ordem") or "Sem OSME"
            dep = dados.get("deposito") or "-"
            seg = dados.get("segmento") or "-"
            venc = dados.get("vencimento") or "-"

            self.lbl_osme_numero.configure(text=ordem, text_color="#34d399")
            self.lbl_osme_status.configure(text="Equipe Destino Localizada", text_color="#10b981")
            self.lbl_deposito_destino.configure(text=f"Depósito: {dep}")
            self.lbl_segmento_destino.configure(text=f"Segmento: {seg}")
            self.lbl_vencimento_destino.configure(text=f"Vencimento: {venc}")
            self._status(f"OSME Destino {ordem} localizada para {dados.get('equipe_orig')}")
        else:
            self.lbl_osme_numero.configure(text="Não Encontrada", text_color="#f87171")
            self.lbl_osme_status.configure(text="Não Cadastrada", text_color="#ef4444")
            self.lbl_deposito_destino.configure(text="Depósito: —")
            self.lbl_segmento_destino.configure(text="Segmento: —")
            self.lbl_vencimento_destino.configure(text="Vencimento: —")

    def obter_data_lancamento(self):
        """Retorna a data de lançamento informada na interface no formato dd.mm.aaaa, ou '' se em branco."""
        val = self.entry_data_lancamento.get().strip() if hasattr(self, "entry_data_lancamento") else ""
        if not val:
            return ""
        if not re.match(r"^\d{2}\.\d{2}\.\d{4}$", val):
            messagebox.showwarning(
                "Data de Lançamento Inválida",
                f"A data de lançamento informada ('{val}') está em formato incorreto.\n\n"
                f"Use o formato dd.mm.aaaa (ex: 01.10.2026) ou deixe o campo vazio para utilizar a data padrão de hoje."
            )
            return None
        try:
            d, m, y = map(int, val.split("."))
            if not (1 <= d <= 31 and 1 <= m <= 12 and 2000 <= y <= 2099):
                raise ValueError()
        except ValueError:
            messagebox.showwarning(
                "Data Inválida",
                f"A data de lançamento '{val}' não corresponde a uma data válida de calendário.\n\n"
                f"Informe no formato dd.mm.aaaa (ex: 01.10.2026)."
            )
            return None
        return val

    def validar_seriais_acao(self):
        """
        Regra fundamental:
        - O serial DEVE constar na aba FORNECIDOS para poder ser transferido.
        - Se constar, extrai a equipe de origem e a OSME da origem.
        - Se não constar, bloqueia o item com alerta.
        - Remove automaticamente qualquer entrada duplicada.
        """
        texto = self.txt_seriais.get("1.0", tk.END).strip()
        if not texto:
            messagebox.showwarning("Transferência entre Equipes", "Insira os números de série para validar.")
            return

        seriais_extraidos = extrair_lista_seriais(texto)
        if not seriais_extraidos:
            messagebox.showwarning("Transferência entre Equipes", "Nenhum serial válido foi identificado.")
            return

        # Deduplicação automática inteligente mantendo a primeira ocorrência
        seriais_unicos = []
        chaves_vistas = set()
        duplicados = []
        for s in seriais_extraidos:
            k = normalizar_serial_chave(s)
            if k in chaves_vistas:
                duplicados.append(s)
            else:
                chaves_vistas.add(k)
                seriais_unicos.append(s)

        if duplicados:
            self._log(f"🧹 {len(duplicados)} serial(is) duplicado(s) removido(s) automaticamente: {', '.join(duplicados[:5])}{'...' if len(duplicados) > 5 else ''}")

        seriais_extraidos = seriais_unicos

        eq_dest_digitada = self.entry_equipe_destino.get().strip().upper()
        eq_dest_norm = normalizar_texto(eq_dest_digitada)
        osme_dest = "—"
        dep_dest = "CL03"
        if eq_dest_norm in self.dados_cadastro:
            osme_dest = self.dados_cadastro[eq_dest_norm].get("ordem") or "—"
            dep_dest = self.dados_cadastro[eq_dest_norm].get("deposito") or "CL03"

        itens_processados = []
        itens_bloqueados_nao_fornecidos = []

        for serial in seriais_extraidos:
            s_norm = normalizar_texto(serial)
            s_chave = normalizar_serial_chave(serial)

            info_forn = (
                self.dados_fornecidos.get(s_chave)
                or self.dados_fornecidos.get(s_norm)
                or (self.dados_fornecidos.get(s_norm.lstrip("0") or "0") if s_norm.isdigit() else None)
                or (self.dados_fornecidos.get(s_norm.zfill(18)) if s_norm.isdigit() else None)
            )

            if info_forn:
                eq_origem = info_forn.get("equipe") or "-"
                eq_orig_norm = normalizar_texto(eq_origem)
                mat = info_forn.get("material") or "—"
                dep_orig = info_forn.get("deposito") or "CL03"

                osme_orig = "—"
                if eq_orig_norm in self.dados_cadastro:
                    osme_orig = self.dados_cadastro[eq_orig_norm].get("ordem") or "—"

                if eq_orig_norm == eq_dest_norm and eq_dest_norm:
                    status_val = f"⚠️ Origem e Destino são iguais ({eq_origem})"
                    tag = "alerta"
                    eh_bloqueado = True
                else:
                    status_val = f"✅ Custódia: {eq_origem} (OSME {osme_orig})"
                    tag = "ok"
                    eh_bloqueado = False

                itens_processados.append({
                    "serial": serial,
                    "material": mat,
                    "equipe_origem": eq_origem,
                    "osme_origem": osme_orig,
                    "deposito_origem": dep_orig,
                    "equipe_destino": eq_dest_digitada or "-",
                    "osme_destino": osme_dest,
                    "deposito_destino": dep_dest,
                    "status_val": status_val,
                    "tag": tag,
                    "eh_bloqueado": eh_bloqueado
                })

            else:
                # Bloqueado: Não está em FORNECIDOS
                itens_bloqueados_nao_fornecidos.append({"serial": serial})
                itens_processados.append({
                    "serial": serial,
                    "material": "—",
                    "equipe_origem": "Não Consta",
                    "osme_origem": "—",
                    "deposito_origem": "—",
                    "equipe_destino": eq_dest_digitada or "-",
                    "osme_destino": osme_dest,
                    "deposito_destino": dep_dest,
                    "status_val": "⛔ BLOQUEADO (Não consta em FORNECIDOS)",
                    "tag": "bloqueado",
                    "eh_bloqueado": True
                })

        self.seriais_validados_atuais = itens_processados
        self._renderizar_saida()

        if itens_bloqueados_nao_fornecidos:
            total_liberados = len(itens_processados) - len(itens_bloqueados_nao_fornecidos)
            self._log(f"⚠️ Atenção: {len(itens_bloqueados_nao_fornecidos)} seriais não constam em FORNECIDOS e foram bloqueados!")
            ModalAlertaNaoFornecidos(
                self.parent_frame,
                itens_bloqueados_nao_fornecidos,
                self.paletas,
                self.obter_tema_atual,
                callback_prosseguir=self.processar_transferencia_sap_acao,
                total_liberados=total_liberados
            )
        else:
            self._log(f"✅ Validação concluída: {len(seriais_extraidos)} seriais prontos para transferência!")

    def consultar_zmmcse001_acao(self):
        """
        Consulta em tempo real a transação ZMMCSE001 no SAP para obter a ordem (AUFNR)
        e a equipe em que cada medidor está alocado no momento.
        """
        if self.processando_sap:
            return

        texto = self.txt_seriais.get("1.0", tk.END).strip()
        seriais_para_consultar = extrair_lista_seriais(texto) if texto else []
        if not seriais_para_consultar and self.seriais_validados_atuais:
            seriais_para_consultar = [it["serial"] for it in self.seriais_validados_atuais]

        if not seriais_para_consultar:
            messagebox.showwarning("Transferência entre Equipes", "Insira os números de série no campo antes de consultar a ZMMCSE001.")
            return

        if not verificar_sap_aberto():
            messagebox.showerror("Transferência entre Equipes", "O SAP GUI precisa estar aberto e conectado para consultar a ZMMCSE001.")
            return

        self.btn_consultar_zmmcse.configure(state="disabled", text="⏳ Consultando SAP...")
        self._status("Consultando ordens de origem na ZMMCSE001...")
        self._log(f"🔬 Iniciando consulta ZMMCSE001 para {len(seriais_para_consultar)} seriais no SAP...")

        threading.Thread(
            target=self._executar_consulta_zmmcse_thread,
            args=(seriais_para_consultar,),
            daemon=True
        ).start()

    def _executar_consulta_zmmcse_thread(self, seriais):
        try:
            dados_sap = self._consultar_zmmcse001_sap(seriais)
            self.parent_frame.after(0, lambda: self._concluir_consulta_zmmcse(dados_sap, seriais))
        except Exception as e:
            self.parent_frame.after(0, lambda e=e: self._erro_consulta_zmmcse(str(e)))

    def _concluir_consulta_zmmcse(self, dados_sap, seriais_consultados=None):
        self.btn_consultar_zmmcse.configure(state="normal", text="🔬 Ordem SAP (ZMMCSE001)")
        if not dados_sap:
            self._status("Nenhum dado retornado pela ZMMCSE001.")
            self._log("⚠️ ZMMCSE001 não retornou registros no grid ALV.")
            messagebox.showinfo("ZMMCSE001 SAP", "A consulta ZMMCSE001 não retornou dados para os seriais informados.")
            return

        atualizados = 0
        # Se os seriais não estavam validados ainda, popula diretamente com os dados do SAP
        if not self.seriais_validados_atuais and seriais_consultados:
            for s in seriais_consultados:
                s_chave = normalizar_serial_chave(s)
                s_norm = normalizar_texto(s)
                info = (
                    dados_sap.get(s_chave)
                    or dados_sap.get(s_norm)
                    or (dados_sap.get(s_norm.lstrip("0") or "0") if s_norm.isdigit() else None)
                    or (dados_sap.get(s_norm.zfill(18)) if s_norm.isdigit() else None)
                    or {}
                )
                recurso = info.get("recurso", "—")
                aufnr = info.get("aufnr", "—")
                mat = info.get("matnr", "—")
                self.seriais_validados_atuais.append({
                    "serial": s,
                    "material": mat,
                    "equipe_origem": recurso,
                    "osme_origem": aufnr,
                    "deposito_origem": "CL03",
                    "status_val": f"✅ SAP ZMMCSE001: {recurso} (OSME {aufnr})",
                    "eh_bloqueado": False,
                    "tag": "liberado"
                })
                atualizados += 1
        else:
            for it in self.seriais_validados_atuais:
                s_chave = normalizar_serial_chave(it["serial"])
                s_norm = normalizar_texto(it["serial"])
                info = (
                    dados_sap.get(s_chave)
                    or dados_sap.get(s_norm)
                    or (dados_sap.get(s_norm.lstrip("0") or "0") if s_norm.isdigit() else None)
                    or (dados_sap.get(s_norm.zfill(18)) if s_norm.isdigit() else None)
                )
                if info:
                    aufnr = info.get("aufnr", "").strip()
                    recurso = info.get("recurso", "").strip()
                    mat = info.get("matnr", "").strip()
                    if aufnr:
                        it["osme_origem"] = aufnr
                    if recurso:
                        it["equipe_origem"] = recurso
                    if mat and (not it.get("material") or it.get("material") == "—"):
                        it["material"] = mat
                    it["status_val"] = f"✅ SAP ZMMCSE001: {it['equipe_origem']} (OSME {it['osme_origem']})"
                    atualizados += 1

        self._renderizar_saida()
        self._status(f"ZMMCSE001: {atualizados} seriais sincronizados com o SAP.")
        self._log(f"✅ Consulta ZMMCSE001 finalizada! {atualizados} itens sincronizados com o SAP.")
        messagebox.showinfo(
            "ZMMCSE001 — Ordens Localizadas no SAP",
            f"Consulta concluída com sucesso!\n\n{atualizados} medidores foram identificados no SAP com suas ordens de origem reais (AUFNR)."
        )

    def _erro_consulta_zmmcse(self, erro_msg):
        self.btn_consultar_zmmcse.configure(state="normal", text="🔬 Ordem SAP (ZMMCSE001)")
        self._status("Erro ao consultar ZMMCSE001.")
        self._log(f"❌ Erro ZMMCSE001: {erro_msg}")
        messagebox.showerror("Transferência entre Equipes — Erro ZMMCSE001", f"Falha ao consultar o SAP:\n\n{erro_msg}")

    def _consultar_zmmcse001_sap(self, seriais_lista):
        """
        Executa a transação ZMMCSE001 no SAP GUI via VBScript e extrai os dados do grid ALV:
        Serial, Material (MATNR), Ordem de Origem (AUFNR) e Recurso (PARNR).
        Retorna dicionário indexado por serial normalizado.
        """
        if not seriais_lista:
            return {}

        # 1. Copia seriais para o clipboard do Windows de forma resiliente
        copiar_texto_clipboard_resiliente("\r\n".join(seriais_lista))

        primeiro_serial = str(seriais_lista[0]).strip()
        qtd_seriais = len(seriais_lista)

        linhas_vbs = [
            'On Error Resume Next',
            'Dim SapGuiAuto, app, connection, session',
            'Set SapGuiAuto = GetObject("SAPGUI")',
            'If Not IsObject(SapGuiAuto) Then',
            '    WScript.Echo "ERRO: SAP GUI não encontrado."',
            '    WScript.Quit 1',
            'End If',
            'Set app = SapGuiAuto.GetScriptingEngine',
            'If Not IsObject(app) Then',
            '    WScript.Echo "ERRO: Scripting Engine do SAP não disponível."',
            '    WScript.Quit 1',
            'End If',
            'Set connection = app.Children(0)',
            'If Not IsObject(connection) Then',
            '    WScript.Echo "ERRO: Nenhuma conexão ativa no SAP."',
            '    WScript.Quit 1',
            'End If',
            'Set session = connection.Children(0)',
            'If Not IsObject(session) Then',
            '    WScript.Echo "ERRO: Nenhuma sessão ativa no SAP."',
            '    WScript.Quit 1',
            'End If',
            '',
            'session.findById("wnd[0]").maximize',
            'session.findById("wnd[0]/tbar[0]/okcd").text = "/nZMMCSE001"',
            'session.findById("wnd[0]").sendVKey 0',
            'WScript.Sleep 800',
            'Dim max_zmm, loop_zmm',
            'max_zmm = 0',
            'Do While session.Children.Count > 1 And max_zmm < 8',
            '    Err.Clear',
            '    session.findById("wnd[1]/tbar[0]/btn[0]").press',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]").sendVKey 0',
            '    End If',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]").close',
            '    End If',
            '    WScript.Sleep 250',
            '    max_zmm = max_zmm + 1',
            'Loop',
            '',
            'session.findById("wnd[0]/usr/btnBT003").press',
            'WScript.Sleep 600',
            '',
            'session.findById("wnd[0]/usr/subSUBSCREEN_9000:ZUDWM_PP350_CTRL_ANALISE_EQUIP:9000/radRD_EQUI").setFocus',
            'session.findById("wnd[0]/usr/subSUBSCREEN_9000:ZUDWM_PP350_CTRL_ANALISE_EQUIP:9000/radRD_EQUI").select',
            'WScript.Sleep 300',
            '',
            '\' 1. Preenche diretamente o campo Número de Série',
            'On Error Resume Next',
            f'session.findById("wnd[0]/usr/subSUBSCREEN_9000:ZUDWM_PP350_CTRL_ANALISE_EQUIP:9000/ctxtS_SERNR-LOW").text = "{primeiro_serial}"',
            f'session.findById("wnd[0]/usr/subSUBSCREEN_9000:ZUDWM_PP350_CTRL_ANALISE_EQUIP:9000/txtS_SERNR-LOW").text = "{primeiro_serial}"',
            'session.findById("wnd[0]").sendVKey 0',
            'WScript.Sleep 200',
            'On Error GoTo 0',
            '',
        ]

        if qtd_seriais > 1:
            linhas_vbs.extend([
                '\' 2. Como há múltiplos seriais, abre a janela de múltipla seleção e cola do clipboard',
                'On Error Resume Next',
                'session.findById("wnd[0]/usr/subSUBSCREEN_9000:ZUDWM_PP350_CTRL_ANALISE_EQUIP:9000/btn%_S_SERNR_%_APP_%-VALU_PUSH").press',
                'WScript.Sleep 400',
                'session.findById("wnd[1]/tbar[0]/btn[24]").press',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[0]").press',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[8]").press',
                'WScript.Sleep 400',
                'On Error GoTo 0',
                ''
            ])

        linhas_vbs.extend([
            'session.findById("wnd[0]/tbar[1]/btn[8]").press',
            'WScript.Sleep 1200',
            '',
            'Dim grid, totalLinhas, l, ser_val, aufnr_val, mat_val, parnr_val',
            'Set grid = session.findById("wnd[0]/usr/cntlCO_ITEM/shellcont/shell")',
            'If IsObject(grid) Then',
            '    totalLinhas = grid.RowCount',
            '    Dim visLinhas',
            '    visLinhas = 10',
            '    On Error Resume Next',
            '    visLinhas = grid.VisibleRowCount',
            '    If visLinhas <= 0 Then visLinhas = 10',
            '    On Error GoTo 0',
            '    ',
            '    Dim ini, fim',
            '    For ini = 0 To totalLinhas - 1 Step visLinhas',
            '        On Error Resume Next',
            '        grid.FirstVisibleRow = ini',
            '        WScript.Sleep 60',
            '        On Error GoTo 0',
            '        ',
            '        fim = ini + visLinhas - 1',
            '        If fim >= totalLinhas Then fim = totalLinhas - 1',
            '        ',
            '        For l = ini To fim',
            '            ser_val = ""',
            '            On Error Resume Next',
            '            ser_val = Trim(grid.getCellValue(l, "SERIAL"))',
            '            If ser_val = "" Then ser_val = Trim(grid.getCellValue(l, "SERNR"))',
            '            If ser_val = "" Then ser_val = Trim(grid.getCellValue(l, "EQUIPMENT"))',
            '            ',
            '            aufnr_val = Trim(grid.getCellValue(l, "AUFNR"))',
            '            If aufnr_val = "" Then aufnr_val = Trim(grid.getCellValue(l, "ORDEM"))',
            '            If aufnr_val = "" Then aufnr_val = Trim(grid.getCellValue(l, "NUM_ORDEM"))',
            '            ',
            '            mat_val = Trim(grid.getCellValue(l, "MATNR"))',
            '            If mat_val = "" Then mat_val = Trim(grid.getCellValue(l, "MATERIAL"))',
            '            ',
            '            parnr_val = Trim(grid.getCellValue(l, "PARNR"))',
            '            If parnr_val = "" Then parnr_val = Trim(grid.getCellValue(l, "RECURSO"))',
            '            If parnr_val = "" Then parnr_val = Trim(grid.getCellValue(l, "EQUIPE"))',
            '            On Error GoTo 0',
            '            ',
            '            If ser_val <> "" Then',
            '                WScript.Echo "ZMMCSE001_DADO:" & ser_val & "|" & aufnr_val & "|" & mat_val & "|" & parnr_val',
            '            End If',
            '        Next',
            '    Next',
            'End If',
            '',
            'session.findById("wnd[0]/tbar[0]/btn[15]").press',
            'WScript.Sleep 200',
            'session.findById("wnd[0]/tbar[0]/btn[15]").press',
            'WScript.Sleep 200',
            'session.findById("wnd[0]/tbar[0]/btn[15]").press'
        ])

        script_vbs = "\r\n".join(linhas_vbs)

        with tempfile.NamedTemporaryFile(mode="w", suffix=".vbs", delete=False, encoding="latin-1") as tf:
            tf.write(script_vbs)
            caminho_vbs = tf.name

        try:
            proc = subprocess.run(
                ["cscript.exe", "//nologo", caminho_vbs],
                capture_output=True,
                text=True,
                timeout=None
            )
            stdout = proc.stdout.strip()
        finally:
            try:
                os.unlink(caminho_vbs)
            except Exception:
                pass

        dados_por_serial = {}
        for line in stdout.splitlines():
            if "ZMMCSE001_DADO:" in line:
                conteudo = line.split("ZMMCSE001_DADO:")[1].strip()
                partes = conteudo.split("|")
                if len(partes) >= 4:
                    s_raw, o_raw, m_raw, r_raw = partes[0].strip(), partes[1].strip(), partes[2].strip(), partes[3].strip()
                    s_norm = normalizar_texto(s_raw)
                    s_chave = normalizar_serial_chave(s_raw)
                    if s_norm:
                        item_sap = {
                            "serial": s_raw,
                            "aufnr": o_raw,
                            "matnr": m_raw,
                            "recurso": r_raw
                        }
                        dados_por_serial[s_norm] = item_sap
                        dados_por_serial[s_chave] = item_sap
                        if s_norm.isdigit():
                            dados_por_serial[s_norm.lstrip("0") or "0"] = item_sap
                            dados_por_serial[s_norm.zfill(18)] = item_sap

        return dados_por_serial

    def _renderizar_saida(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        eq_dest = self.entry_equipe_destino.get().strip().upper()
        eq_dest_norm = normalizar_texto(eq_dest)
        osme_dest = "—"
        dep_dest = "CL03"
        if eq_dest_norm in self.dados_cadastro:
            osme_dest = self.dados_cadastro[eq_dest_norm].get("ordem") or "—"
            dep_dest = self.dados_cadastro[eq_dest_norm].get("deposito") or "CL03"

        for it in self.seriais_validados_atuais:
            it["equipe_destino"] = eq_dest or "-"
            it["osme_destino"] = osme_dest
            it["deposito_destino"] = dep_dest

        # Agrupamento de seriais por material
        agrupados = OrderedDict()
        for it in self.seriais_validados_atuais:
            m = it["material"]
            if m not in agrupados:
                agrupados[m] = []
            agrupados[m].append(it)

        idx_item = 1
        linhas_workforce = []
        totais_por_mat = OrderedDict()

        for mat, itens_mat in agrupados.items():
            totais_por_mat[mat] = itens_mat
            for it in itens_mat:
                self.tree.insert(
                    "",
                    tk.END,
                    values=(
                        idx_item,
                        it["serial"],
                        it["material"],
                        it["equipe_origem"],
                        it["osme_origem"],
                        it["equipe_destino"],
                        it["osme_destino"],
                        it["status_val"]
                    ),
                    tags=(it["tag"],)
                )
                if not it.get("eh_bloqueado"):
                    linhas_workforce.append(f"{it['serial']};")
                idx_item += 1

        self.totais_materiais = totais_por_mat
        self.texto_workforce = "\n".join(linhas_workforce)

        # Atualiza Chips de Resumo
        for w in self.frame_resumo_materiais.winfo_children():
            w.destroy()

        if not agrupados:
            lbl = ctk.CTkLabel(
                self.frame_resumo_materiais,
                text="Valide os seriais para ver o agrupamento consolidado por material...",
                font=("Segoe UI", 10, "italic"),
                text_color="#94a3b8"
            )
            lbl.pack(side="left", padx=10, pady=6)
        else:
            for mat, itens in agrupados.items():
                qtd_ok = sum(1 for x in itens if not x.get("eh_bloqueado"))
                qtd_bloq = sum(1 for x in itens if x.get("eh_bloqueado"))

                txt_badge = f"▸ {mat}: {qtd_ok} UN ✓"
                cor_bg = "#064e3b" if qtd_bloq == 0 else "#7f1d1d"
                cor_fg = "#34d399" if qtd_bloq == 0 else "#fca5a5"
                if qtd_bloq > 0:
                    txt_badge += f" ({qtd_bloq} bloq.)"

                chip = ctk.CTkFrame(self.frame_resumo_materiais, fg_color=cor_bg, corner_radius=6)
                chip.pack(side="left", padx=4, pady=4)
                ctk.CTkLabel(chip, text=f" {txt_badge} ", font=("Segoe UI", 11, "bold"), text_color=cor_fg).pack(padx=6, pady=2)

        total_liberados = sum(1 for x in self.seriais_validados_atuais if not x.get("eh_bloqueado"))
        self.lbl_contagem_tabela.configure(
            text=f"{len(self.seriais_validados_atuais)} itens ({total_liberados} liberados)"
        )

    def _copiar_workforce(self):
        if not self.texto_workforce.strip():
            messagebox.showwarning("Transferência entre Equipes", "Valide primeiro os seriais antes de copiar para o WorkForce.")
            return
        self.parent_frame.clipboard_clear()
        self.parent_frame.clipboard_append(self.texto_workforce)
        self.btn_copiar_workforce.configure(text="✓ Copiado p/ WorkForce!", fg_color="#10b981")
        self.parent_frame.after(2000, lambda: self.btn_copiar_workforce.configure(text="📋 Copiar p/ WorkForce (;)", fg_color="#0284c7"))
        self._log("📋 Lista formatada para WorkForce copiada para a área de transferência!")

    def _abrir_modal_workforce(self):
        if not self.texto_workforce.strip():
            messagebox.showwarning("Transferência entre Equipes", "Valide primeiro os seriais antes de visualizar a lista do WorkForce.")
            return
        ModalVisualizarWorkforce(self.parent_frame, self.texto_workforce, self.totais_materiais, self.paletas, self.obter_tema_atual)

    def abrir_ficha_workforce_manual(self):
        """Dispara a automação Selenium do WorkForce preenchendo todos os campos da ficha."""
        eq_dest = self.entry_wf_recurso.get().strip().upper() if hasattr(self, "entry_wf_recurso") and self.entry_wf_recurso.get().strip() else self.entry_equipe_destino.get().strip().upper()
        re_elet1 = self.entry_re_eletricista.get().strip() if hasattr(self, "entry_re_eletricista") else ""
        re_elet2 = self.entry_wf_re_eletricista_2.get().strip() if hasattr(self, "entry_wf_re_eletricista_2") else ""
        re_almox = self.entry_wf_re_almox.get().strip() if hasattr(self, "entry_wf_re_almox") else ""
        re_almox_entregue = self.entry_wf_re_almox_entregue.get().strip() if hasattr(self, "entry_wf_re_almox_entregue") else ""
        base_wf = self.combo_wf_base.get().strip() if hasattr(self, "combo_wf_base") else "SANTO ANDRE"
        data_f = self.entry_wf_data.get().strip().replace("/", "-") if hasattr(self, "entry_wf_data") else datetime.now().strftime("%d-%m-%Y")
        obs_wf = self.txt_wf_observacao.get("1.0", tk.END).strip() if hasattr(self, "txt_wf_observacao") else "Transferência de medidores entre equipes"
        salvar_auto = bool(self.chk_wf_salvar_auto.get()) if hasattr(self, "chk_wf_salvar_auto") else False

        # Consolidação dos materiais liberados
        materiais_lista = []
        mats_map = OrderedDict()
        for it in self.seriais_validados_atuais:
            if not it.get("eh_bloqueado"):
                m = it["material"]
                mats_map[m] = mats_map.get(m, 0) + 1

        for m, q in mats_map.items():
            if m and m != "—":
                materiais_lista.append((m, str(q)))

        if not materiais_lista:
            if self.txt_seriais.get("1.0", tk.END).strip():
                if messagebox.askyesno(
                    "WorkForce — Transferência",
                    "Os seriais ainda não foram validados.\nDeseja validar os seriais agora antes de preencher a ficha no WorkForce?",
                    parent=self.parent_frame
                ):
                    self.validar_seriais_acao()
                    mats_map = OrderedDict()
                    for it in self.seriais_validados_atuais:
                        if not it.get("eh_bloqueado"):
                            m = it["material"]
                            mats_map[m] = mats_map.get(m, 0) + 1
                    for m, q in mats_map.items():
                        if m and m != "—":
                            materiais_lista.append((m, str(q)))

        creds = carregar_credenciais_workforce()
        if not re_almox:
            re_almox = creds.get("usuario", "19595")
        if not re_almox_entregue:
            re_almox_entregue = re_almox

        dados_wf = {
            "data": data_f,
            "base": base_wf,
            "re_almox": re_almox,
            "re_almox_entregue": re_almox_entregue,
            "re_eletricista_1": re_elet1,
            "re_eletricista_2": re_elet2,
            "recurso": eq_dest,
            "observacao": obs_wf,
            "materiais": materiais_lista,
            "salvar_auto": salvar_auto
        }

        self._status("Iniciando automação do WorkForce...")
        self._log(f"🌐 Disparando preenchimento no WorkForce para recurso {eq_dest} (RE: {re_elet1 or 'não inf.'})...")

        def _callback_wf(msg):
            self.parent_frame.after(0, lambda: self._status(msg))
            self.parent_frame.after(0, lambda: self._log(f"[WorkForce] {msg}"))

        def _worker():
            try:
                sucesso, resultado = executar_fluxo_workforce(
                    dados=dados_wf,
                    callback_status=_callback_wf,
                    login_usuario=creds.get("usuario"),
                    login_senha=creds.get("senha"),
                    mapa_nomes=self.mapa_nomes
                )
                if sucesso:
                    self.parent_frame.after(0, lambda: self._status("Ficha aberta com sucesso no WorkForce!"))
                    self.parent_frame.after(0, lambda: self._log("✅ Ficha e materiais preenchidos com sucesso no WorkForce."))
                else:
                    self.parent_frame.after(0, lambda: self._status("Atenção ao abrir WorkForce."))
                    self.parent_frame.after(0, lambda: self._log(f"⚠️ WorkForce: {resultado}"))
            except Exception as e:
                self.parent_frame.after(0, lambda: self._status("Erro ao abrir WorkForce."))
                self.parent_frame.after(0, lambda: self._log(f"❌ Erro WorkForce: {e}"))

        threading.Thread(target=_worker, daemon=True).start()

    def _modal_credenciais_workforce(self):
        creds = carregar_credenciais_workforce()
        top = ctk.CTkToplevel(self.parent_frame)
        top.title("Credenciais WorkForce")
        top.geometry("340x220")
        top.resizable(False, False)
        top.transient(self.parent_frame.winfo_toplevel())
        top.grab_set()

        ctk.CTkLabel(top, text="🔑 Configuração de Acesso WorkForce", font=("Segoe UI", 12, "bold")).pack(pady=(15, 10))

        f = ctk.CTkFrame(top, fg_color="transparent")
        f.pack(fill="x", padx=20)

        ctk.CTkLabel(f, text="Usuário (RE):", font=("Segoe UI", 11)).grid(row=0, column=0, sticky="w", pady=4)
        e_usr = ctk.CTkEntry(f, width=180)
        e_usr.grid(row=0, column=1, pady=4)
        e_usr.insert(0, creds.get("usuario", "19595"))

        ctk.CTkLabel(f, text="Senha:", font=("Segoe UI", 11)).grid(row=1, column=0, sticky="w", pady=4)
        e_pwd = ctk.CTkEntry(f, width=180, show="*")
        e_pwd.grid(row=1, column=1, pady=4)
        e_pwd.insert(0, creds.get("senha", "979648"))

        def _salvar():
            u = e_usr.get().strip()
            s = e_pwd.get().strip()
            if u and s:
                caminho = obter_caminho_config("config_fichas.json")
                try:
                    with open(caminho, "w", encoding="utf-8") as fp:
                        json.dump({"usuario": u, "senha": s}, fp, indent=4)
                    if hasattr(self, "entry_wf_re_almox"):
                        self.entry_wf_re_almox.delete(0, tk.END)
                        self.entry_wf_re_almox.insert(0, u)
                    top.destroy()
                    messagebox.showinfo("WorkForce", "Credenciais salvas com sucesso!", parent=self.parent_frame)
                except Exception as ex:
                    messagebox.showerror("Erro", f"Falha ao salvar: {ex}", parent=top)

        btn_f = ctk.CTkFrame(top, fg_color="transparent")
        btn_f.pack(fill="x", padx=20, pady=(15, 10))
        ctk.CTkButton(btn_f, text="Salvar", width=120, command=_salvar).pack(side="right", padx=(5, 0))
        ctk.CTkButton(btn_f, text="Cancelar", width=90, fg_color="#64748b", command=top.destroy).pack(side="right")

    def _copiar_tabela(self):
        itens = self.tree.get_children()
        if not itens:
            return
        linhas = ["Item\tSerial\tMaterial\tOrigem\tOSME Origem\tDestino\tOSME Destino\tStatus"]
        for it in itens:
            vals = self.tree.item(it, "values")
            linhas.append("\t".join(str(v) for v in vals))
        texto = "\n".join(linhas)
        self.parent_frame.clipboard_clear()
        self.parent_frame.clipboard_append(texto)
        messagebox.showinfo("Transferência entre Equipes", "Tabela copiada para a área de transferência!")

    def _limpar_campos(self):
        self.txt_seriais.delete("1.0", tk.END)
        if hasattr(self, "entry_data_lancamento"):
            self.entry_data_lancamento.delete(0, tk.END)
        self.seriais_validados_atuais = []
        self.totais_materiais = OrderedDict()
        self.texto_workforce = ""
        self._renderizar_saida()

    def processar_transferencia_sap_acao(self):
        """
        Inicia o processamento no SAP em duas etapas consecutivas:
        1. MIGO Z62: Estorno das ordens de origem.
        2. MIGO Z61: Saída para a ordem de destino.
        """
        if self.processando_sap:
            return

        eq_dest = self.entry_equipe_destino.get().strip().upper()
        if not eq_dest:
            messagebox.showwarning("Transferência entre Equipes", "Informe a equipe de destino.")
            self.entry_equipe_destino.focus()
            return

        eq_dest_norm = normalizar_texto(eq_dest)
        if eq_dest_norm not in self.dados_cadastro:
            messagebox.showwarning("Transferência entre Equipes", f"A equipe de destino '{eq_dest}' não foi encontrada na aba CADASTRO.")
            return

        osme_dest = self.dados_cadastro[eq_dest_norm].get("ordem")
        dep_dest = self.dados_cadastro[eq_dest_norm].get("deposito") or "CL03"
        if not osme_dest or osme_dest in ("—", "-", "Sem OSME"):
            messagebox.showwarning("Transferência entre Equipes", f"A equipe de destino '{eq_dest}' não possui número de OSME cadastrado.")
            return

        itens_liberados = [it for it in self.seriais_validados_atuais if not it.get("eh_bloqueado")]
        if not itens_liberados:
            texto = self.txt_seriais.get("1.0", tk.END).strip()
            if texto:
                self.validar_seriais_acao()
                itens_liberados = [it for it in self.seriais_validados_atuais if not it.get("eh_bloqueado")]
            if not itens_liberados:
                messagebox.showwarning("Transferência entre Equipes", "Nenhum serial liberado para transferência.")
                return

        # Verifica se alguma OSME de origem está faltando
        for it in itens_liberados:
            if not it.get("osme_origem") or it["osme_origem"] in ("—", "-", "Sem OSME"):
                messagebox.showerror(
                    "Transferência entre Equipes",
                    f"A equipe de origem '{it.get('equipe_origem')}' do serial {it.get('serial')} não possui OSME válida cadastrada para devolução Z62."
                )
                return

        # Monta resumo prévio agrupado por material
        mats_map = OrderedDict()
        origens_set = set()
        for it in itens_liberados:
            m = it["material"]
            mats_map[m] = mats_map.get(m, 0) + 1
            origens_set.add(it["equipe_origem"])

        resumo_mats_txt = "\n".join([f"  ▸ {m} — {qtd} un" for m, qtd in mats_map.items()])
        origens_txt = ", ".join(sorted(origens_set))

        data_lcto = self.obter_data_lancamento()
        if data_lcto is None:
            return

        txt_dt = f"• Data de Lançamento (MIGO): {data_lcto} (Fechamento SAP)\n" if data_lcto else "• Data de Lançamento (MIGO): Padrão (Hoje)\n"

        msg_confirma = (
            f"Deseja realmente processar a TRANSFERÊNCIA ENTRE EQUIPES no SAP?\n\n"
            f"Equipe(s) de Origem: {origens_txt}\n"
            f"Equipe de Destino: {eq_dest} (OSME: {osme_dest} | Depósito: {dep_dest})\n\n"
            f"{txt_dt}"
            f"Materiais a transferir:\n{resumo_mats_txt}\n\n"
            f"O SAP executará:\n"
            f"  1. MIGO Z62 (Estorno ao estoque pela OSME de origem)\n"
            f"  2. MIGO Z61 (Saída para a OSME de destino)\n\n"
            f"A planilha de controle será atualizada automaticamente."
        )
        if not messagebox.askyesno("Confirmar Transferência SAP (Z62 ➡️ Z61)", msg_confirma, icon="question"):
            return

        if not verificar_sap_aberto():
            messagebox.showerror("Transferência entre Equipes", "O SAP GUI precisa estar aberto e conectado para processar a transferência.")
            return

        self.processando_sap = True
        self._definir_estado_botoes_processar(executando=True)
        self._status("Processando Transferência no SAP MIGO...")
        self._log(f"Iniciando transferência de {len(itens_liberados)} seriais: {origens_txt} ➡️ {eq_dest}")
        if data_lcto:
            self._log(f"📅 Data de Lançamento configurada: {data_lcto}")

        threading.Thread(
            target=self._executar_transferencia_sap_thread,
            args=(eq_dest, osme_dest, dep_dest, itens_liberados, data_lcto),
            daemon=True
        ).start()

    def _executar_transferencia_sap_thread(self, eq_dest, osme_dest, dep_dest, itens_liberados, data_lancamento=None):
        try:
            # ETAPA 1: Consulta automática da ZMMCSE001 no SAP para obter a ordem real de origem (AUFNR)
            self.parent_frame.after(0, lambda: self._status("Consultando ordens de origem no SAP (ZMMCSE001)..."))
            self.parent_frame.after(0, lambda: self._log("🔬 Consultando ordem exata no SAP via ZMMCSE001..."))

            seriais_alvo = [it["serial"] for it in itens_liberados]
            dados_sap = self._consultar_zmmcse001_sap(seriais_alvo)

            if dados_sap:
                for it in itens_liberados:
                    s_norm = normalizar_texto(it["serial"])
                    if s_norm in dados_sap:
                        info_sap = dados_sap[s_norm]
                        aufnr_sap = info_sap.get("aufnr", "").strip()
                        rec_sap = info_sap.get("recurso", "").strip()
                        mat_sap = info_sap.get("matnr", "").strip()
                        if aufnr_sap:
                            it["osme_origem"] = aufnr_sap
                            self.parent_frame.after(
                                0, lambda s=it["serial"], o=aufnr_sap, r=rec_sap: self._log(f"  ▸ Serial {s}: Ordem SAP {o} (Recurso: {r})")
                            )
                        if rec_sap:
                            it["equipe_origem"] = rec_sap
                        if mat_sap and (not it.get("material") or it.get("material") == "—"):
                            it["material"] = mat_sap

            # Valida se algum item ainda está sem OSME de origem
            sem_ordem = [it["serial"] for it in itens_liberados if not it.get("osme_origem") or it["osme_origem"] in ("—", "-", "Sem OSME")]
            if sem_ordem:
                self.parent_frame.after(
                    0, lambda: self._finalizar_com_aviso_sap(
                        f"Não foi possível localizar a Ordem de Origem (OSME) no SAP (ZMMCSE001) nem na Planilha para os seguintes seriais:\n\n{', '.join(sem_ordem)}"
                    )
                )
                return

            self.parent_frame.after(0, self._renderizar_saida)

            # ETAPA 2: Executa MIGO Z62 (Estorno) e Z61 (Saída para Destino)
            self.parent_frame.after(0, lambda: self._status("Executando MIGO Z62 (Estorno) e Z61 (Saída)..."))
            self.parent_frame.after(0, lambda: self._log("🚀 Executando MIGO Z62 ➡️ Z61 no SAP..."))

            # Agrupa os itens liberados por material e OSME de origem para a Z62
            # Para a Z61, agrupa por material com a OSME de destino
            script_vbs = self._gerar_script_migo_duplo_vbs(eq_dest, osme_dest, dep_dest, itens_liberados, data_lancamento=data_lancamento)

            with tempfile.NamedTemporaryFile(mode="w", suffix=".vbs", delete=False, encoding="latin-1") as tf:
                tf.write(script_vbs)
                caminho_vbs = tf.name

            proc = subprocess.run(
                ["cscript.exe", "//nologo", caminho_vbs],
                capture_output=True,
                text=True,
                timeout=None
            )
            stdout = proc.stdout.strip()
            stderr = proc.stderr.strip()

            try:
                os.unlink(caminho_vbs)
            except Exception:
                pass

            if proc.returncode != 0:
                err_msg = stderr or stdout or "Erro desconhecido ao executar script SAP MIGO."
                self.parent_frame.after(0, lambda: self._finalizar_com_erro(err_msg))
                return

            # Extração dos Documentos SAP Z62 e Z61
            doc_z62 = ""
            doc_z61 = ""

            for line in stdout.splitlines():
                if "DOC_Z62:" in line:
                    doc_z62 = line.split("DOC_Z62:")[1].strip()
                elif "DOC_Z61:" in line:
                    doc_z61 = line.split("DOC_Z61:")[1].strip()
                elif "ERRO_SAP:" in line:
                    erro_sap = line.split("ERRO_SAP:")[1].strip()
                    self.parent_frame.after(0, lambda: self._finalizar_com_aviso_sap(erro_sap))
                    return

            if not doc_z62 or not doc_z61:
                # Fallback: regex por números de documento de 10 dígitos
                matches = re.findall(r"\b(49\d{8}|50\d{8}|\d{10})\b", stdout)
                if len(matches) >= 2:
                    doc_z62, doc_z61 = matches[0], matches[1]
                elif len(matches) == 1:
                    doc_z61 = matches[0]

            if not doc_z61:
                msg_completa = stdout if stdout else "O SAP não retornou número de documento."
                self.parent_frame.after(0, lambda: self._finalizar_com_aviso_sap(msg_completa))
                return

            self._log(f"🎉 SAP Concluído com Sucesso! Doc Z62: {doc_z62} | Doc Z61: {doc_z61}")

            # Atualização da Planilha (exclui registro anterior em FORNECIDOS e reinsere com nova saída)
            sucesso_plan, modo_plan = self._atualizar_planilha_transferencia(eq_dest, dep_dest, doc_z61, itens_liberados)

            origens_set = sorted(set(it["equipe_origem"] for it in itens_liberados))
            origem_resumo = ", ".join(origens_set)

            self.parent_frame.after(
                0,
                lambda: self._concluir_sucesso_transferencia(
                    doc_z62, doc_z61, origem_resumo, eq_dest, osme_dest, itens_liberados, sucesso_plan, modo_plan
                )
            )

        except Exception as e:
            self.parent_frame.after(0, lambda e=e: self._finalizar_com_erro(str(e)))

    def _gerar_script_migo_duplo_vbs(self, eq_dest, osme_dest, dep_dest, itens_liberados, data_lancamento=None):
        """
        Gera VBScript automatizado que realiza:
        Etapa 1: MIGO Z62 (Estorno da OSME de Origem ao Depósito).
        Etapa 2: MIGO Z61 (Saída do Depósito para a OSME de Destino).
        """
        # Agrupamento de itens para Z62: por (Material, OSME Origem, Depósito Origem)
        grupos_z62 = OrderedDict()
        for it in itens_liberados:
            chave = (it["material"], it["osme_origem"], it["deposito_origem"])
            if chave not in grupos_z62:
                grupos_z62[chave] = []
            grupos_z62[chave].append(it["serial"])

        # Agrupamento de itens para Z61: por Material
        grupos_z61 = OrderedDict()
        for it in itens_liberados:
            m = it["material"]
            if m not in grupos_z61:
                grupos_z61[m] = []
            grupos_z61[m].append(it["serial"])

        linhas_vbs = [
            'On Error Resume Next',
            'Dim SapGuiAuto, app, connection, session',
            'Set SapGuiAuto = GetObject("SAPGUI")',
            'If Not IsObject(SapGuiAuto) Then',
            '    WScript.Echo "ERRO: SAP GUI não encontrado."',
            '    WScript.Quit 1',
            'End If',
            'Set app = SapGuiAuto.GetScriptingEngine',
            'If Not IsObject(app) Then',
            '    WScript.Echo "ERRO: Scripting Engine do SAP não disponível."',
            '    WScript.Quit 1',
            'End If',
            'Set connection = app.Children(0)',
            'If Not IsObject(connection) Then',
            '    WScript.Echo "ERRO: Nenhuma conexão ativa no SAP."',
            '    WScript.Quit 1',
            'End If',
            'Set session = connection.Children(0)',
            'If Not IsObject(session) Then',
            '    WScript.Echo "ERRO: Nenhuma sessão ativa no SAP."',
            '    WScript.Quit 1',
            'End If',
            '',
            'Sub FecharModais(maxVezes)',
            '    Dim mCount',
            '    mCount = 0',
            '    On Error Resume Next',
            '    Do While session.Children.Count > 1 And mCount < maxVezes',
            '        Err.Clear',
            '        session.findById("wnd[1]/tbar[0]/btn[0]").press',
            '        If Err.Number <> 0 Then',
            '            Err.Clear',
            '            session.findById("wnd[1]").sendVKey 0',
            '        End If',
            '        If Err.Number <> 0 Then',
            '            Err.Clear',
            '            session.findById("wnd[1]/usr/btnSPOP-OPTION1").press',
            '        End If',
            '        If Err.Number <> 0 Then',
            '            Err.Clear',
            '            session.findById("wnd[1]").close',
            '        End If',
            '        Err.Clear',
            '        WScript.Sleep 300',
            '        mCount = mCount + 1',
            '    Loop',
            '    Err.Clear',
            'End Sub',
            '',
            'Sub PreencherDataLancamento(cId, dtLcto)',
            '    If Trim(dtLcto) = "" Then Exit Sub',
            '    On Error Resume Next',
            '    session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & cId & "/subSUB_HEADER:SAPLMIGO:0101/subSUB_HEADER:SAPLMIGO:0100/tabsTS_GOHEAD/tabpOK_GOHEAD_GENERAL").select',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & cId & "/subSUB_HEADER:SAPLMIGO:0102/btnBUTTON_HEADER").press',
            '        WScript.Sleep 250',
            '        session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & cId & "/subSUB_HEADER:SAPLMIGO:0101/subSUB_HEADER:SAPLMIGO:0100/tabsTS_GOHEAD/tabpOK_GOHEAD_GENERAL").select',
            '    End If',
            '    Err.Clear',
            '    session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & cId & "/subSUB_HEADER:SAPLMIGO:0101/subSUB_HEADER:SAPLMIGO:0100/tabsTS_GOHEAD/tabpOK_GOHEAD_GENERAL/ssubSUB_TS_GOHEAD_GENERAL:SAPLMIGO:0112/ctxtGOHEAD-BUDAT").text = dtLcto',
            '    session.findById("wnd[0]").sendVKey 0',
            '    WScript.Sleep 350',
            '    FecharModais 4',
            '    Err.Clear',
            'End Sub',
            '',
            'session.findById("wnd[0]").maximize',
            'Dim carrier',
            'carrier = "0003"',
            '',
            '\' =========================================================================',
            '\' ETAPA 1: MIGO Z62 (ESTORNO AO DEPÓSITO DAS ORDENS DE ORIGEM)',
            '\' =========================================================================',
            'session.findById("wnd[0]/tbar[0]/okcd").text = "/nMIGO"',
            'session.findById("wnd[0]").sendVKey 0',
            'WScript.Sleep 800',
            'FecharModais 8',
            '',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:0003/subSUB_FIRSTLINE:SAPLMIGO:0010/ctxtGODEFAULT_TV-BWART").setFocus',
            'If Err.Number <> 0 Then',
            '    Err.Clear',
            '    carrier = "0002"',
            'End If',
            '',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-ACTION").setFocus',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-ACTION").key = "A01"',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-REFDOC").setFocus',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-REFDOC").key = "R10"',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/ctxtGODEFAULT_TV-BWART").text = "Z62"',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/ctxtGODEFAULT_TV-BWART").setFocus',
            'session.findById("wnd[0]").sendVKey 0',
            'WScript.Sleep 500',
        ]

        dt_lcto_limpa = str(data_lancamento or "").strip()
        if dt_lcto_limpa:
            linhas_vbs.extend([
                f'PreencherDataLancamento carrier, "{dt_lcto_limpa}"',
                'WScript.Sleep 300',
            ])

        # Itens da Z62
        idx_z62 = 0
        total_grupos_z62 = len(grupos_z62)
        for (mat_cod, osme_origem, dep_origem), seriais_list in grupos_z62.items():
            qtd = str(len(seriais_list))
            dep_clean = str(dep_origem or "CL03").strip()
            recebedor = f"EMPR 1{dep_clean}"

            linhas_vbs.extend([
                f"' --- Z62 Item {idx_z62 + 1}: Material {mat_cod} | Qtd {qtd} | OSME Origem {osme_origem} ---"
            ])

            seriais_clip = vb_escape_clipboard("\r\n".join(seriais_list))

            linhas_vbs.extend([
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_MATERIAL").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_MATERIAL/ssubSUB_TS_GOITEM_MATERIAL:SAPLMIGO:0310/ctxtGOITEM-MAKTX").text = "{mat_cod}"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 300',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_QUANTITIES").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_QUANTITIES/ssubSUB_TS_GOITEM_QUANTITIES:SAPLMIGO:0315/txtGOITEM-ERFMG").text = "{qtd}"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 300',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT.").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT./ssubSUB_TS_GOITEM_DESTINATION:SAPLMIGO:0325/ctxtGOITEM-NAME1").text = "EMPR"',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT./ssubSUB_TS_GOITEM_DESTINATION:SAPLMIGO:0325/ctxtGOITEM-LGOBE").text = "{dep_clean}"',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT./ssubSUB_TS_GOITEM_DESTINATION:SAPLMIGO:0325/txtGOITEM-WEMPF").text = "{recebedor}"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 300',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_ACCOUNT").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_ACCOUNT/ssubSUB_TS_GOITEM_ACCOUNT:SAPLMIGO:0345/ssubSUB_ACCOUNTINGBLOCK:SAPLKACB:1003/ctxtCOBL-AUFNR").text = "{osme_origem}"',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_ACCOUNT/ssubSUB_TS_GOITEM_ACCOUNT:SAPLMIGO:0345/ssubSUB_ACCOUNTINGBLOCK:SAPLKACB:1003/ctxtCOBL-VORNR_AUF").text = "0010"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 400',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_SERIAL").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_SERIAL/ssubSUB_TS_GOITEM_SERIAL:SAPLMIGO:0360/btnOK_SER_SEARCH").press',
                'WScript.Sleep 400',
                'session.findById("wnd[0]/usr/btn%_SERNR_%_APP_%-VALU_PUSH").press',
                'WScript.Sleep 300',
                f'CreateObject("htmlfile").parentWindow.clipboardData.setData "Text", {seriais_clip}',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[24]").press',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[0]").press',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[8]").press',
                'WScript.Sleep 400',
                'session.findById("wnd[0]/tbar[1]/btn[8]").press',
                'WScript.Sleep 500',
                'session.findById("wnd[0]/usr/cntlGRID1/shellcont/shell").setCurrentCell -1, ""',
                'session.findById("wnd[0]/usr/cntlGRID1/shellcont/shell").selectAll',
                'WScript.Sleep 300',
                'session.findById("wnd[0]/tbar[1]/btn[47]").press',
                'WScript.Sleep 500',
                '',
                '\' Marcar Item OK (confirma o item atual no SAP e evita sobreposição)',
                'On Error Resume Next',
                'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/subSUB_DETAIL_TAKE:SAPLMIGO:0304/chkGODYNPRO-DETAIL_TAKE").selected = True',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 400',
                'On Error GoTo 0',
            ])

            if idx_z62 + 1 < total_grupos_z62:
                linhas_vbs.extend([
                    '\' Avançar para o próximo item no MIGO Z62 sem sobrepor o anterior',
                    'On Error Resume Next',
                    'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/btnOK_NEXT_ITEM").press',
                    'If Err.Number <> 0 Then',
                    '    Err.Clear',
                    '    session.findById("wnd[0]/tbar[1]/btn[13]").press',
                    'End If',
                    'On Error GoTo 0',
                    'WScript.Sleep 600',
                    '\' Garante detalhe aberto para o próximo item',
                    'On Error Resume Next',
                    'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0302/btnBUTTON_DETAIL").press',
                    'On Error GoTo 0',
                    'WScript.Sleep 300',
                ])

            idx_z62 += 1

        # Finalização da Z62
        linhas_vbs.extend([
            'session.findById("wnd[0]/tbar[1]/btn[7]").press',
            'WScript.Sleep 600',
            'session.findById("wnd[0]/tbar[1]/btn[23]").press',
            'WScript.Sleep 1000',
            'Dim max_migo2, loop_migo2',
            'max_migo2 = 0',
            'Do While session.Children.Count > 1 And max_migo2 < 12',
            '    Err.Clear',
            '    session.findById("wnd[1]/tbar[0]/btn[0]").press',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]").sendVKey 0',
            '    End If',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]/usr/btnSPOP-OPTION1").press',
            '    End If',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]").close',
            '    End If',
            '    WScript.Sleep 350',
            '    max_migo2 = max_migo2 + 1',
            'Loop',
            'Dim msg_sbar_z62',
            'msg_sbar_z62 = session.findById("wnd[0]/sbar").text',
            'Dim doc_z62, re_z62, matches_z62',
            'doc_z62 = ""',
            'Set re_z62 = New RegExp',
            're_z62.Pattern = "\\b(49\\d{8}|50\\d{8}|\\d{10})\\b"',
            'Set matches_z62 = re_z62.Execute(msg_sbar_z62)',
            'If matches_z62.Count > 0 Then',
            '    doc_z62 = matches_z62(0).Value',
            '    WScript.Echo "DOC_Z62: " & doc_z62',
            'Else',
            '    Dim txtErrModal_z62',
            '    txtErrModal_z62 = ""',
            '    If session.Children.Count > 1 Then',
            '        On Error Resume Next',
            '        txtErrModal_z62 = session.findById("wnd[1]/usr/txtMESSTXT1").text',
            '        If txtErrModal_z62 = "" Then txtErrModal_z62 = session.findById("wnd[1]").text',
            '    End If',
            '    If txtErrModal_z62 <> "" Then',
            '        WScript.Echo "ERRO_SAP: " & txtErrModal_z62',
            '    Else',
            '        WScript.Echo "ERRO_SAP: Não foi possível capturar o doc Z62: " & msg_sbar_z62',
            '    End If',
            '    WScript.Quit 1',
            'End If',
            '',
            '\' =========================================================================',
            '\' ETAPA 2: MIGO Z61 (SAÍDA DO DEPÓSITO PARA A ORDEM DE DESTINO)',
            '\' =========================================================================',
            'session.findById("wnd[0]/tbar[0]/okcd").text = "/nMIGO"',
            'session.findById("wnd[0]").sendVKey 0',
            'WScript.Sleep 800',
            'FecharModais 8',
            '',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-ACTION").setFocus',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-ACTION").key = "A07"',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-REFDOC").setFocus',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/cmbGODYNPRO-REFDOC").key = "R10"',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/ctxtGODEFAULT_TV-BWART").text = "Z61"',
            'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_FIRSTLINE:SAPLMIGO:0010/ctxtGODEFAULT_TV-BWART").setFocus',
            'session.findById("wnd[0]").sendVKey 0',
            'WScript.Sleep 500',
        ])

        if dt_lcto_limpa:
            linhas_vbs.extend([
                f'PreencherDataLancamento carrier, "{dt_lcto_limpa}"',
                'WScript.Sleep 300',
            ])

        linhas_vbs.extend([
            ''
        ])

        # Itens da Z61
        dep_dest_clean = str(dep_dest or "CL03").strip()
        recebedor_dest = f"EMPR 1{dep_dest_clean}"
        idx_z61 = 0
        total_grupos_z61 = len(grupos_z61)

        for mat_cod, seriais_list in grupos_z61.items():
            qtd = str(len(seriais_list))

            linhas_vbs.extend([
                f"' --- Z61 Item {idx_z61 + 1}: Material {mat_cod} | Qtd {qtd} | OSME Destino {osme_dest} ---"
            ])

            seriais_clip_z61 = vb_escape_clipboard("\r\n".join(seriais_list))

            linhas_vbs.extend([
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_MATERIAL").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_MATERIAL/ssubSUB_TS_GOITEM_MATERIAL:SAPLMIGO:0310/ctxtGOITEM-MAKTX").text = "{mat_cod}"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 300',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_QUANTITIES").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_QUANTITIES/ssubSUB_TS_GOITEM_QUANTITIES:SAPLMIGO:0315/txtGOITEM-ERFMG").text = "{qtd}"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 300',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT.").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT./ssubSUB_TS_GOITEM_DESTINATION:SAPLMIGO:0325/ctxtGOITEM-NAME1").text = "EMPR"',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT./ssubSUB_TS_GOITEM_DESTINATION:SAPLMIGO:0325/ctxtGOITEM-LGOBE").text = "{dep_dest_clean}"',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_DESTINAT./ssubSUB_TS_GOITEM_DESTINATION:SAPLMIGO:0325/txtGOITEM-WEMPF").text = "{recebedor_dest}"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 300',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_ACCOUNT").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_ACCOUNT/ssubSUB_TS_GOITEM_ACCOUNT:SAPLMIGO:0345/ssubSUB_ACCOUNTINGBLOCK:SAPLKACB:1003/ctxtCOBL-AUFNR").text = "{osme_dest}"',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_ACCOUNT/ssubSUB_TS_GOITEM_ACCOUNT:SAPLMIGO:0345/ssubSUB_ACCOUNTINGBLOCK:SAPLKACB:1003/ctxtCOBL-VORNR_AUF").text = "010"',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 400',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_SERIAL").select',
                f'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/tabsTS_GOITEM/tabpOK_GOITEM_SERIAL/ssubSUB_TS_GOITEM_SERIAL:SAPLMIGO:0360/btnOK_SER_SEARCH").press',
                'WScript.Sleep 400',
                'session.findById("wnd[0]/usr/btn%_SERNR_%_APP_%-VALU_PUSH").press',
                'WScript.Sleep 400',
                f'CreateObject("htmlfile").parentWindow.clipboardData.setData "Text", {seriais_clip_z61}',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[24]").press',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[0]").press',
                'WScript.Sleep 300',
                'session.findById("wnd[1]/tbar[0]/btn[8]").press',
                'WScript.Sleep 400',
                'session.findById("wnd[0]/tbar[1]/btn[8]").press',
                'WScript.Sleep 500',
                'session.findById("wnd[0]/usr/cntlGRID1/shellcont/shell").setCurrentCell -1, ""',
                'session.findById("wnd[0]/usr/cntlGRID1/shellcont/shell").selectAll',
                'WScript.Sleep 300',
                'session.findById("wnd[0]/tbar[1]/btn[47]").press',
                'WScript.Sleep 500',
                '',
                '\' Marcar Item OK (confirma o item atual no SAP e evita sobreposição)',
                'On Error Resume Next',
                'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/subSUB_DETAIL_TAKE:SAPLMIGO:0304/chkGODYNPRO-DETAIL_TAKE").selected = True',
                'session.findById("wnd[0]").sendVKey 0',
                'WScript.Sleep 400',
                'On Error GoTo 0',
            ])

            if idx_z61 + 1 < total_grupos_z61:
                linhas_vbs.extend([
                    '\' Avançar para o próximo item no MIGO Z61 sem sobrepor o anterior',
                    'On Error Resume Next',
                    'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0301/subSUB_DETAIL:SAPLMIGO:0300/btnOK_NEXT_ITEM").press',
                    'If Err.Number <> 0 Then',
                    '    Err.Clear',
                    '    session.findById("wnd[0]/tbar[1]/btn[13]").press',
                    'End If',
                    'On Error GoTo 0',
                    'WScript.Sleep 600',
                    '\' Garante detalhe aberto para o próximo item',
                    'On Error Resume Next',
                    'session.findById("wnd[0]/usr/ssubSUB_MAIN_CARRIER:SAPLMIGO:" & carrier & "/subSUB_ITEMDETAIL:SAPLMIGO:0302/btnBUTTON_DETAIL").press',
                    'On Error GoTo 0',
                    'WScript.Sleep 300',
                ])

            idx_z61 += 1

        # Finalização da Z61
        linhas_vbs.extend([
            'session.findById("wnd[0]/tbar[1]/btn[7]").press',
            'WScript.Sleep 600',
            'session.findById("wnd[0]/tbar[1]/btn[23]").press',
            'WScript.Sleep 1000',
            'Dim max_migo4, loop_migo4',
            'max_migo4 = 0',
            'Do While session.Children.Count > 1 And max_migo4 < 12',
            '    Err.Clear',
            '    session.findById("wnd[1]/tbar[0]/btn[0]").press',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]").sendVKey 0',
            '    End If',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]/usr/btnSPOP-OPTION1").press',
            '    End If',
            '    If Err.Number <> 0 Then',
            '        Err.Clear',
            '        session.findById("wnd[1]").close',
            '    End If',
            '    WScript.Sleep 350',
            '    max_migo4 = max_migo4 + 1',
            'Loop',
            'Dim msg_sbar_z61',
            'msg_sbar_z61 = session.findById("wnd[0]/sbar").text',
            'Dim doc_z61, re_z61, matches_z61',
            'doc_z61 = ""',
            'Set re_z61 = New RegExp',
            're_z61.Pattern = "\\b(49\\d{8}|50\\d{8}|\\d{10})\\b"',
            'Set matches_z61 = re_z61.Execute(msg_sbar_z61)',
            'If matches_z61.Count > 0 Then',
            '    doc_z61 = matches_z61(0).Value',
            '    WScript.Echo "DOC_Z61: " & doc_z61',
            'Else',
            '    Dim txtErrModal_z61',
            '    txtErrModal_z61 = ""',
            '    If session.Children.Count > 1 Then',
            '        On Error Resume Next',
            '        txtErrModal_z61 = session.findById("wnd[1]/usr/txtMESSTXT1").text',
            '        If txtErrModal_z61 = "" Then txtErrModal_z61 = session.findById("wnd[1]").text',
            '    End If',
            '    If txtErrModal_z61 <> "" Then',
            '        WScript.Echo "ERRO_SAP: " & txtErrModal_z61',
            '    Else',
            '        WScript.Echo "ERRO_SAP: Não foi possível capturar o doc Z61: " & msg_sbar_z61',
            '    End If',
            '    WScript.Quit 1',
            'End If'
        ])

        return "\r\n".join(linhas_vbs)

    def _atualizar_planilha_transferencia(self, eq_dest, dep_dest, doc_z61, itens_liberados):
        """
        Regra da planilha:
        Apaga o registro anterior dos seriais em FORNECIDOS e reinsere com os novos dados:
        Equipe de Destino, Data de Hoje e Novo Doc SAP (Doc Z61).
        """
        seriais_transferidos_set = set(normalizar_texto(it["serial"]) for it in itens_liberados)
        data_hoje = datetime.now().strftime("%d/%m/%Y")
        sucesso = False
        modo = "openpyxl"

        # 0. Backup Preventivo Automático
        bak_criado = criar_backup_preventivo(self.caminho_planilha, prefixo="PRE_TRANSF_EQUIPES")
        if bak_criado:
            self._log(f"🛡️ Backup preventivo criado: {os.path.basename(bak_criado)}")

        # 1. Tentativa via Excel ativo (win32com)
        try:
            excel = win32com.client.GetActiveObject("Excel.Application")
            wb_ativo = None
            for w in excel.Workbooks:
                if os.path.abspath(w.FullName).lower() == os.path.abspath(self.caminho_planilha).lower():
                    wb_ativo = w
                    break

            if wb_ativo:
                modo = "Excel Aberto (win32com)"
                ws_forn = None
                for sh in wb_ativo.Sheets:
                    if sh.Name.upper() == "FORNECIDOS":
                        ws_forn = sh
                        break

                if ws_forn:
                    last_row = ws_forn.Cells(ws_forn.Rows.Count, 1).End(3).Row
                    # Remove linhas antigas de baixo para cima
                    for r in range(last_row, 1, -1):
                        val_serie = str(ws_forn.Cells(r, 1).Value or "")
                        if normalizar_texto(val_serie) in seriais_transferidos_set:
                            ws_forn.Rows(r).Delete()

                    # Anexa as novas linhas no final
                    new_last = ws_forn.Cells(ws_forn.Rows.Count, 1).End(3).Row
                    for it in itens_liberados:
                        new_last += 1
                        ws_forn.Cells(new_last, 1).Value = str(it["serial"]).strip()
                        ws_forn.Cells(new_last, 2).Value = str(it["material"]).strip()
                        ws_forn.Cells(new_last, 3).Value = str(dep_dest).strip()
                        ws_forn.Cells(new_last, 4).Value = str(eq_dest).strip()
                        ws_forn.Cells(new_last, 5).Value = data_hoje
                        ws_forn.Cells(new_last, 6).Value = str(doc_z61).strip()

                    wb_ativo.Save()
                    sucesso = True
        except Exception:
            pass

        # 2. Se o Excel não estiver aberto ou falhar, grava via openpyxl
        if not sucesso:
            try:
                modo = "Arquivo Fechado (openpyxl)"
                wb = load_workbook(self.caminho_planilha, keep_vba=self.caminho_planilha.lower().endswith(".xlsm"))
                if "FORNECIDOS" in wb.sheetnames:
                    ws_forn = wb["FORNECIDOS"]
                    idx_ser = None
                    for c_idx, val in enumerate(next(ws_forn.iter_rows(max_row=1, values_only=True), ()), start=1):
                        if normalizar_texto(val) in ("SERIE", "SERIAL", "SERNR"):
                            idx_ser = c_idx
                            break
                    if idx_ser is None:
                        idx_ser = 1

                    # Identifica e exclui as linhas antigas de baixo para cima
                    linhas_para_excluir = []
                    for r_idx, row in enumerate(ws_forn.iter_rows(min_row=2, values_only=True), start=2):
                        if row and len(row) >= idx_ser and row[idx_ser - 1]:
                            s_norm = normalizar_texto(row[idx_ser - 1])
                            if s_norm in seriais_transferidos_set:
                                linhas_para_excluir.append(r_idx)

                    for r_del in reversed(linhas_para_excluir):
                        ws_forn.delete_rows(r_del)

                    # Anexa os novos registros
                    for it in itens_liberados:
                        ws_forn.append([
                            str(it["serial"]).strip(),
                            str(it["material"]).strip(),
                            str(dep_dest).strip(),
                            str(eq_dest).strip(),
                            data_hoje,
                            str(doc_z61).strip()
                        ])

                    wb.save(self.caminho_planilha)
                    wb.close()
                    sucesso = True
            except Exception as e:
                self._log(f"⚠️ Erro ao atualizar Excel via openpyxl: {e}")

        # Se não foi atualizado com sucesso (planilha em uso na rede ou aberta em Somente Leitura)
        if not sucesso:
            motivo = "Planilha em uso por outro usuário no Excel da rede (Arquivo Bloqueado)"
            try:
                if wb_ativo is not None and getattr(wb_ativo, "ReadOnly", False):
                    motivo = "Planilha aberta no Excel local em modo 'Somente Leitura' (bloqueada na rede)"
            except Exception:
                pass

            pid = adicionar_pendencia(
                tipo="transferencia_equipes",
                caminho_planilha=self.caminho_planilha,
                dados_operacao={
                    "eq_dest": eq_dest,
                    "dep_dest": dep_dest,
                    "doc_z61": doc_z61,
                    "data": data_hoje,
                    "itens": [
                        {
                            "serial": it.get("serial"),
                            "material": it.get("material"),
                            "deposito": dep_dest
                        }
                        for it in itens_liberados
                    ]
                },
                motivo=motivo
            )
            modo = "Fila de Sincronização Automática (Planilha Bloqueada)"
            self._log(f"⏳ {motivo}. Operação #{pid} salva na fila de pendências para gravação automática.")
            self.parent_frame.after(0, lambda: self._atualizar_banner_pendencias())

        # Atualiza o cache em memória
        for it in itens_liberados:
            s_norm = normalizar_texto(it["serial"])
            self.dados_fornecidos[s_norm] = {
                "serial": it["serial"],
                "material": it["material"],
                "deposito": dep_dest,
                "equipe": eq_dest,
                "data": data_hoje,
                "doc_sap": doc_z61
            }

        # Salva contingência local em JSON
        try:
            pasta_dados = obter_diretorio_dados()
            caminho_cont = os.path.join(pasta_dados, "contingencia_transferencias_equipes.json")
            historico = []
            if os.path.exists(caminho_cont):
                try:
                    with open(caminho_cont, "r", encoding="utf-8") as f:
                        historico = json.load(f)
                except Exception:
                    historico = []

            historico.append({
                "data_hora": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "equipe_destino": eq_dest,
                "deposito_destino": dep_dest,
                "doc_z61": doc_z61,
                "total_itens": len(itens_liberados),
                "seriais": [
                    {
                        "serial": it["serial"],
                        "material": it["material"],
                        "origem": it["equipe_origem"],
                        "osme_origem": it["osme_origem"]
                    }
                    for it in itens_liberados
                ]
            })
            with open(caminho_cont, "w", encoding="utf-8") as f:
                json.dump(historico, f, indent=2, ensure_ascii=False)
        except Exception:
            pass

        return sucesso, modo

    def _iniciar_timer_pendencias(self):
        """Inicia o ciclo de verificação em segundo plano das pendências da planilha."""
        self._verificar_timer_pendencias()

    def _verificar_timer_pendencias(self):
        """Verifica a cada 20 segundos se a planilha foi liberada para gravar pendências."""
        try:
            qtd = obter_quantidade_pendencias()
            self._atualizar_banner_pendencias(qtd)
            if qtd > 0:
                sucessos, restantes, msg = processar_todas_pendencias()
                if sucessos > 0:
                    self._log(f"🎉 Sincronização Automática: {sucessos} operação(ões) gravada(s) na planilha com sucesso!")
                    self._status(f"🎉 {sucessos} operação(ões) gravada(s) na planilha automaticamente!")
                    self._atualizar_banner_pendencias(restantes)
                    self._iniciar_leitura_planilha_thread()
        except Exception:
            pass
        self.parent_frame.after(20000, self._verificar_timer_pendencias)

    def _atualizar_banner_pendencias(self, qtd=None):
        """Exibe ou oculta o banner de pendências com a contagem atual."""
        if qtd is None:
            qtd = obter_quantidade_pendencias()
        if not hasattr(self, "f_banner_pendencias"):
            return
        if qtd > 0:
            self.lbl_banner_pendencias.configure(
                text=f"⏳ Há {qtd} operação(ões) aguardando liberação da planilha pelo outro usuário."
            )
            if not self.f_banner_pendencias.winfo_ismapped():
                self.f_banner_pendencias.pack(fill="x", padx=14, pady=(0, 10))
        else:
            if self.f_banner_pendencias.winfo_ismapped():
                self.f_banner_pendencias.pack_forget()

    def _sincronizar_pendencias_manual(self):
        """Acionamento manual do botão de sincronização com feedback visual."""
        self.btn_sinc_banner.configure(text="⏳ Gravando...", state="disabled")
        self.parent_frame.update_idletasks()
        sucessos, restantes, msg = processar_todas_pendencias()
        self._atualizar_banner_pendencias(restantes)
        self.btn_sinc_banner.configure(text="🔄 Sincronizar Agora", state="normal")
        if sucessos > 0 and restantes == 0:
            messagebox.showinfo(
                "Sincronização Concluída",
                f"🎉 Sucesso total!\n\nForam gravados {sucessos} registro(s) pendente(s) na planilha com êxito!"
            )
            self._iniciar_leitura_planilha_thread()
        elif sucessos > 0 and restantes > 0:
            messagebox.showinfo(
                "Sincronização Parcial",
                f"✓ {sucessos} gravados com sucesso.\nRestam {restantes} que continuam em uso na rede."
            )
            self._iniciar_leitura_planilha_thread()
        else:
            messagebox.showwarning(
                "Planilha em Uso",
                "A planilha continua aberta por outro usuário na rede.\n\n"
                "O N.E.X.U.S continuará tentando gravar automaticamente a cada 20 segundos em segundo plano."
            )

    def _abrir_modal_pendencias(self):
        """Abre o modal para inspecionar, copiar dados ou descartar pendências."""
        ModalVisualizarPendencias(
            self.parent_frame,
            callback_apos_sincronizacao=lambda: [self._atualizar_banner_pendencias(), self._iniciar_leitura_planilha_thread()],
            paletas=self.paletas,
            obter_tema_atual=self.obter_tema_atual
        )

    def _concluir_sucesso_transferencia(self, doc_z62, doc_z61, origens, eq_dest, osme_dest, itens, sucesso_plan, modo_plan):
        self.processando_sap = False
        self._definir_estado_botoes_processar(executando=False)
        self._status(f"Transferência concluída! Z62: {doc_z62} | Z61: {doc_z61}")
        self._log(f"✅ Transferência finalizada com sucesso. Doc Z62: {doc_z62} | Doc Z61: {doc_z61}")

        # Se a opção de abrir WorkForce automaticamente estiver ativa e houver RE preenchido, dispara em segundo plano
        if getattr(self, "chk_abrir_wf", None) and self.chk_abrir_wf.get() and self.entry_re_eletricista.get().strip():
            self._status("Transferência concluída. Abrindo Ficha no WorkForce...")
            self.abrir_ficha_workforce_manual()

        ModalSucessoTransferenciaEquipes(
            self.parent_frame,
            doc_z62,
            doc_z61,
            origens,
            eq_dest,
            osme_dest,
            itens,
            sucesso_plan,
            modo_plan,
            self.paletas,
            self.obter_tema_atual,
            callback_abrir_workforce=self.abrir_ficha_workforce_manual
        )
        self._limpar_campos()

    def _finalizar_com_aviso_sap(self, alerta_msg):
        self.processando_sap = False
        self._definir_estado_botoes_processar(executando=False)
        self._status("Aviso retornado pelo SAP.")
        self._log(f"⚠️ Aviso do SAP: {alerta_msg}")
        messagebox.showwarning(
            "Transferência entre Equipes — Aviso SAP",
            f"O SAP retornou a seguinte mensagem na barra de status:\n\n{alerta_msg}\n\nNenhuma alteração foi realizada na planilha."
        )

    def _finalizar_com_erro(self, err_msg):
        self.processando_sap = False
        self._definir_estado_botoes_processar(executando=False)
        self._status("Erro durante execução no SAP.")
        self._log(f"❌ Erro SAP: {err_msg}")
        messagebox.showerror("Transferência entre Equipes — Erro", f"Ocorreu um erro durante a execução no SAP:\n\n{err_msg}")

    def preparar_transferencia_via_chat(self, destino=None, seriais=None, data_lancamento=None):
        """
        Preenche os dados na tela de Transferência entre Equipes e valida previamente.
        Retorna dicionário com status e resumo dos itens preparados.
        """
        if destino:
            self.entry_equipe_destino.delete(0, tk.END)
            self.entry_equipe_destino.insert(0, str(destino).strip().upper())
            self._consultar_equipe_destino_atual()

        if data_lancamento and hasattr(self, "entry_data_lancamento"):
            self.entry_data_lancamento.delete(0, tk.END)
            self.entry_data_lancamento.insert(0, str(data_lancamento).strip())

        if seriais:
            texto_seriais = "\n".join(seriais) if isinstance(seriais, list) else str(seriais)
            self.txt_seriais.delete("1.0", tk.END)
            self.txt_seriais.insert("1.0", texto_seriais)
            self.validar_seriais_acao()

        todos_itens = self.seriais_validados_atuais
        dest_atual = self.entry_equipe_destino.get().strip().upper()
        osme_dest = self.lbl_osme_numero.cget("text").strip()
        dep_dest = self.lbl_deposito_destino.cget("text").replace("Depósito:", "").strip()

        origens_set = set(it["equipe_origem"] for it in todos_itens if not it.get("eh_bloqueado"))
        origens_txt = ", ".join(sorted(origens_set)) if origens_set else "-"

        resumo_mats = []
        mats_map = OrderedDict()
        for it in todos_itens:
            if not it.get("eh_bloqueado"):
                m = it["material"]
                mats_map[m] = mats_map.get(m, 0) + 1
        for m, q in mats_map.items():
            resumo_mats.append(f"{m} — {q} un")

        return {
            "sucesso": True,
            "origens": origens_txt,
            "destino": dest_atual,
            "osme_destino": osme_dest,
            "deposito_destino": dep_dest,
            "total_itens": len(todos_itens),
            "resumo_materiais": resumo_mats,
            "tem_bloqueados": any(it.get("eh_bloqueado") for it in todos_itens)
        }

    def executar_transferencia_via_chat(self):
        """Dispara a transferência no SAP diretamente via confirmação no Chatbot."""
        itens_liberados = [it for it in self.seriais_validados_atuais if not it.get("eh_bloqueado")]
        if not itens_liberados:
            return {"erro": "Nenhum serial liberado para transferência na aba. Valide os seriais primeiro."}

        eq_dest = self.entry_equipe_destino.get().strip().upper()
        osme_dest = self.lbl_osme_numero.cget("text").strip()
        dep_dest = self.lbl_deposito_destino.cget("text").replace("Depósito:", "").strip() or "CL03"

        if not eq_dest or osme_dest in ("—", "-", "Sem OSME", "", "Não Encontrada"):
            return {"erro": "Equipe ou OSME de destino inválida na aba Transferência entre Equipes."}

        if not verificar_sap_aberto():
            return {"erro": "O SAP GUI precisa estar aberto e conectado em uma sessão ativa."}

        if self.processando_sap:
            return {"erro": "Já existe um processamento de transferência no SAP em andamento."}

        data_lcto = self.obter_data_lancamento() if hasattr(self, "entry_data_lancamento") else ""
        if data_lcto is None:
            data_lcto = ""

        self.processando_sap = True
        self._definir_estado_botoes_processar(executando=True)
        self._status("Processando Transferência no SAP MIGO (via Chatbot)...")
        origens_txt = ", ".join(sorted(set(it["equipe_origem"] for it in itens_liberados))) or "-"
        self._log(f"🚀 Iniciando transferência de {len(itens_liberados)} seriais: {origens_txt} ➡️ {eq_dest} (via Chatbot)...")
        if data_lcto:
            self._log(f"📅 Data de Lançamento configurada: {data_lcto}")

        threading.Thread(
            target=self._executar_transferencia_sap_thread,
            args=(eq_dest, osme_dest, dep_dest, itens_liberados, data_lcto),
            daemon=True
        ).start()
        return {
            "sucesso": True,
            "status": f"Execução de Transferência iniciada no SAP (Z62 ➡️ Z61) para {eq_dest} com {len(itens_liberados)} seriais."
        }

    def aplicar_tema(self, tema_nome):
        cores = self.paletas.get(tema_nome, self.paletas.get("Nexus", DEFAULT_PALETAS["Nexus"]))
        bg = cores.get("bg", "#090d16")
        container = cores.get("container", "#111827")
        inp = cores.get("input", "#1a2436")
        texto = cores.get("texto", "#f8fafc")
        accent = cores.get("accent", "#00d2ff")

        try:
            self.parent_frame.configure(fg_color=bg)
            if hasattr(self, "f_barra_fixa") and self.f_barra_fixa:
                self.f_barra_fixa.configure(fg_color=container)
            if hasattr(self, "frame_topo") and self.frame_topo:
                self.frame_topo.configure(fg_color=container)
            if hasattr(self, "frame_planilha") and self.frame_planilha:
                self.frame_planilha.configure(fg_color=container)
            if hasattr(self, "frame_entradas") and self.frame_entradas:
                self.frame_entradas.configure(fg_color=container)
            if hasattr(self, "card_osme_info") and self.card_osme_info:
                self.card_osme_info.configure(fg_color=inp)
            if hasattr(self, "card_origens") and self.card_origens:
                self.card_origens.configure(fg_color=inp)
            if hasattr(self, "lbl_osme_numero") and self.lbl_osme_numero:
                self.lbl_osme_numero.configure(text_color=accent)
        except Exception:
            pass

        estilo = ttk.Style()
        estilo.configure(
            "Treeview",
            background=container,
            foreground=texto,
            fieldbackground=inp,
            bordercolor=container,
            relief="flat"
        )
        estilo.map(
            "Treeview",
            background=[("selected", accent)],
            foreground=[("selected", "white")]
        )


def vb_escape_clipboard(texto):
    """Escapa quebras de linha para inserção segura no VBScript clipboard."""
    partes = texto.split("\r\n")
    if len(partes) == 1:
        return f'"{partes[0]}"'
    vbs_partes = [f'"{p}"' for p in partes]
    return " & vbCrLf & ".join(vbs_partes)


def construir_aba_transferencia_equipes(parent_frame, label_status_global=None, paletas=None, obter_tema_atual=None):
    return AbaTransferenciaEquipes(parent_frame, label_status_global, paletas, obter_tema_atual)
